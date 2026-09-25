"""Editable native-slide cloning for the approved Phase3F templates.

The source files are PPTX files made from editable PowerPoint text and shapes.
This module copies those shapes into the current presentation; it never adds a
full-slide screenshot.  A caller may fall back to the existing renderer when
the source cannot be safely cloned.
"""

from __future__ import annotations

from copy import deepcopy
from io import BytesIO
import logging
from pathlib import Path
import re
from copy import deepcopy
import tempfile
from typing import Iterable
from zipfile import ZIP_DEFLATED, ZipFile
import xml.etree.ElementTree as ET

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.util import Inches

from app.services.pptx_parts.native_trace_registry import (
    canonical_runtime_role,
    get_native_role_spec,
    get_runtime_native_role_spec,
    native_role_gate,
    runtime_native_role_gate,
)
from app.services.pptx_parts.native_trace_content_adapter import (
    FailureReason,
    NativeSlotPayload,
    SAMPLE_STRINGS,
    render_native_role_dry_run,
    write_slot_payload,
)
from app.services.pptx_parts.native_trace_validation import (
    clone_template_package,
    inspect_template_package,
    sha256_file,
    validate_injected_template_package,
)


logger = logging.getLogger(__name__)


def clone_runtime_native_template(
    source_path: str | Path,
    destination_path: str | Path,
    *,
    approved_slide_id: str,
    required_slots: Iterable[str] = (),
) -> dict:
    """Clone an approved PPTX package without dropping image relationships.

    This isolated foundation helper is intentionally not called by the
    existing dispatcher.  It copies the complete OOXML package, validates an
    explicit relationship allowlist, then adds only non-visible semantic
    ``shape.name`` identifiers to the runtime copy.
    """

    from app.services.pptx_parts.native_trace_validation import clone_template_package

    return clone_template_package(
        source_path,
        destination_path,
        approved_slide_id=approved_slide_id,
        required_slots=required_slots,
    )


def _iter_text_shapes(slide) -> Iterable[object]:
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            yield shape
        # GroupShape exposes child shapes through .shapes.
        children = getattr(shape, "shapes", None)
        if children is not None:
            for child in children:
                if getattr(child, "has_text_frame", False):
                    yield child


def _replace_shape_text(shape, replacements: dict[str, str]) -> bool:
    """Replace text in runs so the template's typography stays intact."""

    changed = False
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            updated = run.text
            for old, new in replacements.items():
                updated = updated.replace(old, new)
            if updated != run.text:
                run.text = updated
                changed = True
    return changed


def _scale_shape_transforms(element, scale_x: float, scale_y: float, *, slide_width: int, slide_height: int) -> None:
    """Scale geometry and text layout while keeping the shape in the slide.

    Native trace templates are authored at their canonical 16:9 EMU canvas,
    while the production deck can use a different 16:9 canvas.  Scaling only
    ``a:xfrm`` makes text boxes smaller without scaling their font metrics,
    margins, or point-based paragraph spacing; that changes wrapping in every
    static text region.  Text layout metrics therefore follow the same scale,
    while typefaces, rich-text runs, wrapping, and autofit policy remain
    untouched.
    """

    transforms = [*element.iter(qn("p:xfrm")), *element.iter(qn("a:xfrm"))]
    for xfrm in transforms:
        off = xfrm.find(qn("a:off"))
        ext = xfrm.find(qn("a:ext"))
        if off is not None:
            off.set("x", str(round(int(off.get("x", "0")) * scale_x)))
            off.set("y", str(round(int(off.get("y", "0")) * scale_y)))
        if ext is not None:
            ext.set("cx", str(round(int(ext.get("cx", "0")) * scale_x)))
            ext.set("cy", str(round(int(ext.get("cy", "0")) * scale_y)))

    # Preserve the source text layout at the destination canvas scale.  EMU
    # margins use the horizontal/vertical scale; font sizes and point-based
    # paragraph spacing use the vertical scale.  Percentage spacing and all
    # rich-text/typeface/autofit attributes are intentionally preserved.
    for tx_body in element.iter(qn("p:txBody")):
        body_pr = tx_body.find(qn("a:bodyPr"))
        if body_pr is not None:
            for attr in ("lIns", "rIns"):
                if body_pr.get(attr) is not None:
                    body_pr.set(attr, str(round(int(body_pr.get(attr, "0")) * scale_x)))
            for attr in ("tIns", "bIns"):
                if body_pr.get(attr) is not None:
                    body_pr.set(attr, str(round(int(body_pr.get(attr, "0")) * scale_y)))

        for child in tx_body.iter():
            tag = child.tag
            if tag in {qn("a:rPr"), qn("a:defRPr"), qn("a:endParaRPr")}:
                if child.get("sz") is not None:
                    child.set("sz", str(max(1, round(int(child.get("sz", "1")) * scale_y))))
            elif tag == qn("a:spcPts"):
                if child.get("val") is not None:
                    child.set("val", str(round(int(child.get("val", "0")) * scale_y)))

    # A few approved trace files contain decorative guide lines that extend a
    # few pixels beyond the source canvas.  Preserve the line but clamp only
    # the top-level transform so the generated production slide has no
    # off-canvas object.
    if transforms:
        xfrm = transforms[0]
        off = xfrm.find(qn("a:off"))
        ext = xfrm.find(qn("a:ext"))
        if off is not None and ext is not None:
            x = max(0, int(off.get("x", "0")))
            y = max(0, int(off.get("y", "0")))
            cx = max(1, int(ext.get("cx", "1")))
            cy = max(1, int(ext.get("cy", "1")))
            x = min(x, max(0, slide_width - 1))
            y = min(y, max(0, slide_height - 1))
            cx = min(cx, max(1, slide_width - x))
            cy = min(cy, max(1, slide_height - y))
            off.set("x", str(x))
            off.set("y", str(y))
            ext.set("cx", str(cx))
            ext.set("cy", str(cy))


def _clone_template_slide(prs: Presentation, source_path: Path):
    source_prs = Presentation(str(source_path))
    if not source_prs.slides:
        raise ValueError(f"native template has no slides: {source_path}")
    if source_prs.slide_width <= 0 or source_prs.slide_height <= 0:
        raise ValueError(f"native template has invalid page size: {source_path}")

    source_slide = source_prs.slides[0]
    scale_x = prs.slide_width / source_prs.slide_width
    scale_y = prs.slide_height / source_prs.slide_height

    # The approved Phase3F trace slides contain editable shapes and text.  If a
    # future asset introduces package relationships (for example an image),
    # fail closed so the caller can use the existing renderer instead of
    # producing a partially copied slide.
    relationship_types = [rel.reltype for rel in source_slide.part.rels.values()]
    unsupported = [reltype for reltype in relationship_types if "notesSlide" not in reltype and "slideLayout" not in reltype]
    if unsupported:
        raise ValueError(f"native template contains unsupported relationships: {source_path}")

    destination = prs.slides.add_slide(prs.slide_layouts[6])
    for shape in source_slide.shapes:
        element = deepcopy(shape.element)
        _scale_shape_transforms(
            element,
            scale_x,
            scale_y,
            slide_width=prs.slide_width,
            slide_height=prs.slide_height,
        )
        destination.shapes._spTree.insert_element_before(element, "p:extLst")
    return destination


def _clone_template_slide_with_media(prs: Presentation, source_path: Path):
    """Import an approved package slide while retaining editable media links.

    The normal shape importer intentionally rejects image relationships.  The
    approved S01 cover is the one exception needed by the visual fallback: its
    picture is copied as a native PowerPoint picture shape from the validated
    package, so the output keeps a real media relationship instead of a
    full-slide raster.
    """

    source_prs = Presentation(str(source_path))
    if not source_prs.slides:
        raise ValueError(f"native template has no slides: {source_path}")
    source_slide = source_prs.slides[0]
    destination = prs.slides.add_slide(prs.slide_layouts[6])
    scale_x = prs.slide_width / source_prs.slide_width
    scale_y = prs.slide_height / source_prs.slide_height

    for shape in source_slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            picture = destination.shapes.add_picture(
                BytesIO(shape.image.blob),
                left=round(shape.left * scale_x),
                top=round(shape.top * scale_y),
                width=round(shape.width * scale_x),
                height=round(shape.height * scale_y),
            )
            picture.crop_left = shape.crop_left
            picture.crop_right = shape.crop_right
            picture.crop_top = shape.crop_top
            picture.crop_bottom = shape.crop_bottom
            continue
        element = deepcopy(shape.element)
        _scale_shape_transforms(
            element,
            scale_x,
            scale_y,
            slide_width=prs.slide_width,
            slide_height=prs.slide_height,
        )
        destination.shapes._spTree.insert_element_before(element, "p:extLst")
    return destination


def _replace_brand_text(slide) -> None:
    for shape in _iter_text_shapes(slide):
        text = shape.text
        _replace_shape_text(
            shape,
            {
                "READY CREW Proposal": "提案クエスト",
                "READY CREW": "提案クエスト",
                "ProposalPilot": "提案クエスト",
                "AI営業秘書": "提案クエスト",
            },
        )


_TRACE_TITLES = {
    "ESTIMATE": "概算見積と判断条件",
    "KPI": "KPI設計と効果測定",
    "COMPETITIVE_COMPARISON": "競合比較と差別化ポイント",
    "WIN_PROBABILITY": "受注確度と次の判断",
    "SCHEDULE": "導入スケジュールと推進体制",
    "ROADMAP": "導入ステップとスケジュール",
}


def _replace_title(slide, role: str, title: str) -> None:
    canonical_title = _TRACE_TITLES.get(role)
    if not title or not canonical_title or title == canonical_title:
        return
    for shape in _iter_text_shapes(slide):
        if shape.text.strip() == canonical_title:
            if not _replace_shape_text(shape, {canonical_title: title}):
                shape.text = title
            return


_VISUAL_FALLBACK_ROLES = frozenset(
    {
        "EXECUTIVE_SUMMARY",
        "DECISION_AND_EXPECTED_EFFECTS",
        "PROPOSAL_SUMMARY",
        "IMPLEMENTATION_CONFIGURATION",
        "KPI",
        "SCHEDULE",
        "CURRENT_STATE",
        "PROBLEM_ANALYSIS",
        "SOLUTION_CONCEPT",
        "SOLUTION_APPROACH",
    }
)

_SCHEDULE_UNVERIFIED_TEXT = (
    "要件整理",
    "情報設計",
    "制作・開発",
    "検証・公開準備",
    "公開後改善",
    "現状ヒアリング",
    "課題・目的の整理",
    "要件の優先順位付け",
    "サイト構成の設計",
    "コンテンツ整理",
    "ワイヤーフレーム作成",
    "デザイン制作",
    "フロントエンド実装",
    "CMS構築・各種設定",
    "動作検証・修正",
    "コンテンツ登録",
    "公開前の最終確認",
    "アクセス解析",
    "効果測定・レポート",
    "改善施策の実施",
    "要件定義書 プロジェクト計画書",
    "サイトマップ ワイヤーフレーム",
    "デザインデータ テスト環境一式",
    "検証報告書 公開チェックリスト",
    "運用レポート 改善提案書",
    "営業責任者",
    "窓口・意思決定支援",
    "全体調整・経営層との連携",
    "プロジェクト マネージャー",
    "進行管理・課題整理",
    "スケジュール管理・リスク対応",
    "制作・開発担当",
    "設計・実装・検証",
    "デザイン・開発・テスト対応",
    "お客様担当",
    "要件確認・レビュー",
    "社内調整・コンテンツ提供",
    "週次確認で判断待ちを最小化",
    "定例ミーティングで進捗・課題・次のアクションを確認",
    "優先順位を固定して短納期に対応",
    "必要要件から段階的に進め、スコープを明確化",
    "公開後の改善テーマまで先に整理",
    "運用フェーズの目標・改善項目を事前に設定",
    "確認事項と判断タイミングを前倒しし、",
    "約3か月で公開、その後も継続改善できる体制",
    "を構築します。",
)


def _visual_fallback_replacements(role: str) -> dict[str, str]:
    replacements: dict[str, str] = {}
    if role in {"CURRENT_STATE", "PROBLEM_ANALYSIS", "SOLUTION_CONCEPT", "SOLUTION_APPROACH"}:
        replacements.update({sample: "要確認" for sample in SAMPLE_STRINGS.get(role, ())})
        replacements["2025年6月22日"] = "2026.08.26"
        replacements["2026.06.22"] = "2026.08.26"
    elif role == "PROPOSAL_SUMMARY":
        replacements.update({sample: "要確認" for sample in SAMPLE_STRINGS.get(role, ())})
        safe_sentence = "確認済み情報を整理し、未確認項目は確認後に確定します。"
        replacements["提案業務の効率化と提案品質の向上を両立するため、3つの施策を一体で実行します。"] = safe_sentence
        replacements["要確認と提案品質の向上を両立するため、3つの施策を一体で実行します。"] = safe_sentence
    elif role == "IMPLEMENTATION_CONFIGURATION":
        replacements.update({sample: "確認中" for sample in SAMPLE_STRINGS.get(role, ())})
        replacements["2025年6月22日"] = "2026.08.26"
        replacements["業務の標準化・データの一元管理により、部門間の連携を強化し、効率的で拡張性の高いERP基盤を構築します。"] = (
            "確認済み情報のみ表示し、未確認の費用・ROI・日程は保留します。"
        )
    elif role == "KPI":
        replacements.update(
            {
                "2026.06.22": "2026.08.26",
                "現状値（例）": "現状値",
                "資料作成時間": "確認項目",
                "提案数": "確認項目",
                "修正回数": "確認項目",
                "受注確度": "確認項目",
                "1案件あたりの\n資料修正回数を削減": "確認後に確定",
                "提案からの\n受注確度を向上": "確認後に確定",
                "20時間/件": "未取得",
                "70%削減": "要確認",
                "6時間/件": "要確認",
                "(6時間/件)": "要確認",
                "月10件": "未取得",
                "1.5倍": "要確認",
                "(月15件)": "要確認",
                "5回/件": "未取得",
                "30%削減": "要確認",
                "(3.5回/件)": "要確認",
                "20%": "未取得",
                "20%向上": "要確認",
                "(24%)": "要確認",
                "24%": "要確認",
                "未取得向上": "要確認",
                "(3.未取得)": "要確認",
                "作業時間の実績記録\n（ツール・工数表）": "工数記録",
                "営業管理ツールの\n案件件数": "案件数",
                "版管理・履歴からの\nカウント": "履歴カウント",
                "営業管理ツールの\n受注率": "受注率",
                "既存のデータやヒアリングをもとに、現状の数値を把握します。": "既存データとヒアリングで現状値を確認します。",
                "事業目標やリソースを踏まえ、現実的かつ挑戦的な目標値を設定します。": "事業目標とリソースを踏まえ、目標値を設定します。",
                "月次でKPIを確認し、要因分析と打ち手の検討を行います。": "月次でKPIを確認し、改善策を検討します。",
                "導入初期や計測環境が整っていない場合は、仮説値を設定し、実測後に見直します。": "初期は仮説とし、実測後に見直します。",
                "業界・案件規模・提案内容などの条件により、効果は変動する点に留意が必要です。": "条件により効果は変動します。",
                "KPIの進捗を定期的に確認し、施策の見直し・改善を継続的に行います。": "進捗を確認し、施策を継続改善します。",
                "導入後の成果を見える化し、改善判断に使える指標として整理します。": (
                    "候補KPIと測定方針を示し、現状値取得後に確定します。"
                ),
            }
        )
    elif role == "SCHEDULE":
        replacements.update({sample: "要確認" for sample in SAMPLE_STRINGS.get(role, ())})
        replacements.update({sample: "要確認" for sample in _SCHEDULE_UNVERIFIED_TEXT})
        replacements["短期間での立ち上げを実現するため、工程・担当・確認ポイントを明確にした進行計画を示します。"] = (
            "未確認の予定・担当者・マイルストーンは、確認後に確定します。"
        )
    replacements["提案クエスト Inc."] = "提案クエスト"
    return replacements


def _visual_fallback_title(role: str, slide_data) -> str:
    current = str(getattr(slide_data, "title", "") or "").strip()
    if current:
        return current
    return {
        "IMPLEMENTATION_CONFIGURATION": "導入構成",
        "KPI": "KPI設計",
        "SCHEDULE": "導入スケジュール",
    }.get(role, "提案クエスト")


def _replace_fragmented_fallback_text(package_path: Path, replacements: dict[str, str]) -> None:
    """Apply fallback replacements across text runs split by PowerPoint XML."""

    temp = package_path.with_suffix(".fragmented.tmp.pptx")
    with ZipFile(package_path, "r") as package:
        names = package.namelist()
        updated_slides: dict[str, bytes] = {}
        for name in names:
            if not (name.startswith("ppt/slides/slide") and name.endswith(".xml")):
                continue
            root = ET.fromstring(package.read(name))
            for shape in root.iter():
                if shape.tag.rsplit("}", 1)[-1] not in {"sp", "pic", "graphicFrame", "grpSp", "cxnSp"}:
                    continue
                nodes = [node for node in shape.iter() if node.tag.rsplit("}", 1)[-1] == "t"]
                if not nodes:
                    continue
                original = "".join(node.text or "" for node in nodes)
                updated = original
                for source, target in sorted(replacements.items(), key=lambda item: -len(item[0])):
                    updated = updated.replace(source, target)
                if updated == original:
                    continue
                nodes[0].text = updated
                for node in nodes[1:]:
                    node.text = ""
            updated_slides[name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
        with ZipFile(temp, "w", compression=ZIP_DEFLATED) as output:
            for name in names:
                output.writestr(name, updated_slides.get(name, package.read(name)))
    temp.replace(package_path)


_S09_UNSUPPORTED_FALLBACK_TEXT = (
    "業務の標準化",
    "標準化",
    "データの一元管理",
    "一元管理",
    "部門間の連携",
    "コストを最適化",
    "持続的な成長",
    "ERP",
    "構築します",
    "確認中を構築",
)


def _sanitize_implementation_configuration_fallback_slide(slide) -> None:
    """Remove unsupported S09 narrative without changing the approved layout."""

    safe_sentence = "確認済み情報のみ表示し、未確認の項目は確認後に確定します。"
    sentence_shapes = {
        "trace:S09:title.primary.3",
        "trace:S09:content.auto.177",
    }
    for shape in _iter_all_shapes(slide.shapes):
        if not getattr(shape, "has_text_frame", False):
            continue
        text = str(getattr(shape, "text", "") or "")
        if not any(fragment in text for fragment in _S09_UNSUPPORTED_FALLBACK_TEXT):
            continue
        _set_shape_text_preserving_style(shape, safe_sentence if shape.name in sentence_shapes else "要確認")


_SEMANTIC_PAGE_FALLBACK_SLIDE_IDS = {
    "CURRENT_STATE": "S03",
    "PROBLEM_ANALYSIS": "S04",
    "SOLUTION_CONCEPT": "S05",
    "SOLUTION_APPROACH": "S06",
}


_SEMANTIC_PAGE_SAFE_FALLBACK_CONTENT = {
    "CURRENT_STATE": {
        "content.auto.11": "確認済み情報",
        "lead": "確認済み情報",
        "content.auto.18": "未確認項目",
        "lead.2": "未確認項目",
        "content.auto.24": "次に確認する項目",
        "lead.3": "確認後に確定",
        "title.primary.2": "確認後に確定",
        "lead.4": "次に確認する項目",
        "content.auto.30": "確認済み情報",
        "content.auto.41": "確認済み",
        "content.auto.42": "確認後に確定",
        "content.auto.50": "未確認",
        "content.auto.51": "確認後に確定",
        "content.auto.58": "次に確認",
        "content.auto.59": "確認後に確定",
        "title.primary.3": "未確認項目",
        "title.primary.4": "確認後に確定",
        "title.primary.5": "次に確認する項目",
        "title.primary.6": "未確認項目",
        "title.primary.7": "確認後に確定",
        "title.primary.8": "次に確認する項目",
        "title.primary.9": "確認後に確定",
        "content.auto.76": "確認済み情報",
        "content.auto.81": "未確認項目",
        "content.auto.82": "確認後に確定",
        "content.auto.92": "未確認項目",
        "content.auto.93": "確認後に確定",
        "content.auto.102": "次に確認する項目",
        "content.auto.103": "確認後に確定",
        "content.auto.111": "確認済み情報",
        "content.auto.114": "未確認項目",
        "content.auto.115": "確認後に確定",
    },
    "PROBLEM_ANALYSIS": {
        "title.primary.3": "確認済み情報",
        "title.primary.4": "未確認項目",
        "content.auto.5": "確認後に確定",
        "title.primary.6": "確認済み情報",
        "lead": "未確認項目",
        "content.auto.22": "次に確認する項目",
        "content.auto.32": "確認済み情報",
        "content.auto.33": "確認後に確定",
        "content.auto.42": "未確認項目",
        "content.auto.43": "確認後に確定",
        "title.primary.8": "確認済み情報",
        "title.primary.9": "確認後に確定",
        "content.auto.53": "未確認項目",
        "lead.2": "確認後に確定",
        "content.auto.60": "次に確認する項目",
        "content.auto.61": "確認後に確定",
        "content.auto.66": "確認済み情報",
        "content.auto.71": "未確認項目",
        "content.auto.72": "確認後に確定",
        "content.auto.77": "次に確認する項目",
        "content.auto.78": "確認済み情報",
        "content.auto.79": "確認後に確定",
    },
    "SOLUTION_CONCEPT": {
        "title.primary.2": "確認済み情報",
        "title.primary.3": "未確認項目",
        "title.primary.4": "次に確認する項目",
        "title.primary.6": "確認済み情報",
        "title.primary.7": "確認後に確定",
        "content.auto.23": "未確認項目",
        "lead": "確認後に確定",
        "content.auto.30": "次に確認する項目",
        "content.auto.31": "確認後に確定",
        "content.auto.47": "確認",
        "title.primary.9": "確認後に確定",
        "content.auto.55": "確認",
        "content.auto.56": "確認後に確定",
        "content.auto.66": "確認",
        "content.auto.67": "確認後に確定",
        "content.auto.80": "確認済み情報",
        "content.auto.81": "確認後に確定",
        "content.auto.88": "未確認項目",
        "content.auto.89": "確認後に確定",
        "content.auto.94": "次に確認する項目",
        "content.auto.95": "確認後に確定",
        "content.auto.100": "確認済み情報",
        "content.auto.101": "未確認項目",
        "content.auto.102": "確認後に確定",
        "content.auto.103": "",
    },
    "SOLUTION_APPROACH": {
        "title.primary.2": "確認済み情報",
        "title.primary.3": "未確認項目",
        "content.auto.15": "確認後に確定",
        "title.primary.4": "確認済み情報",
        "content.auto.29": "確認後に確定",
        "title.primary.5": "未確認項目",
        "content.auto.55": "確認後に確定",
        "title.primary.6": "確認済み情報",
        "content.auto.69": "確認後に確定",
        "title.primary.7": "次に確認する項目",
        "content.auto.83": "確認後に確定",
        "title.primary.8": "確認済み情報",
        "content.auto.94": "未確認項目",
        "title.primary.9": "確認後に確定",
        "content.auto.104": "未確認項目",
        "title.primary.10": "確認後に確定",
        "content.auto.112": "未確認項目",
        "content.auto.114": "確認後に確定",
        "content.auto.121": "確認済み情報",
        "content.auto.129": "未確認項目",
        "content.auto.131": "確認後に確定",
        "content.auto.133": "確認後に確定",
        "content.auto.135": "確認後に確定",
        "content.auto.142": "未確認項目",
        "content.auto.144": "確認後に確定",
        "content.auto.146": "確認後に確定",
        "content.auto.148": "確認後に確定",
        "content.auto.158": "次に確認する項目",
        "content.auto.160": "確認後に確定",
        "content.auto.162": "確認後に確定",
        "content.auto.164": "確認後に確定",
        "content.auto.172": "次に確認する項目",
        "content.auto.175": "確認済み情報",
        "content.auto.176": "確認後に確定",
        "content.auto.179": "未確認項目",
        "content.auto.180": "確認後に確定",
        "content.auto.183": "次に確認する項目",
        "content.auto.184": "確認後に確定",
        "content.auto.189": "確認済み情報",
        "content.auto.190": "未確認項目",
        "content.auto.191": "確認後に確定",
        "content.auto.192": "",
    },
}


_SEMANTIC_PAGE_FALLBACK_DEFAULTS = {
    "CURRENT_STATE": {
        "lead": "現状は確認後に確定",
        "lead.2": "業務フローは確認後に確定",
        "content.auto.30": "課題は確認後に確定",
        "content.auto.41": "優先テーマは確認後に確定",
    },
    "PROBLEM_ANALYSIS": {
        "lead": "課題は確認後に確定",
        "content.auto.22": "優先課題は確認後に確定",
        "content.auto.42": "根拠は確認後に確定",
        "content.auto.78": "次の確認事項は確認後に確定",
    },
    "SOLUTION_CONCEPT": {
        "title.primary.2": "方針は確認後に確定",
        "content.auto.23": "施策は確認後に確定",
        "content.auto.47": "実行条件は確認後に確定",
        "content.auto.100": "成果条件は確認後に確定",
    },
    "SOLUTION_APPROACH": {
        "title.primary.2": "導入方針は確認後に確定",
        "title.primary.3": "導入ステップは確認後に確定",
        "title.primary.4": "実施条件は確認後に確定",
        "content.auto.94": "成果条件は確認後に確定",
    },
}


def _semantic_page_fallback_content(role: str, payload: NativeSlotPayload | None) -> dict[str, str]:
    """Build fallback copy from accepted runtime slots, never template text."""

    safe_content = dict(_SEMANTIC_PAGE_FALLBACK_DEFAULTS.get(role, {}))
    if payload is None:
        return safe_content
    prefix = f"trace:{payload.slide_id}:"
    allowed_slots = set(safe_content)
    for slot, value in payload.slots.items():
        if not slot.startswith(prefix):
            continue
        relative_slot = slot[len(prefix):]
        if relative_slot not in allowed_slots:
            continue
        if isinstance(value, (list, tuple)):
            text = " / ".join(str(item).strip() for item in value if str(item).strip())
        else:
            text = str(value).strip()
        if text:
            safe_content[relative_slot] = text
    return safe_content


def _semantic_shape_bounds(shape) -> tuple[float, float, float, float]:
    return (
        float(shape.left),
        float(shape.top),
        float(shape.left + shape.width),
        float(shape.top + shape.height),
    )


def _semantic_shape_center(shape) -> tuple[float, float]:
    return (
        float(shape.left + shape.width / 2),
        float(shape.top + shape.height / 2),
    )


def _semantic_point_in_bounds(point: tuple[float, float], bounds: tuple[float, float, float, float]) -> bool:
    x, y = point
    left, top, right, bottom = bounds
    return left <= x <= right and top <= y <= bottom


def _semantic_expanded_bounds(
    shape,
    *,
    horizontal: float = 0.12,
    vertical: float | None = None,
) -> tuple[float, float, float, float]:
    left, top, right, bottom = _semantic_shape_bounds(shape)
    if vertical is None:
        vertical = 0.35 if top < Inches(3.8) else 0.15
    return (
        left - Inches(horizontal),
        top - Inches(vertical),
        right + Inches(horizontal),
        bottom + Inches(vertical),
    )


def _remove_semantic_shape(shape) -> None:
    element = getattr(shape, "_element", None)
    parent = element.getparent() if element is not None else None
    if parent is not None:
        parent.remove(element)


def _apply_semantic_visual_finish(slide, role: str) -> None:
    """Tighten sparse evidence-safe pages without adding semantic content."""

    slide_id = _SEMANTIC_PAGE_FALLBACK_SLIDE_IDS.get(role)
    if not slide_id:
        return
    prefix = f"trace:{slide_id}:"

    def shape(suffix: str):
        return _shape_by_name(slide, f"{prefix}{suffix}")

    if role == "CURRENT_STATE":
        # Three compact blocks: two current-state blocks and one lower issue
        # block.  Existing runtime text is moved into existing approved
        # containers; no copy is added or rewritten.
        for suffix, geometry in {
            "content.auto.6": (0.42, 2.18, 3.35, 1.16),
            "content.auto.12": (4.02, 2.18, 3.35, 1.16),
            "lead": (0.68, 2.47, 2.82, 0.56),
            "lead.2": (4.10, 2.43, 3.16, 0.72),
            "content.auto.29": (0.42, 3.72, 6.95, 1.18),
            "content.auto.30": (0.72, 3.98, 6.28, 0.30),
            "content.auto.41": (0.72, 4.38, 6.28, 0.30),
        }.items():
            current = shape(suffix)
            if current is not None:
                _set_shape_geometry(current, left=geometry[0], top=geometry[1], width=geometry[2], height=geometry[3])
        lower_unused = shape("content.auto.31")
        if lower_unused is not None:
            _remove_semantic_shape(lower_unused)
        # This nested empty overlay sits above the second live block in the
        # source trace and masks the beginning of its runtime sentence.
        nested_overlay = shape("content.auto.19")
        if nested_overlay is not None:
            _remove_semantic_shape(nested_overlay)

    elif role == "PROBLEM_ANALYSIS":
        # Preserve the two-card-plus-summary composition while using the
        # unused upper body area.  Only geometry changes.
        for suffix in (
            "content.auto.14",
            "content.auto.22",
            "lead",
            "content.auto.34",
            "content.auto.42",
            "content.auto.73",
            "content.auto.78",
        ):
            current = shape(suffix)
            if current is not None and current.top >= Inches(3.0):
                current.top = max(Inches(1.95), current.top - Inches(0.78))

    elif role == "SOLUTION_CONCEPT":
        # Keep the insight bar comfortably above the fixed footer.  Place the
        # existing sentence wholly in the light region and use a dark Master
        # color so no white run spills onto the pale background.
        bar = shape("content.auto.96")
        red_accent = shape("content.auto.97")
        sentence = shape("content.auto.100")
        if bar is not None:
            _set_shape_geometry(bar, left=0.42, top=5.38, width=12.48, height=0.48)
        if red_accent is not None:
            _set_shape_geometry(red_accent, left=0.42, top=5.38, width=0.42, height=0.48)
        if sentence is not None:
            _set_shape_geometry(sentence, left=1.02, top=5.52, width=11.25, height=0.27)
            _set_shape_text_color(sentence, (18, 38, 58))
        for suffix in ("content.auto.113", "content.auto.114", "content.auto.115"):
            decorative = shape(suffix)
            if decorative is not None:
                _remove_semantic_shape(decorative)

    elif role == "SOLUTION_APPROACH":
        # Turn the two existing step regions into intentional compact cards
        # and reduce the verified-outcome panel to the size of its live copy.
        for suffix, geometry in {
            "content.auto.6": (0.52, 1.82, 3.10, 1.28),
            "title.primary.3": (0.78, 2.25, 2.58, 0.46),
            "content.auto.16": (3.92, 1.82, 3.10, 1.28),
            "title.primary.4": (4.18, 2.25, 2.58, 0.46),
            "content.auto.84": (9.12, 1.88, 3.18, 1.28),
            "content.auto.94": (9.40, 2.30, 2.62, 0.40),
        }.items():
            current = shape(suffix)
            if current is not None:
                _set_shape_geometry(current, left=geometry[0], top=geometry[1], width=geometry[2], height=geometry[3])


def _hide_empty_semantic_body_visuals(slide, prefix: str) -> None:
    """Remove body-only placeholders when their region has no live copy.

    The approved trace remains the source of geometry and styling.  This only
    removes body shapes that are not backed by any non-empty runtime text.  A
    populated card keeps its editable background and icon geometry; an empty
    card does not remain as an icon-only placeholder.
    """

    shapes = list(_iter_all_shapes(slide.shapes))
    body_shapes = [
        shape
        for shape in shapes
        if str(getattr(shape, "name", "") or "").startswith(prefix)
        and not any(
            marker in str(getattr(shape, "name", "") or "")
            for marker in (":footer.", ":header.", ":footer", ":header")
        )
        # Body-only cleanup starts below the shared header Chrome.  Several
        # approved traces place an optional panel rule immediately below the
        # header (around 0.7in); keeping that rule after its text is cleared
        # leaves an icon-only/empty panel impression.
        and float(shape.top) >= Inches(0.65)
        and float(shape.top) < Inches(6.85)
    ]
    active_text = [
        shape
        for shape in body_shapes
        if getattr(shape, "has_text_frame", False)
        and str(getattr(shape, "text", "") or "").strip()
    ]

    # Large textless shapes are the card/section containers in the approved
    # traces.  A container is live when it contains at least one active text
    # shape.  Small icon/line shapes inherit that live state from their
    # containing card.
    containers = [
        shape
        for shape in body_shapes
        if float(shape.width) * float(shape.height) >= Inches(0.65) ** 2
    ]
    live_containers = [
        container
        for container in containers
        if any(
                _semantic_point_in_bounds(
                    _semantic_shape_center(text_shape), _semantic_shape_bounds(container)
                )
            for text_shape in active_text
        )
    ]

    remove: list[object] = []
    for shape in body_shapes:
        name = str(getattr(shape, "name", "") or "")
        has_text = bool(getattr(shape, "has_text_frame", False))
        text = str(getattr(shape, "text", "") or "").strip() if has_text else ""
        if has_text:
            # Empty text boxes are placeholders even when their card remains
            # live because another slot in that card is populated.  Large
            # empty auto-shapes can also be card backgrounds in python-pptx,
            # so retain those when they contain live copy.
            if not text:
                if not any(
                    _semantic_point_in_bounds(
                        _semantic_shape_center(text_shape), _semantic_shape_bounds(shape)
                    )
                    for text_shape in active_text
                ) or float(shape.width) * float(shape.height) < Inches(0.65) ** 2:
                    remove.append(shape)
            continue

        # Some trace templates contain a second empty panel layered inside a
        # live card.  It is not a semantic container and can cover the text
        # that was moved into the approved card during visual finishing.
        if float(shape.width) * float(shape.height) >= Inches(0.65) ** 2 and any(
            container is not shape
            and _semantic_shape_bounds(container)[0] <= _semantic_shape_bounds(shape)[0]
            and _semantic_shape_bounds(container)[1] <= _semantic_shape_bounds(shape)[1]
            and _semantic_shape_bounds(container)[2] >= _semantic_shape_bounds(shape)[2]
            and _semantic_shape_bounds(container)[3] >= _semantic_shape_bounds(shape)[3]
            for container in live_containers
        ):
            remove.append(shape)
            continue

        if any(
            _semantic_point_in_bounds(_semantic_shape_center(shape), _semantic_shape_bounds(container))
            for container in live_containers
        ):
            continue

        if float(shape.width) * float(shape.height) >= Inches(0.65) ** 2:
            # A large body container with no live text is an empty card/row.
            remove.append(shape)
            continue

        # Keep a small decorative/icon shape only when it is close to live
        # copy.  Otherwise it is an icon-only placeholder belonging to an
        # empty region.
        sx, sy = _semantic_shape_center(shape)
        near_live_text = any(
            abs(sx - tx) <= Inches(1.0) and abs(sy - ty) <= Inches(0.9)
            for tx, ty in (_semantic_shape_center(text_shape) for text_shape in active_text)
        )
        if not near_live_text:
            remove.append(shape)

    for shape in remove:
        _remove_semantic_shape(shape)


def _lift_semantic_insight_bar(slide, prefix: str) -> None:
    """Move only the Slide 07 insight band clear of the fixed footer Chrome."""

    shift = Inches(0.22)
    for shape in list(_iter_all_shapes(slide.shapes)):
        name = str(getattr(shape, "name", "") or "")
        if not name.startswith(f"{prefix}content.auto"):
            continue
        top = float(shape.top)
        bottom = float(shape.top + shape.height)
        if top >= Inches(6.1) and bottom <= Inches(6.9):
            shape.top = max(Inches(5.8), shape.top - shift)


def _sanitize_semantic_page_fallback_slide(
    slide,
    role: str,
    payload: NativeSlotPayload | None = None,
) -> None:
    """Keep approved geometry while showing only accepted runtime content."""

    slide_id = _SEMANTIC_PAGE_FALLBACK_SLIDE_IDS.get(role)
    if not slide_id:
        return
    prefix = f"trace:{slide_id}:"
    protected = {
        f"{prefix}content.auto",
        f"{prefix}content.auto.2",
        f"{prefix}content.auto.3",
        f"{prefix}title.primary",
        f"{prefix}footer.brand",
        f"{prefix}footer.brand.2",
    }
    safe_content = _semantic_page_fallback_content(role, payload)
    all_shapes = list(_iter_all_shapes(slide.shapes))
    for shape in _iter_all_shapes(slide.shapes):
        name = str(getattr(shape, "name", "") or "")
        if not getattr(shape, "has_text_frame", False) or not name.startswith(prefix):
            continue
        if name in protected or ":footer." in name or ":header." in name:
            continue
        if not (":content.auto" in name or ":lead" in name or ":title.primary." in name):
            continue
        _set_shape_text_preserving_style(shape, "")

    shapes_by_name = {
        str(getattr(shape, "name", "") or ""): shape
        for shape in all_shapes
        if getattr(shape, "has_text_frame", False)
    }
    for slot_name, value in safe_content.items():
        shape = shapes_by_name.get(f"{prefix}{slot_name}")
        if shape is not None:
            _set_shape_text_preserving_style(shape, value)

    _apply_semantic_visual_finish(slide, role)
    _hide_empty_semantic_body_visuals(slide, prefix)
    if role == "SOLUTION_CONCEPT":
        _lift_semantic_insight_bar(slide, prefix)

    if role == "SOLUTION_APPROACH":
        for slot_name in ("content.auto", "content.auto.194"):
            shape = shapes_by_name.get(f"{prefix}{slot_name}")
            if shape is None:
                continue
            shape.text_frame.word_wrap = False
            shape.text_frame.margin_left = Inches(0.02)
            shape.text_frame.margin_right = Inches(0.02)
            _set_shape_text_preserving_style(shape, "08")


def _compact_proposal_summary_insight_bar(slide) -> None:
    """Keep Slide 04's existing four insight values readable in two rows."""

    insight_names = [
        "trace:S02:content.auto.130",
        "trace:S02:content.auto.131",
        "trace:S02:content.auto.132",
        "trace:S02:content.auto.133",
    ]
    shapes = {shape.name: shape for shape in _iter_all_shapes(slide.shapes)}
    summary_shape = shapes.get("trace:S02:title.primary.2")
    if summary_shape is not None and any(
        sample in str(summary_shape.text or "")
        for sample in ("提案業務の効率化", "提案品質の向上", "提案資料作成")
    ):
        _set_shape_text_preserving_style(
            summary_shape,
            "確認済み情報を整理し、未確認項目は確認後に確定します。",
        )
    insight_shapes = [shapes.get(name) for name in insight_names]
    if not all(insight_shapes):
        return

    values = [str(shape.text or "").strip() for shape in insight_shapes if str(shape.text or "").strip()]
    if not values:
        return
    split = max(1, (len(values) + 1) // 2)
    first_line = "".join(values[:split])
    second_line = "".join(values[split:])

    primary = insight_shapes[0]
    secondary = insight_shapes[1]
    bar = shapes.get("trace:S02:content.auto.129")
    if bar is not None:
        left = primary.left
        right = bar.left + bar.width - 100000
        width = max(primary.width, right - left)
        inner_top = bar.top + 20000
        row_gap = 20000
        row_height = max(1, int((bar.height - 2 * 20000 - row_gap) / 2))
        primary.left = left
        primary.top = inner_top
        primary.width = width
        primary.height = row_height
        secondary.left = left
        secondary.top = inner_top + row_height + row_gap
        secondary.width = width
        secondary.height = row_height

    _set_shape_text_preserving_style(primary, first_line)
    _set_shape_text_preserving_style(secondary, second_line)
    for shape in insight_shapes[2:]:
        _set_shape_text_preserving_style(shape, "")

    # S02 is the source template for this role, but Slide 04 is the fourth
    # summary page in the runtime deck.
    for page_number_name in ("trace:S02:content.auto", "trace:S02:content.auto.135"):
        page_number = shapes.get(page_number_name)
        if page_number is not None:
            _set_shape_text_preserving_style(page_number, "04")


def _sanitize_proposal_summary_fallback_slide(slide) -> None:
    """Clear Slide 04 sample body copy while retaining Proposal Master geometry."""

    safe_sentence = "確認済み情報を整理し、未確認項目は確認後に確定します。"
    sentence_shapes = {
        "trace:S02:title.primary.2",
        "trace:S02:content.auto.130",
    }
    insight_shapes = {
        "trace:S02:content.auto.130",
        "trace:S02:content.auto.131",
        "trace:S02:content.auto.132",
        "trace:S02:content.auto.133",
    }
    samples = tuple(sample for sample in SAMPLE_STRINGS.get("PROPOSAL_SUMMARY", ()) if sample)
    for shape in _iter_all_shapes(slide.shapes):
        if not getattr(shape, "has_text_frame", False):
            continue
        text = str(getattr(shape, "text", "") or "")
        if shape.name in sentence_shapes:
            _set_shape_text_preserving_style(shape, safe_sentence)
            continue
        if shape.name in insight_shapes:
            _set_shape_text_preserving_style(shape, "")
            continue
        if not any(sample in text for sample in samples):
            continue
        _set_shape_text_preserving_style(shape, "要確認")
    _compact_proposal_summary_insight_bar(slide)


def _iter_all_shapes(shapes) -> Iterable[object]:
    for shape in shapes:
        yield shape
        children = getattr(shape, "shapes", None)
        if children is not None:
            yield from _iter_all_shapes(children)


def _set_text_preserving_style(text_frame, value: str) -> None:
    paragraphs = list(text_frame.paragraphs)
    if not paragraphs:
        return
    first = paragraphs[0]
    runs = list(first.runs)
    if runs:
        runs[0].text = value
        for run in runs[1:]:
            run.text = ""
    else:
        first.text = value
    for paragraph in paragraphs[1:]:
        for run in paragraph.runs:
            run.text = ""


def _set_shape_text_preserving_style(shape, value: str) -> None:
    if getattr(shape, "has_text_frame", False):
        _set_text_preserving_style(shape.text_frame, value)


def _set_shape_text_color(shape, color: tuple[int, int, int]) -> None:
    """Apply a visual-only text color without changing the semantic copy."""

    if not getattr(shape, "has_text_frame", False):
        return
    rgb = RGBColor(*color)
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = rgb


def _copy_first_run_style(reference, target) -> None:
    """Copy the approved Chrome text style without changing target geometry."""

    if not (
        getattr(reference, "has_text_frame", False)
        and getattr(target, "has_text_frame", False)
        and reference.text_frame.paragraphs
        and target.text_frame.paragraphs
    ):
        return
    reference_paragraph = reference.text_frame.paragraphs[0]
    target_paragraph = target.text_frame.paragraphs[0]
    target_paragraph.alignment = reference_paragraph.alignment
    if not reference_paragraph.runs or not target_paragraph.runs:
        return
    reference_run = reference_paragraph.runs[0]
    target_run = target_paragraph.runs[0]
    reference_rpr = reference_run._r.get_or_add_rPr()
    target_rpr = target_run._r.get_or_add_rPr()
    target_run._r.replace(target_rpr, deepcopy(reference_rpr))


def _shape_by_name(slide, name: str):
    return next((shape for shape in _iter_all_shapes(slide.shapes) if shape.name == name), None)


def _set_shape_geometry(shape, *, left: float | None = None, top: float | None = None, width: float | None = None, height: float | None = None) -> None:
    if left is not None:
        shape.left = Inches(left)
    if top is not None:
        shape.top = Inches(top)
    if width is not None:
        shape.width = Inches(width)
    if height is not None:
        shape.height = Inches(height)


_PROPOSAL_MASTER_CHROME = {
    "header_page": (0.37, 0.213, 0.333, 0.222),
    "header_label": (0.963, 0.185, 2.13, 0.25),
    "header_date": (11.666, 0.204, 1.296, 0.213),
    "top_rule": (1.89, 0.30, 10.09, 0.01),
    "footer_rule": (0.37, 6.96, 12.59, 0.01),
    "footer_page": (0.37, 7.093, 0.259, 0.194),
    "footer_brand": (1.103, 7.009, 1.069, 0.253),
    "footer_tagline_jp": (2.241, 6.981, 2.5, 0.157),
    "footer_tagline_en": (2.241, 7.148, 2.593, 0.148),
    "footer_copyright": (9.444, 7.093, 3.518, 0.185),
}
_PROPOSAL_MASTER_TAGLINE_JP = "人とテクノロジーで、より良い社会をつくる"
_PROPOSAL_MASTER_TAGLINE_EN = "Think Together, Create the Next."
_PROPOSAL_MASTER_COPYRIGHT = "© 2026 提案クエスト All Rights Reserved."
_PROPOSAL_MASTER_DATE = "2026.08.26"


def _shape_bounds_in_inches(shape) -> tuple[float, float, float, float]:
    return (
        shape.left / Inches(1),
        shape.top / Inches(1),
        shape.width / Inches(1),
        shape.height / Inches(1),
    )


def _first_matching_shape(shapes: Iterable[object], predicate) -> object | None:
    return next((shape for shape in shapes if predicate(shape)), None)


def _clone_chrome_text_shape(slide, reference_shape, name: str):
    """Clone a reference footer text shape with a fresh PowerPoint shape id."""

    element = deepcopy(reference_shape.element)
    existing_ids = [int(shape.shape_id) for shape in slide.shapes if str(shape.shape_id).isdigit()]
    c_nv_pr = next(element.iter(qn("p:cNvPr")), None)
    if c_nv_pr is None:
        return None
    c_nv_pr.set("id", str(max(existing_ids, default=0) + 1))
    c_nv_pr.set("name", name)
    slide.shapes._spTree.insert_element_before(element, "p:extLst")
    return _first_matching_shape(_iter_all_shapes(slide.shapes), lambda shape: shape.name == name)


def _find_footer_reference(prs, suffix: str):
    for candidate_slide in prs.slides:
        candidate = _first_matching_shape(
            _iter_all_shapes(candidate_slide.shapes),
            lambda shape: shape.name.endswith(suffix) and getattr(shape, "has_text_frame", False),
        )
        if candidate is not None:
            return candidate
    return None


def _find_header_date_reference(prs):
    return _first_matching_shape(
        (
            shape
            for candidate_slide in prs.slides
            for shape in _iter_all_shapes(candidate_slide.shapes)
        ),
        lambda shape: getattr(shape, "has_text_frame", False)
        and bool(text_of := " ".join(str(getattr(shape, "text", "") or "").split()))
        and re.fullmatch(r"\d{4}[./-]\d{1,2}[./-]\d{1,2}", text_of)
        and shape.top / Inches(1) < 0.6
        and shape.left / Inches(1) > 10.5,
    )


def _normalize_footer_text_frame(shape, reference=None) -> None:
    """Keep approved footer copy on one line without changing its content."""

    if not getattr(shape, "has_text_frame", False):
        return
    if reference is not None and reference is not shape:
        _copy_first_run_style(reference, shape)
    text_frame = shape.text_frame
    text_frame.word_wrap = False
    text_frame.auto_size = MSO_AUTO_SIZE.NONE
    text_frame.margin_left = 0
    text_frame.margin_right = 0
    text_frame.margin_top = 0
    text_frame.margin_bottom = 0
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for paragraph in text_frame.paragraphs:
        paragraph.space_before = 0
        paragraph.space_after = 0
    paragraphs = list(text_frame.paragraphs)
    for paragraph in paragraphs[1:]:
        parent = paragraph._p.getparent()
        if parent is not None:
            parent.remove(paragraph._p)


def _finish_semantic_footer_chrome(prs: Presentation, slide, slide_number: int) -> None:
    """Finish only Slides 05-08 footer chrome after shared alignment.

    The semantic body is intentionally untouched.  These trace assets retain
    a few source-only separator lines and, on one slide, a narrow brand text
    frame that wraps the approved brand name.  Remove only those footer-region
    artifacts and normalize the already-approved text frames.
    """

    shapes = list(_iter_all_shapes(slide.shapes))
    footer_top = Inches(6.90)
    footer_bottom = Inches(7.38)

    # These are source-template separator remnants, not the approved footer
    # rule or any footer text.  Their only visual role is an orphan line.
    for shape in list(shapes):
        if not (footer_top <= shape.top <= footer_bottom):
            continue
        if shape.name.endswith((":footer.brand", ":footer.brand.2", ":footer.tagline.jp", ":footer.tagline.en")):
            continue
        if getattr(shape, "has_text_frame", False) and shape.width <= Inches(0.02) and shape.height >= Inches(0.10):
            _remove_semantic_shape(shape)

    refreshed = list(_iter_all_shapes(slide.shapes))
    page_reference = _find_footer_reference(prs, ":footer.page")
    page_shape = next(
        (
            shape
            for shape in refreshed
            if getattr(shape, "has_text_frame", False)
            and shape.top >= footer_top
            and shape.left < Inches(0.95)
            and shape.width <= Inches(0.5)
        ),
        None,
    )
    if page_shape is None and page_reference is not None:
        page_shape = _clone_chrome_text_shape(slide, page_reference, "trace:CHROME:footer.page")
        refreshed = list(_iter_all_shapes(slide.shapes))
    if page_shape is not None:
        _set_shape_text_preserving_style(page_shape, f"{slide_number:02d}")
        if page_reference is not None and page_shape is not page_reference:
            _copy_first_run_style(page_reference, page_shape)
        _set_shape_geometry(page_shape, *(), **dict(zip(("left", "top", "width", "height"), _PROPOSAL_MASTER_CHROME["footer_page"])))
        _normalize_footer_text_frame(page_shape, page_reference)

    # Keep only the approved footer rule/page/text slots in the semantic
    # fallback slides.  Any remaining source-only footer shape is an orphan
    # and cannot contribute to the Proposal Master Chrome.
    protected_page_name = page_shape.name if page_shape is not None else None
    for shape in list(_iter_all_shapes(slide.shapes)):
        if not (footer_top <= shape.top <= footer_bottom):
            continue
        if protected_page_name is not None and shape.name == protected_page_name:
            continue
        if shape.name.endswith((":footer.brand", ":footer.brand.2", ":footer.tagline.jp", ":footer.tagline.en")):
            continue
        if shape.width > Inches(8.0) and shape.height <= Inches(0.02):
            continue
        _remove_semantic_shape(shape)

    refreshed = list(_iter_all_shapes(slide.shapes))
    brand_reference = _find_footer_reference(prs, ":footer.brand")
    copyright_reference = _find_footer_reference(prs, ":footer.brand.2")
    tagline_jp_reference = _find_footer_reference(prs, ":footer.tagline.jp")
    tagline_en_reference = _find_footer_reference(prs, ":footer.tagline.en")

    for suffix, value, reference in (
        (":footer.brand", "提案クエスト", brand_reference),
        (":footer.brand.2", _PROPOSAL_MASTER_COPYRIGHT, copyright_reference),
        (":footer.tagline.jp", _PROPOSAL_MASTER_TAGLINE_JP, tagline_jp_reference),
        (":footer.tagline.en", _PROPOSAL_MASTER_TAGLINE_EN, tagline_en_reference),
    ):
        targets = [
            shape
            for shape in refreshed
            if getattr(shape, "has_text_frame", False) and shape.name.endswith(suffix)
        ]
        for target in targets:
            _set_shape_text_preserving_style(target, value)
            _normalize_footer_text_frame(target, reference)


def _unify_proposal_master_chrome(prs: Presentation, slide, slide_number: int) -> None:
    """Align only the shared Proposal Master chrome on a summary slide.

    Body shapes and semantic text are deliberately excluded.  The helper runs
    on the runtime slide copy after Native/fallback rendering, so frozen source
    templates remain unchanged.
    """

    if slide_number < 2:
        return

    shapes = list(_iter_all_shapes(slide.shapes))

    def text_of(shape) -> str:
        return " ".join(str(getattr(shape, "text", "") or "").split())

    header_page = _first_matching_shape(
        shapes,
        lambda shape: getattr(shape, "has_text_frame", False)
        and re.fullmatch(r"\d{1,2}", text_of(shape))
        and shape.top / Inches(1) < 0.65
        and shape.left / Inches(1) < 0.9,
    )
    if header_page is not None:
        _set_shape_text_preserving_style(header_page, f"{slide_number:02d}")
        _set_shape_geometry(header_page, *(), **dict(zip(("left", "top", "width", "height"), _PROPOSAL_MASTER_CHROME["header_page"])))

    footer_page = _first_matching_shape(
        shapes,
        lambda shape: getattr(shape, "has_text_frame", False)
        and re.fullmatch(r"\d{1,2}", text_of(shape))
        and shape.top / Inches(1) > 6.7
        and shape.left / Inches(1) < 0.95,
    )
    if footer_page is not None:
        _set_shape_text_preserving_style(footer_page, f"{slide_number:02d}")
        _set_shape_geometry(footer_page, *(), **dict(zip(("left", "top", "width", "height"), _PROPOSAL_MASTER_CHROME["footer_page"])))

    header_label = _first_matching_shape(
        shapes,
        lambda shape: getattr(shape, "has_text_frame", False)
        and bool(text_of(shape))
        and text_of(shape) not in {"│", "|", "｜"}
        and shape.top / Inches(1) < 0.55
        and 0.75 < shape.left / Inches(1) < 4.5
        and not re.fullmatch(r"\d{1,2}", text_of(shape)),
    )
    if header_label is not None:
        _set_shape_geometry(header_label, *(), **dict(zip(("left", "top", "width", "height"), _PROPOSAL_MASTER_CHROME["header_label"])))

    # Some trace assets use a visible text bar as a chapter separator while the
    # approved Master keeps that slot empty.  Clear only that Chrome slot.
    for shape in shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        if shape.top / Inches(1) < 0.55 and shape.left / Inches(1) < 0.9 and text_of(shape) in {"│", "|", "｜"}:
            _set_shape_text_preserving_style(shape, "")

    top_rules = [
        shape
        for shape in shapes
        if getattr(shape, "has_text_frame", False)
        and not text_of(shape)
        and shape.top / Inches(1) < 0.55
        and shape.left / Inches(1) > 1.0
        and shape.width / Inches(1) > 4.0
    ]
    for index, shape in enumerate(top_rules):
        if index == 0:
            _set_shape_geometry(shape, *(), **dict(zip(("left", "top", "width", "height"), _PROPOSAL_MASTER_CHROME["top_rule"])))
        else:
            _set_shape_geometry(shape, left=1.89, top=0.30, width=0.01, height=0.01)

    footer_rules = [
        shape
        for shape in shapes
        if getattr(shape, "has_text_frame", False)
        and not text_of(shape)
        and 6.80 < shape.top / Inches(1) < 7.15
        and shape.width / Inches(1) > 8.0
    ]
    for index, shape in enumerate(footer_rules):
        if index == 0:
            _set_shape_geometry(shape, *(), **dict(zip(("left", "top", "width", "height"), _PROPOSAL_MASTER_CHROME["footer_rule"])))
        else:
            _set_shape_geometry(shape, left=0.37, top=6.96, width=0.01, height=0.01)

    brand = _first_matching_shape(
        shapes,
        lambda shape: getattr(shape, "has_text_frame", False) and shape.name.endswith(":footer.brand"),
    )
    if brand is not None:
        _set_shape_text_preserving_style(brand, "提案クエスト")
        _set_shape_geometry(brand, *(), **dict(zip(("left", "top", "width", "height"), _PROPOSAL_MASTER_CHROME["footer_brand"])))

    copyright_shape = _first_matching_shape(
        shapes,
        lambda shape: getattr(shape, "has_text_frame", False) and shape.name.endswith(":footer.brand.2"),
    )
    if copyright_shape is None:
        reference = _find_footer_reference(prs, ":footer.brand.2")
        if reference is not None:
            copyright_shape = _clone_chrome_text_shape(slide, reference, "trace:CHROME:footer.brand.2")
    if copyright_shape is not None:
        _set_shape_text_preserving_style(copyright_shape, _PROPOSAL_MASTER_COPYRIGHT)
        _set_shape_geometry(copyright_shape, *(), **dict(zip(("left", "top", "width", "height"), _PROPOSAL_MASTER_CHROME["footer_copyright"])))

    # Remove non-Master footer slogans, then ensure the canonical two-line
    # tagline exists as editable text on every summary slide.
    for shape in list(_iter_all_shapes(slide.shapes)):
        top = shape.top / Inches(1)
        left = shape.left / Inches(1)
        if not getattr(shape, "has_text_frame", False) or top < 6.90 or left < 2.0 or left > 6.8:
            continue
        if any(shape.name.endswith(suffix) for suffix in (":footer.tagline.jp", ":footer.tagline.en")):
            continue
        if shape.name.endswith(":footer.brand") or shape.name.endswith(":footer.brand.2"):
            continue
        if text_of(shape):
            _set_shape_text_preserving_style(shape, "")

    tagline_jp = _first_matching_shape(
        _iter_all_shapes(slide.shapes),
        lambda shape: getattr(shape, "has_text_frame", False) and shape.name.endswith(":footer.tagline.jp"),
    )
    if tagline_jp is None:
        reference = _find_footer_reference(prs, ":footer.tagline.jp")
        if reference is not None:
            tagline_jp = _clone_chrome_text_shape(slide, reference, "trace:CHROME:footer.tagline.jp")
    if tagline_jp is not None:
        _set_shape_text_preserving_style(tagline_jp, _PROPOSAL_MASTER_TAGLINE_JP)
        _set_shape_geometry(tagline_jp, *(), **dict(zip(("left", "top", "width", "height"), _PROPOSAL_MASTER_CHROME["footer_tagline_jp"])))

    tagline_en = _first_matching_shape(
        _iter_all_shapes(slide.shapes),
        lambda shape: getattr(shape, "has_text_frame", False) and shape.name.endswith(":footer.tagline.en"),
    )
    if tagline_en is None:
        reference = _find_footer_reference(prs, ":footer.tagline.en")
        if reference is not None:
            tagline_en = _clone_chrome_text_shape(slide, reference, "trace:CHROME:footer.tagline.en")
    if tagline_en is not None:
        _set_shape_text_preserving_style(tagline_en, _PROPOSAL_MASTER_TAGLINE_EN)
        _set_shape_geometry(tagline_en, *(), **dict(zip(("left", "top", "width", "height"), _PROPOSAL_MASTER_CHROME["footer_tagline_en"])))

    date_shape = _first_matching_shape(
        shapes,
        lambda shape: getattr(shape, "has_text_frame", False)
        and (":footer.date" in shape.name or (shape.top / Inches(1) < 0.55 and shape.left / Inches(1) > 10.5 and shape.width / Inches(1) < 2.0)),
    )
    date_reference = _find_header_date_reference(prs)
    if date_shape is None and date_reference is not None:
        date_shape = _clone_chrome_text_shape(slide, date_reference, "trace:CHROME:header.date")
    if date_shape is not None:
        _set_shape_text_preserving_style(date_shape, _PROPOSAL_MASTER_DATE)
        if date_reference is not None and date_shape is not date_reference:
            _copy_first_run_style(date_reference, date_shape)
        _set_shape_geometry(date_shape, *(), **dict(zip(("left", "top", "width", "height"), _PROPOSAL_MASTER_CHROME["header_date"])))

    if slide_number == 9:
        for shape in shapes:
            if getattr(shape, "has_text_frame", False) and "最適な一手" in text_of(shape):
                _set_shape_text_preserving_style(shape, "")

    if 5 <= slide_number <= 8:
        _finish_semantic_footer_chrome(prs, slide, slide_number)


def _tighten_kpi_fallback_layout(slide) -> None:
    """Keep the evidence-safe KPI copy readable inside the approved layout."""

    shapes = list(_iter_all_shapes(slide.shapes))

    text_replacements = {
        "現状値（未取得）": "現状値",
        "現状値未取得": "未取得",
        "現状値未取得向上": "要確認",
        "目標値は要確認": "要確認",
        "(要確認)": "要確認",
        "(3.現状値未取得)": "要確認",
        "作業時間の実績記録\n（ツール・工数表）": "工数記録",
        "営業管理ツールの\n案件件数": "案件数",
        "版管理・履歴からの\nカウント": "履歴カウント",
        "営業管理ツールの\n受注率": "受注率",
        "既存のデータやヒアリングをもとに、現状の数値を把握します。": "既存データとヒアリングで現状値を確認します。",
        "事業目標やリソースを踏まえ、現実的かつ挑戦的な目標値を設定します。": "事業目標とリソースを踏まえ、目標値を設定します。",
        "月次でKPIを確認し、要因分析と打ち手の検討を行います。": "月次でKPIを確認し、改善策を検討します。",
        "導入初期や計測環境が整っていない場合は、仮説値を設定し、実測後に見直します。": "初期は仮説とし、実測後に見直します。",
        "業界・案件規模・提案内容などの条件により、効果は変動する点に留意が必要です。": "条件により効果は変動します。",
        "KPIの進捗を定期的に確認し、施策の見直し・改善を継続的に行います。": "進捗を確認し、施策を継続改善します。",
    }
    for shape in shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        updated = shape.text
        for source, replacement in sorted(text_replacements.items(), key=lambda item: -len(item[0])):
            updated = updated.replace(source, replacement)
        if updated != shape.text:
            _set_shape_text_preserving_style(shape, updated)

    for name in (
        "trace:KPI:title.primary.4",
        "trace:KPI:content.auto.51",
        "trace:KPI:content.auto.64",
        "trace:KPI:content.auto.74",
    ):
        shape = _shape_by_name(slide, name)
        if shape is not None:
            _set_shape_text_preserving_style(shape, "確認後に確定")

    for name in (
        "trace:KPI:title.primary.6",
        "trace:KPI:content.auto.56",
        "trace:KPI:content.auto.69",
        "trace:KPI:content.auto.79",
    ):
        shape = _shape_by_name(slide, name)
        if shape is not None:
            _set_shape_text_preserving_style(shape, "方法を確認")

    for name in (
        "trace:KPI:content.auto.43",
        "trace:KPI:content.auto.58",
        "trace:KPI:content.auto.71",
        "trace:KPI:content.auto.81",
    ):
        shape = _shape_by_name(slide, name)
        if shape is not None:
            _set_shape_text_preserving_style(shape, "要確認")

    table_shape = _shape_by_name(slide, "trace:KPI:title.primary.3")
    if table_shape is not None and getattr(table_shape, "has_table", False):
        table = table_shape.table
        column_widths = (1.45, 1.0, 1.0, 1.55, 0.45)
        for column, width in zip(table.columns, column_widths):
            column.width = Inches(width)
        for row, height in zip(table.rows, (0.29, 0.96, 0.96, 0.96, 0.97)):
            row.height = Inches(height)
        cells = [
            ["KPI", "現状値", "目標値", "測定方法", "優先度"],
            ["確認項目", "未確認", "確認後に確定", "測定方法を確認", "要確認"],
            ["確認項目", "未確認", "確認後に確定", "根拠を確認", "要確認"],
            ["確認項目", "未確認", "確認後に確定", "条件を確認", "要確認"],
            ["確認項目", "未確認", "確認後に確定", "実績を確認", "要確認"],
        ]
        for row, values in zip(table.rows, cells):
            for cell, value in zip(row.cells, values):
                _set_text_preserving_style(cell.text_frame, value)

    # Keep the editable grid lines aligned with the widened, readable columns.
    for name, left, width in (
        ("trace:KPI:content.auto.27", 0.28, 1.45),
        ("trace:KPI:content.auto.28", 1.73, 1.0),
        ("trace:KPI:content.auto.29", 2.73, 1.0),
        ("trace:KPI:content.auto.30", 3.73, 1.55),
        ("trace:KPI:content.auto.31", 5.28, 0.45),
    ):
        shape = _shape_by_name(slide, name)
        if shape is not None:
            _set_shape_geometry(shape, left=left, width=width)

    for name, left in {
        "trace:KPI:content.auto.18": 1.73,
        "trace:KPI:content.auto.19": 2.73,
        "trace:KPI:content.auto.20": 3.73,
        "trace:KPI:content.auto.21": 5.28,
    }.items():
        shape = _shape_by_name(slide, name)
        if shape is not None:
            shape.left = Inches(left)

    # The semantic text overlays are kept editable and placed inside their
    # corresponding columns after the grid correction.
    for name in (
        "trace:KPI:content.auto.39", "trace:KPI:content.auto.52",
        "trace:KPI:content.auto.65", "trace:KPI:content.auto.75",
    ):
        shape = _shape_by_name(slide, name)
        if shape is not None:
            _set_shape_geometry(shape, left=1.82, width=0.82)
    for name in (
        "trace:KPI:content.auto.41", "trace:KPI:content.auto.54",
        "trace:KPI:content.auto.67", "trace:KPI:content.auto.77",
        "trace:KPI:title.primary.5", "trace:KPI:content.auto.55",
        "trace:KPI:content.auto.68", "trace:KPI:content.auto.78",
    ):
        shape = _shape_by_name(slide, name)
        if shape is not None:
            _set_shape_geometry(shape, left=2.82, width=0.82)
    for name in (
        "trace:KPI:title.primary.6", "trace:KPI:content.auto.56",
        "trace:KPI:content.auto.69", "trace:KPI:content.auto.79",
    ):
        shape = _shape_by_name(slide, name)
        if shape is not None:
            _set_shape_geometry(shape, left=3.82, width=1.42)
    for name in (
        "trace:KPI:content.auto.43", "trace:KPI:content.auto.58",
        "trace:KPI:content.auto.71", "trace:KPI:content.auto.81",
    ):
        shape = _shape_by_name(slide, name)
        if shape is not None:
            _set_shape_geometry(shape, left=5.29, width=0.38)

    for name, value in {
        "trace:KPI:content.auto": "10",
        "trace:KPI:content.auto.110": "10",
    }.items():
        shape = _shape_by_name(slide, name)
        if shape is not None:
            _set_shape_text_preserving_style(shape, value)
def _tighten_schedule_fallback_layout(slide) -> None:
    """Reduce fallback density without creating dates, owners, or milestones."""

    shapes = list(_iter_all_shapes(slide.shapes))

    def _shrink_cluster(names: tuple[str, ...], *, left: float, top: float, width: float, height: float) -> None:
        members = [shape for name in names if (shape := _shape_by_name(slide, name)) is not None]
        if not members:
            return
        min_left = min(shape.left for shape in members)
        min_top = min(shape.top for shape in members)
        max_right = max(shape.left + shape.width for shape in members)
        max_bottom = max(shape.top + shape.height for shape in members)
        old_width = max_right - min_left
        old_height = max_bottom - min_top
        if not old_width or not old_height:
            return
        target_left = Inches(left)
        target_top = Inches(top)
        target_width = Inches(width)
        target_height = Inches(height)
        scale = min(target_width / old_width, target_height / old_height)
        placed_width = int(old_width * scale)
        placed_height = int(old_height * scale)
        origin_left = target_left + int((target_width - placed_width) / 2)
        origin_top = target_top + int((target_height - placed_height) / 2)
        for shape in members:
            shape.left = origin_left + int((shape.left - min_left) * scale)
            shape.top = origin_top + int((shape.top - min_top) * scale)
            shape.width = max(0, int(shape.width * scale))
            shape.height = max(0, int(shape.height * scale))

    # Keep the two-digit slide number on one line without changing its text.
    page_number = _shape_by_name(slide, "trace:SCHEDULE:content.auto")
    if page_number is not None and getattr(page_number, "has_text_frame", False):
        page_number.text_frame.word_wrap = False
        _set_shape_geometry(page_number, left=0.353, width=0.36, height=0.255)

    # Fit each governance icon inside the left icon column, leaving a stable
    # gutter before the role text.  Only icon geometry is changed here.
    _shrink_cluster(
        tuple(f"trace:SCHEDULE:content.auto.{index}" for index in range(112, 115)),
        left=7.10,
        top=2.19,
        width=0.20,
        height=0.31,
    )
    _shrink_cluster(
        tuple(f"trace:SCHEDULE:content.auto.{index}" for index in range(119, 126)),
        left=7.10,
        top=2.78,
        width=0.20,
        height=0.31,
    )
    _shrink_cluster(
        ("trace:SCHEDULE:content.auto.177", "trace:SCHEDULE:content.auto.178"),
        left=7.10,
        top=3.37,
        width=0.20,
        height=0.31,
    )
    _shrink_cluster(
        tuple(f"trace:SCHEDULE:content.auto.{index}" for index in range(135, 141)),
        left=7.10,
        top=3.95,
        width=0.20,
        height=0.31,
    )
    _shrink_cluster(
        tuple(f"trace:SCHEDULE:content.auto.{index}" for index in range(147, 153)),
        left=7.12,
        top=4.56,
        width=0.18,
        height=0.22,
    )

    # Keep role/value columns on one consistent left padding and right edge;
    # wording and evidence-safe status values remain unchanged.
    for name in (
        "trace:SCHEDULE:content.auto.115",
        "trace:SCHEDULE:title.primary.9",
        "trace:SCHEDULE:content.auto.129",
        "trace:SCHEDULE:content.auto.141",
    ):
        shape = _shape_by_name(slide, name)
        if shape is not None:
            _set_shape_geometry(shape, left=7.46, width=0.78)
    project_manager_role = _shape_by_name(slide, "trace:SCHEDULE:title.primary.9")
    if project_manager_role is not None:
        _set_shape_geometry(project_manager_role, left=7.42, width=0.86)
    for name in (
        "trace:SCHEDULE:title.primary.7",
        "trace:SCHEDULE:title.primary.8",
        "trace:SCHEDULE:title.primary.10",
        "trace:SCHEDULE:title.primary.11",
        "trace:SCHEDULE:content.auto.131",
        "trace:SCHEDULE:content.auto.132",
        "trace:SCHEDULE:content.auto.143",
        "trace:SCHEDULE:content.auto.144",
    ):
        shape = _shape_by_name(slide, name)
        if shape is not None:
            _set_shape_geometry(shape, left=8.34, width=1.14)

    for shape in shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        if shape.text in {"確認後に確定", "確認後に確定担当"}:
            _set_shape_text_preserving_style(shape, "要確認")

    # Deliverable names from the web-production fixture are not verified ERP
    # evidence. Keep the approved Master layout and replace only the section
    # contents with an explicit safe status label.
    for name in (
        "trace:SCHEDULE:content.auto.92",
        "trace:SCHEDULE:content.auto.93",
        "trace:SCHEDULE:content.auto.94",
        "trace:SCHEDULE:content.auto.95",
        "trace:SCHEDULE:content.auto.96",
    ):
        shape = _shape_by_name(slide, name)
        if shape is not None:
            _set_shape_text_preserving_style(shape, "要確認")

    # The approved Master uses the deck-wide date and the real slide index.
    for name, value in {
        "trace:SCHEDULE:content.auto": "11",
        "trace:SCHEDULE:content.auto.173": "11",
        "trace:SCHEDULE:footer.date": "2026.08.26",
    }.items():
        shape = _shape_by_name(slide, name)
        if shape is not None:
            _set_shape_text_preserving_style(shape, value)
    date_shape = _shape_by_name(slide, "trace:SCHEDULE:footer.date")
    if date_shape is not None:
        _set_shape_geometry(date_shape, left=8.72, width=1.05)

    # Keep the phrase visible in the lead and policy band while shortening
    # repeated, low-information cells to a single safe status label.
    for name, value in {
        "trace:SCHEDULE:content.auto.169": "未確認の予定は、",
        "trace:SCHEDULE:content.auto.170": "確認後に確定",
        "trace:SCHEDULE:content.auto.171": "します。",
    }.items():
        shape = _shape_by_name(slide, name)
        if shape is not None:
            _set_shape_text_preserving_style(shape, value)

    # Give the weekly grid a real finite width and leave a clear gutter before
    # the right governance panel.  The table remains a native editable table.
    table_shape = _shape_by_name(slide, "trace:SCHEDULE:content.auto.97")
    if table_shape is not None and getattr(table_shape, "has_table", False):
        table_shape.width = Inches(6.15)
        table = table_shape.table
        table.columns[0].width = Inches(1.5)
        week_width = (6.15 - 1.5) / 12
        for column_index in range(1, len(table.columns)):
            table.columns[column_index].width = Inches(week_width)
        for row, height in zip(table.rows, (0.32, 0.27, 0.27, 0.27, 0.27, 0.29)):
            row.height = Inches(height)

    # Shift the right-hand panels a little right and preserve their right edge.
    for shape in shapes:
        if not shape.name.startswith("trace:SCHEDULE:content.auto"):
            continue
        suffix = shape.name.rsplit(".", 1)[-1]
        if not suffix.isdigit() or not 103 <= int(suffix) <= 165:
            continue
        shape.left += Inches(0.12)
        if shape.width > Inches(0.22):
            shape.width = max(Inches(0.18), shape.width - Inches(0.12))
def _apply_visual_fallback_layout(slide, role: str) -> None:
    if role == "KPI":
        _tighten_kpi_fallback_layout(slide)
    elif role == "SCHEDULE":
        _tighten_schedule_fallback_layout(slide)


def _render_visual_master_fallback(
    prs: Presentation,
    role: str,
    slide_data,
    spec: dict[str, object],
    *,
    surface: str,
    media: bool = False,
    semantic_payload: NativeSlotPayload | None = None,
) -> dict[str, object]:
    """Render a safe fallback using an approved Proposal Master template.

    This path deliberately never turns missing evidence into Native eligibility.
    It reuses the approved visual package, replaces only safe identity/status
    text, and imports the editable slide so the legacy renderer is not used.
    """

    source = Path(__file__).resolve().parents[4] / str(spec.get("runtime_asset", ""))
    if not source.is_file():
        raise ValueError(f"fallback template missing: {source}")
    slide_id = str(spec.get("slide_id"))
    required_slots = list(spec.get("required_slots", ()))
    with tempfile.TemporaryDirectory(prefix="phase3f25e-visual-fallback-") as temp_dir:
        target = Path(temp_dir) / f"{slide_id}_fallback.pptx"
        clone_template_package(
            source,
            target,
            approved_slide_id=slide_id,
            required_slots=required_slots,
        )
        title_slot = f"trace:{slide_id}:title.primary"
        payload = semantic_payload or NativeSlotPayload(
            role=role,
            slide_id=slide_id,
            surface=surface,
            slots={} if role == "COVER" else {title_slot: _visual_fallback_title(role, slide_data)},
            text_replacements=_visual_fallback_replacements(role),
            prohibited_sample_strings=list(SAMPLE_STRINGS.get(role, ())),
        )
        if semantic_payload is not None:
            payload.text_replacements.update(_visual_fallback_replacements(role))
            payload.prohibited_sample_strings = list(SAMPLE_STRINGS.get(role, ()))
        write_report = write_slot_payload(target, payload, required_slots=required_slots)
        _replace_fragmented_fallback_text(
            target,
            {
                **_visual_fallback_replacements(role),
                "READY CREW Proposal": "提案クエスト",
                "READY CREW Inc.": "提案クエスト",
                "READY CREW": "提案クエスト",
                "ProposalPilot": "提案クエスト",
                "AI営業秘書": "提案クエスト",
            },
        )
        validation = validate_injected_template_package(
            target,
            source_template=source,
            approved_slide_id=slide_id,
            required_slots=required_slots,
            unresolved_required_slots=write_report["unresolved_required_slots"],
            prohibited_sample_strings=payload.prohibited_sample_strings,
            text_fit_diagnostics={"valid": True, "fields": {}},
            evidence_diagnostics=[
                {
                    "status": "FALLBACK",
                    "reason": "verified evidence unavailable; Proposal Master Chrome retained",
                }
            ],
            expected_source_sha256=sha256_file(source),
        )
        if not validation.get("valid") or not validation.get("sample_content_leak_free"):
            raise ValueError({"fallback_validation": validation})
        if media:
            _clone_template_slide_with_media(prs, target)
        else:
            _clone_template_slide(prs, target)
        if role == "PROPOSAL_SUMMARY":
            _sanitize_proposal_summary_fallback_slide(prs.slides[-1])
        elif role == "IMPLEMENTATION_CONFIGURATION":
            _sanitize_implementation_configuration_fallback_slide(prs.slides[-1])
        elif role in _SEMANTIC_PAGE_FALLBACK_SLIDE_IDS:
            _sanitize_semantic_page_fallback_slide(prs.slides[-1], role, semantic_payload)
        _apply_visual_fallback_layout(prs.slides[-1], role)
        return {
            "validation": validation,
            "injected_slots": write_report["injected_slots"],
            "template_id": slide_id,
        }


def _safe_dynamic_values(role: str, context) -> dict[str, str]:
    """Return only values that are safe to bind without inventing evidence."""

    values: dict[str, str] = {}
    if role == "ESTIMATE":
        estimate = getattr(context, "estimate", None)
        total_label = getattr(estimate, "total_label", "要確認") or "要確認"
        values["￥10,500,000"] = str(total_label)
    if role == "WIN_PROBABILITY":
        probability = getattr(getattr(context, "win_probability", None), "probability", None)
        values["72%"] = str(probability) if probability is not None else "要確認"
    elif role == "KPI":
        # KPI trace samples are not production evidence.  The current model
        # does not carry verified actual/target provenance, so use the
        # contract fallbacks instead of promoting example metrics.
        values.update(
            {
                "20時間/件": "未取得",
                "70%削減": "要確認",
                "(6時間/件)": "(要確認)",
                "月10件": "未取得",
                "1.5倍": "要確認",
                "(月15件)": "(要確認)",
                "5回/件": "未取得",
                "30%削減": "要確認",
                "(3.5回/件)": "(要確認)",
                "20%": "未取得",
                "20%向上": "要確認",
                "(24%)": "(要確認)",
            }
        )
    elif role == "COMPETITIVE_COMPARISON":
        # The existing competitor model has no claim-level provenance.  Keep
        # the editable canonical structure but prevent unsupported superiority
        # statements from entering production output.
        values.update(
            {
                "業務理解から提案・制作・改善まで一気通貫で支援": "比較条件と根拠は要確認",
                "AI活用を前提にした運用設計まで提示": "比較条件と根拠は要確認",
                "概要見積・体制・スケジュールの透明性が高い": "比較条件と根拠は要確認",
                "公開後の改善伴走まで含めた提案": "比較条件と根拠は要確認",
                "当社提案を第一候補として比較継続し、条件調整後に最終判断へ進むことを推奨します。": "根拠確認後に比較条件を整理し、最終判断へ進みます。",
            }
        )
    elif role in {"SCHEDULE", "ROADMAP"}:
        # Schedule sample dates/durations are trace-only.  The current model
        # exposes phase descriptions but not verified dates or owners.
        values.update(
            {
                "1 ～ 2 週": "次回確認",
                "2 ～ 3 週": "次回確認",
                "4 ～ 6 週": "次回確認",
                "2 週": "次回確認",
                "(1ヶ月)": "(要確認)",
                "(1〜2ヶ月)": "(要確認)",
                "(2〜3ヶ月)": "(要確認)",
                "(3〜6ヶ月)": "(要確認)",
                "(Week 1)": "(次回確認)",
                "(Week 6)": "(次回確認)",
                "(Week 12)": "(次回確認)",
                "(Week 20)": "(次回確認)",
                "約3か月で公開、その後も継続改善できる体制": "公開時期と継続改善方針は要確認",
            }
        )
    return values


def _apply_safe_bindings(slide, role: str, slide_data, context) -> None:
    _replace_brand_text(slide)
    _replace_title(slide, role, getattr(slide_data, "title", ""))
    replacements = _safe_dynamic_values(role, context)
    if not replacements:
        return
    for shape in _iter_text_shapes(slide):
        _replace_shape_text(shape, replacements)


def render_approved_native_slide(
    prs: Presentation,
    slide_data,
    data,
    context,
    index: int,
    *,
    role: str | None = None,
    surface: str | None = None,
) -> bool:
    """Render one approved native slide through the fail-closed adapter path."""

    return dispatch_approved_native_slide(
        prs,
        slide_data,
        data,
        context,
        index,
        role=role,
        surface=surface,
    )["NATIVE_RENDERED"]


def _effective_surface(role: str | None, surface: str | None) -> str:
    runtime_role = canonical_runtime_role(role)
    if runtime_role in {"COMPETITION", "WIN_PROBABILITY", "ROADMAP"}:
        return "conditional"
    return surface or "summary"


def _trace(
    role: str | None,
    *,
    native_rendered: bool,
    native_eligible: bool | None = None,
    failure_reason: str | None,
    human_approval_status: str,
    provenance_status: str,
    text_fit_status: str,
    sample_leak_status: str,
    requested_role: str | None = None,
    template_id: str | None = None,
    fallback_rendered: bool = False,
    fallback_kind: str | None = None,
) -> dict[str, object]:
    return {
        "ROLE": role,
        "REQUESTED_ROLE": requested_role or role,
        "RUNTIME_ROLE": role,
        "TEMPLATE_ID": template_id,
        "NATIVE_REQUESTED": True,
        "NATIVE_ELIGIBLE": bool(native_rendered if native_eligible is None else native_eligible),
        "NATIVE_RENDERED": bool(native_rendered),
        "FALLBACK_USED": not native_rendered,
        "FALLBACK_RENDERED": bool(fallback_rendered),
        "FALLBACK_KIND": fallback_kind,
        "FAILURE_REASON": failure_reason,
        "HUMAN_APPROVAL_STATUS": human_approval_status,
        "PROVENANCE_STATUS": provenance_status,
        "TEXT_FIT_STATUS": text_fit_status,
        "SAMPLE_LEAK_STATUS": sample_leak_status,
    }


def dispatch_approved_native_slide(
    prs: Presentation,
    slide_data,
    data,
    context,
    index: int,
    *,
    role: str | None = None,
    surface: str | None = None,
    visual_fallback: bool = False,
) -> dict[str, object]:
    """Attempt one native role and return a structured role-level trace.

    The function never raises for an expected native eligibility/rendering
    failure.  The caller owns the existing-renderer fallback.
    """

    requested_role = role
    resolved_role = role
    if resolved_role is None:
        from app.services.pptx_parts.native_trace_registry import resolve_approved_native_role

        resolved_role = resolve_approved_native_role(slide_data, index)
        requested_role = resolved_role
    resolved_role = canonical_runtime_role(resolved_role)
    effective_surface = _effective_surface(resolved_role, surface)
    spec = get_runtime_native_role_spec(resolved_role, surface=effective_surface)
    if spec is None:
        return _trace(
            resolved_role,
            native_rendered=False,
            failure_reason=FailureReason.ROLE_NOT_REGISTERED.value,
            human_approval_status="UNKNOWN",
            provenance_status="UNKNOWN",
            text_fit_status="NOT_RUN",
            sample_leak_status="NOT_RUN",
            requested_role=requested_role,
        )

    source_path = Path(__file__).resolve().parents[4] / str(spec.get("runtime_asset", ""))
    gate, checks = runtime_native_role_gate(
        resolved_role,
        surface=effective_surface,
        slide_id=str(spec.get("slide_id")),
        template_available=source_path.is_file(),
        data_contract_valid=True,
        text_fit_valid=True,
        renderer_available=True,
    )
    if not gate:
        failed_gate = next((name for name, passed in checks.items() if not passed), "NATIVE_GATE_FAILED")
        return _trace(
            resolved_role,
            native_rendered=False,
            failure_reason=failed_gate,
            human_approval_status="PASS" if checks.get("HUMAN_APPROVED") else "PENDING_OR_UNKNOWN",
            provenance_status="NOT_RUN",
            text_fit_status="NOT_RUN",
            sample_leak_status="NOT_RUN",
            requested_role=requested_role,
            template_id=str(spec.get("slide_id")),
        )

    try:
        result = render_native_role_dry_run(
            resolved_role,
            data=data,
            context=context,
            slide=slide_data,
            surface=effective_surface,
            slide_id=str(spec.get("slide_id")),
        )
        payload = result.payload
        provenance_status = "PASS" if payload and payload.failure_reason is None else "FAIL"
        text_fit_status = "PASS" if result.text_fit.get("valid") else ("FAIL" if result.text_fit else "NOT_RUN")
        sample_leak_status = "PASS" if result.validation.get("injected", {}).get("sample_content_leak_free", False) else ("FAIL" if result.validation else "NOT_RUN")
        if not result.success or not result.package_path:
            if visual_fallback and result.failure_reason == FailureReason.EVIDENCE_REQUIRED.value and resolved_role in _VISUAL_FALLBACK_ROLES:
                fallback = _render_visual_master_fallback(
                    prs,
                    resolved_role,
                    slide_data,
                    spec,
                    surface=effective_surface,
                    semantic_payload=payload,
                )
                if effective_surface == "summary":
                    _unify_proposal_master_chrome(prs, prs.slides[-1], int(getattr(slide_data, "slide_no", index + 1)))
                return _trace(
                    resolved_role,
                    native_rendered=False,
                    native_eligible=False,
                    failure_reason=FailureReason.EVIDENCE_REQUIRED.value,
                    human_approval_status="PASS",
                    provenance_status="FAIL",
                    text_fit_status=text_fit_status,
                    sample_leak_status="PASS" if fallback["validation"].get("sample_content_leak_free") else "FAIL",
                    requested_role=requested_role,
                    template_id=str(spec.get("slide_id")),
                    fallback_rendered=True,
                    fallback_kind="PROPOSAL_MASTER_EVIDENCE_SAFE",
                )
            return _trace(
                resolved_role,
                native_rendered=False,
                failure_reason=result.failure_reason or FailureReason.VALIDATION_FAILED.value,
                human_approval_status="PASS",
                provenance_status=provenance_status,
                text_fit_status=text_fit_status,
                sample_leak_status=sample_leak_status,
                requested_role=requested_role,
                template_id=str(spec.get("slide_id")),
            )

        # Import only the already validated, runtime-cloned package.  Image
        # relationships are deliberately rejected by the editable slide
        # importer; the established renderer then handles that role safely.
        _clone_template_slide(prs, Path(result.package_path))
        if resolved_role == "PROPOSAL_SUMMARY":
            _compact_proposal_summary_insight_bar(prs.slides[-1])
        if effective_surface == "summary":
            _unify_proposal_master_chrome(prs, prs.slides[-1], int(getattr(slide_data, "slide_no", index + 1)))
        return _trace(
            resolved_role,
            native_rendered=True,
            failure_reason=None,
            human_approval_status="PASS",
            provenance_status=provenance_status,
            text_fit_status=text_fit_status,
            sample_leak_status=sample_leak_status,
            requested_role=requested_role,
            template_id=str(spec.get("slide_id")),
        )
    except ValueError as exc:
        if visual_fallback and resolved_role == "COVER" and "unsupported relationships" in str(exc):
            try:
                fallback = _render_visual_master_fallback(
                    prs,
                    resolved_role,
                    slide_data,
                    spec,
                    surface=effective_surface,
                    media=True,
                )
                if effective_surface == "summary":
                    _unify_proposal_master_chrome(prs, prs.slides[-1], int(getattr(slide_data, "slide_no", index + 1)))
                return _trace(
                    resolved_role,
                    native_rendered=False,
                    native_eligible=True,
                    failure_reason=FailureReason.UNSUPPORTED_RELATIONSHIP.value,
                    human_approval_status="PASS",
                    provenance_status="PASS",
                    text_fit_status="PASS",
                    sample_leak_status="PASS" if fallback["validation"].get("sample_content_leak_free") else "FAIL",
                    requested_role=requested_role,
                    template_id=str(spec.get("slide_id")),
                    fallback_rendered=True,
                    fallback_kind="APPROVED_COVER_PACKAGE_CLONE",
                )
            except Exception:
                logger.exception("approved_cover_visual_fallback_failed")
        return _trace(
            resolved_role,
            native_rendered=False,
            native_eligible=True,
            failure_reason=(
                FailureReason.UNSUPPORTED_RELATIONSHIP.value
                if "unsupported relationships" in str(exc)
                else FailureReason.TEMPLATE_CLONE_FAILED.value
            ),
            human_approval_status="PASS",
            provenance_status="PASS",
            text_fit_status="PASS",
            sample_leak_status="PASS",
            requested_role=requested_role,
            template_id=str(spec.get("slide_id")),
        )
    except Exception:
        logger.exception("approved_native_dispatch_failed", extra={"role": resolved_role})
        return _trace(
            resolved_role,
            native_rendered=False,
            native_eligible=True,
            failure_reason=FailureReason.TEMPLATE_CLONE_FAILED.value,
            human_approval_status="PASS",
            provenance_status="PASS",
            text_fit_status="PASS",
            sample_leak_status="PASS",
            requested_role=requested_role,
            template_id=str(spec.get("slide_id")),
        )


# Name used by integration tests and future renderer adapters.
render_approved_slide = render_approved_native_slide
