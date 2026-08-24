from __future__ import annotations

from datetime import date
from typing import Callable

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

from app.models import PowerPointData, PowerPointSlide, WinProbability
from app.services.pptx_parts.models import PptxContext
from app.services.pptx_parts.content import (
    _contains_any,
    _trim,
    build_case_triplets_from_items,
    build_solution_rows,
    concept_statement,
    derive_kpi_rows,
    display_case_title,
    ensure_items,
    has_competitor_context,
    journey_action,
    kpi_metric_for,
    merge_understanding_items,
    projected_probability_for,
    rank_color_for,
    rank_light_color_for,
    rank_probability_for,
    risk_label_for,
    risk_score_for_probability,
    sitemap_note,
    unique_items,
)
from app.services.pptx_parts.drawing import (
    add_bullet_list,
    add_card,
    add_case_line,
    add_factor_column,
    add_footer,
    add_header,
    add_icon_badge,
    add_insight_band,
    add_section_label,
    add_shape,
    add_side_panel,
    add_step_flow,
    add_table,
    add_text,
    add_title,
    add_visual_frame,
    blank_slide,
    chapter_icon,
    chapter_message,
    content_title,
    set_background,
)
from app.services.pptx_design.components import (
    add_architecture_diagram,
    add_estimate_overview,
    add_metric_card,
    add_next_action_cards,
    add_timeline,
)
from app.services.pptx_design.icons import icon_labels_for_category
from app.services.pptx_design.diagrams import architecture_nodes_for_category
from app.services.pptx_design_system.typography import normalize_customer_facing_text, split_label_body
from app.services.pptx_layout_integration import layout_id_from_layout_key
from app.services.pptx_quality import extract_numbers
from app.services.pptx_theme import COLORS, MARGIN_X, SECTION_COLORS, SLIDE_HEIGHT, SLIDE_WIDTH, resolve_template_colors


def _split_metric_text(value: str) -> tuple[str, str]:
    if ":" in value:
        label, body = value.split(":", 1)
        return label.strip(), body.strip()
    if "：" in value:
        label, body = value.split("：", 1)
        return label.strip(), body.strip()
    return value, "評価基準を合意"


def _layout_items(slide_data: PowerPointSlide, count: int, fallback: list[str] | None = None) -> list[str]:
    fallback_items = fallback or [f"{slide_data.title or 'Review point'} {index + 1}" for index in range(count)]
    return ensure_items(unique_items(slide_data.bullets, count), fallback_items, count)


def _layout_numbers(slide_data: PowerPointSlide, count: int = 4) -> list[str]:
    body = "\n".join([slide_data.title, *slide_data.bullets])
    return unique_items(extract_numbers(body), count)


def _layout_accent(index: int = 0) -> str:
    return SECTION_COLORS[index % len(SECTION_COLORS)]


def _add_layout_header(slide, slide_data: PowerPointSlide, section: str, accent: str | None = None) -> None:
    add_header(slide, slide_data.title or section.title(), section, accent=accent or _layout_accent())


def render_title_only(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_section_label(slide, "要点", 0.9, 0.62, fill=COLORS["navy"], color=COLORS["white"])
    add_title(slide, slide_data.title, 1.05, 2.35, 11.15, 0.9, size=34, color=COLORS["navy"])
    if slide_data.bullets:
        add_insight_band(slide, "主旨", slide_data.bullets[0], 1.05, 4.05, 11.05, 0.78)
    add_footer(slide, slide_data.slide_no)


def render_title_body(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _add_layout_header(slide, slide_data, "MESSAGE", COLORS["blue"])
    add_bullet_list(slide, slide_data.bullets, 0.95, 1.8, 7.8, 3.85, max_items=5, size=15)
    add_visual_frame(slide, slide_data.visual_suggestion or "diagram placeholder", 9.05, 1.82, 3.2, 3.82)
    add_footer(slide, slide_data.slide_no)


def render_two_column(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _add_layout_header(slide, slide_data, "TWO COLUMN", COLORS["teal"])
    items = _layout_items(slide_data, 6)
    midpoint = max(1, (len(items) + 1) // 2)
    add_card(slide, "主要論点", "", 0.95, 1.72, 5.45, 3.85, COLORS["blue"], COLORS["white"])
    add_card(slide, "補足論点", "", 6.78, 1.72, 5.45, 3.85, COLORS["teal"], COLORS["white"])
    add_bullet_list(slide, items[:midpoint], 1.24, 2.35, 4.88, 2.65, max_items=4, size=13)
    add_bullet_list(slide, items[midpoint:], 7.08, 2.35, 4.88, 2.65, max_items=4, size=13)
    add_footer(slide, slide_data.slide_no)


def render_three_column(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _add_layout_header(slide, slide_data, "THREE COLUMN", COLORS["purple"])
    items = _layout_items(slide_data, 3)
    for idx, item in enumerate(items[:3]):
        add_card(slide, f"重点テーマ {idx + 1}", item, 0.95 + idx * 3.88, 1.86, 3.42, 3.35, _layout_accent(idx), COLORS["white"], str(idx + 1))
    add_footer(slide, slide_data.slide_no)


def render_hero(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    add_cover_slide(prs, slide_data, data, context)


def render_kpi_cards(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _add_layout_header(slide, slide_data, "KPI", COLORS["green"])
    numbers = _layout_numbers(slide_data, 4)
    items = _layout_items(slide_data, 4)
    if numbers:
        for idx, value in enumerate(numbers[:4]):
            add_metric_card(slide, f"評価指標 {idx + 1}", value, 0.95 + idx * 2.88, 1.72, 2.5, 1.45, _layout_accent(idx))
        add_bullet_list(slide, items, 1.0, 3.72, 11.2, 1.65, max_items=5, size=13)
    else:
        for idx, item in enumerate(items[:4]):
            add_card(slide, f"KPI設計 {idx + 1}", item, 0.95 + idx * 2.88, 1.72, 2.5, 2.35, _layout_accent(idx), COLORS["white"])
    add_footer(slide, slide_data.slide_no)


def render_comparison_table(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _add_layout_header(slide, slide_data, "COMPARISON", COLORS["blue"])
    items = _layout_items(slide_data, 6)
    labels = ["現状", "課題", "目指す姿", "本提案の強み"]
    for idx, label in enumerate(labels):
        x = 0.92 + idx * 3.0
        accent = _layout_accent(idx)
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.82, 2.62, 3.08, fill=COLORS["white"], line=COLORS["line"])
        add_shape(slide, MSO_SHAPE.RECTANGLE, x, 1.82, 2.62, 0.09, fill=accent, line=accent)
        add_icon_badge(slide, str(idx + 1), x + 0.88, 2.18, accent, size=0.66)
        add_text(slide, label, x + 0.24, 3.0, 2.14, 0.28, size=15, color=accent, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, _trim(items[idx], 42), x + 0.28, 3.52, 2.06, 0.68, size=13, color=COLORS["text"], bold=True, align=PP_ALIGN.CENTER)
        if idx < 3:
            add_shape(slide, MSO_SHAPE.CHEVRON, x + 2.52, 3.02, 0.32, 0.46, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    add_table(
        slide,
        ["判断根拠", "次回確認事項"],
        [[_trim(items[4] if len(items) > 4 else slide_data.speaker_notes, 58), _trim(items[5] if len(items) > 5 else "お客様と確認", 48)]],
        0.95,
        5.35,
        11.38,
        0.72,
        [6.3, 5.08],
    )
    add_footer(slide, slide_data.slide_no)


def render_timeline(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _add_layout_header(slide, slide_data, "工程", COLORS["teal"])
    items = _layout_items(slide_data, 5)
    add_timeline(slide, items, 0.95, 2.02, 11.55)
    for idx, item in enumerate(items[:5]):
        add_card(slide, f"工程 {idx + 1}", item, 0.95 + idx * 2.3, 4.28, 2.05, 1.05, _layout_accent(idx), COLORS["white"])
    add_footer(slide, slide_data.slide_no)


def render_roadmap(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _add_layout_header(slide, slide_data, "ロードマップ", COLORS["orange"])
    items = _layout_items(slide_data, 5)
    add_timeline(slide, items, 0.92, 1.95, 11.6)
    add_insight_band(slide, "意思決定ポイント", items[-1], 0.95, 5.42, 11.45, 0.72)
    add_footer(slide, slide_data.slide_no)


def render_flow(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _add_layout_header(slide, slide_data, "FLOW", COLORS["purple"])
    items = _layout_items(slide_data, 5)
    add_step_flow(slide, items, 0.95, 2.05, 11.35, 1.65)
    add_side_panel(slide, "確認観点", items[:4], 8.85, 4.18, 3.38, 1.68, COLORS["teal"])
    add_footer(slide, slide_data.slide_no)


def render_matrix(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _add_layout_header(slide, slide_data, "MATRIX", COLORS["orange"])
    items = _layout_items(slide_data, 4)
    positions = [(0.95, 1.78), (6.72, 1.78), (0.95, 4.12), (6.72, 4.12)]
    labels = ["優先度が高い", "効果が大きい", "確認が必要", "優先度を調整"]
    for idx, (x, y) in enumerate(positions):
        add_card(slide, labels[idx], items[idx], x, y, 5.32, 1.42, _layout_accent(idx), COLORS["white"])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 6.48, 1.58, 0.03, 4.22, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0.9, 3.75, 11.52, 0.03, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    add_footer(slide, slide_data.slide_no)


def render_quote(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _add_layout_header(slide, slide_data, "QUOTE", COLORS["teal"])
    quote = slide_data.bullets[0] if slide_data.bullets else slide_data.title
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 1.25, 2.05, 10.78, 2.36, fill=COLORS["canvas"], line=COLORS["line"])
    add_text(slide, _trim(quote, 92), 1.78, 2.68, 9.72, 0.68, size=24, color=COLORS["navy"], bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, slide_data.slide_no)


def render_image_left(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _add_layout_header(slide, slide_data, "IMAGE LEFT", COLORS["blue"])
    add_visual_frame(slide, slide_data.visual_suggestion or "image placeholder", 0.95, 1.78, 4.55, 3.9)
    add_bullet_list(slide, slide_data.bullets, 6.05, 1.9, 5.8, 3.55, max_items=5, size=14)
    add_footer(slide, slide_data.slide_no)


def render_image_right(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _add_layout_header(slide, slide_data, "IMAGE RIGHT", COLORS["blue"])
    add_bullet_list(slide, slide_data.bullets, 0.95, 1.9, 5.8, 3.55, max_items=5, size=14)
    add_visual_frame(slide, slide_data.visual_suggestion or "image placeholder", 7.18, 1.78, 4.55, 3.9)
    add_footer(slide, slide_data.slide_no)


def render_large_number(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _add_layout_header(slide, slide_data, "NUMBER", COLORS["green"])
    numbers = _layout_numbers(slide_data, 1)
    items = _layout_items(slide_data, 4)
    if numbers:
        add_text(slide, numbers[0], 0.95, 2.02, 4.2, 0.9, size=46, color=COLORS["green"], bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, _trim(items[0], 52), 0.95, 3.05, 4.2, 0.48, size=18, color=COLORS["navy"], bold=True, align=PP_ALIGN.CENTER)
        add_bullet_list(slide, items[1:], 5.85, 2.0, 5.75, 2.35, max_items=4, size=14)
    else:
        add_bullet_list(slide, items, 1.0, 2.1, 11.1, 2.6, max_items=4, size=18)
    add_footer(slide, slide_data.slide_no)


def render_checklist(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _add_layout_header(slide, slide_data, "CHECKLIST", COLORS["teal"])
    items = _layout_items(slide_data, 6)
    for idx, item in enumerate(items[:6]):
        x = 0.95 + (idx % 2) * 5.78
        y = 1.72 + (idx // 2) * 1.32
        add_shape(slide, MSO_SHAPE.OVAL, x, y + 0.12, 0.34, 0.34, fill=_layout_accent(idx), line=_layout_accent(idx))
        add_text(slide, str(idx + 1), x + 0.01, y + 0.2, 0.32, 0.1, size=8, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, _trim(item, 54), x + 0.52, y + 0.1, 4.9, 0.38, size=15, color=COLORS["text"], bold=True)
    add_footer(slide, slide_data.slide_no)


def render_closing(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_section_label(slide, "次の行動", 0.9, 0.72, fill=COLORS["navy"], color=COLORS["white"])
    add_title(slide, slide_data.title or "次のアクション", 1.02, 1.34, 11.15, 0.64, size=32, color=COLORS["navy"])
    add_next_action_cards(slide, _layout_items(slide_data, 4), 0.92, 2.42, 11.55)
    add_insight_band(slide, "意思決定", "次のアクションと担当を合意し、実行開始までの迷いをなくします。", 0.95, 5.48, 11.42, 0.72)
    add_footer(slide, slide_data.slide_no)





def render_executive_message(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _add_layout_header(slide, slide_data, "経営判断", COLORS["blue"])
    items = _layout_items(slide_data, 4, [context.concept, "現状課題を整理", "期待効果を確認", "次の意思決定を合意"])
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.95, 1.62, 11.35, 1.46, fill=COLORS["navy"], line=COLORS["navy"])
    add_text(slide, "この提案の結論", 1.28, 1.88, 2.1, 0.24, size=13, color=COLORS["teal_light"], bold=True)
    add_text(slide, _trim(items[0], 62), 3.35, 1.78, 8.45, 0.5, size=25, color=COLORS["white"], bold=True)
    labels = ["背景", "判断材料", "期待効果"]
    for idx, item in enumerate(items[1:4]):
        add_card(slide, labels[idx], item, 0.95 + idx * 4.02, 3.58, 3.36, 1.46, SECTION_COLORS[idx], COLORS["white"], number=str(idx + 1))
    add_insight_band(slide, "意思決定ポイント", "範囲、KPI、予算、開始時期を合意できれば、次工程へ進められます。", 0.95, 5.76, 11.4, 0.58)
    add_footer(slide, slide_data.slide_no)


def render_current_state_map(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _add_layout_header(slide, slide_data, "現状整理", COLORS["teal"])
    items = _layout_items(slide_data, 5, ["現在の業務・顧客接点", "発生している課題", "事業への影響", "改善機会", "確認事項"])
    labels = ["現状", "課題", "影響", "改善機会", "確認事項"]
    for idx, item in enumerate(items[:5]):
        x = 0.9 + idx * 2.32
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.82, 2.02, 3.15, fill=COLORS["white"], line=COLORS["line"])
        add_shape(slide, MSO_SHAPE.RECTANGLE, x, 1.82, 2.02, 0.09, fill=SECTION_COLORS[idx % len(SECTION_COLORS)], line=SECTION_COLORS[idx % len(SECTION_COLORS)])
        add_text(slide, labels[idx], x + 0.18, 2.12, 1.66, 0.24, size=15, color=SECTION_COLORS[idx % len(SECTION_COLORS)], bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, _trim(item, 42), x + 0.2, 2.82, 1.62, 0.78, size=13, color=COLORS["text"], bold=True, align=PP_ALIGN.CENTER)
        if idx < 4:
            add_shape(slide, MSO_SHAPE.CHEVRON, x + 1.92, 3.08, 0.28, 0.38, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    add_insight_band(slide, "提案へのつながり", "現状の事実から課題、影響、改善機会へ順に整理し、提案の必要性を明確にします。", 0.95, 5.72, 11.4, 0.62)
    add_footer(slide, slide_data.slide_no)


def render_problem_structure(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _add_layout_header(slide, slide_data, "課題構造", COLORS["red"])
    items = _layout_items(slide_data, 5, ["表面化している課題", "業務上の制約", "根本原因", "事業影響", "解決の方向性"])
    add_card(slide, "表面課題", items[0], 0.95, 1.78, 3.0, 1.2, COLORS["orange"], COLORS["white"])
    add_card(slide, "業務要因", items[1], 0.95, 3.42, 3.0, 1.2, COLORS["purple"], COLORS["white"])
    add_shape(slide, MSO_SHAPE.CHEVRON, 4.12, 2.75, 0.46, 0.52, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    add_card(slide, "根本原因", items[2], 4.82, 2.28, 3.22, 1.56, COLORS["red"], COLORS["white"], number="1")
    add_shape(slide, MSO_SHAPE.CHEVRON, 8.24, 2.75, 0.46, 0.52, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    add_card(slide, "事業影響", items[3], 8.95, 1.78, 3.0, 1.2, COLORS["blue"], COLORS["white"])
    add_card(slide, "解決の方向性", items[4], 8.95, 3.42, 3.0, 1.2, COLORS["green"], COLORS["white"])
    add_insight_band(slide, "課題認識", "症状ではなく原因に手を打つことで、導入後の効果を測定しやすくします。", 0.95, 5.72, 11.4, 0.58)
    add_footer(slide, slide_data.slide_no)


def render_before_after_transformation(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _add_layout_header(slide, slide_data, "Before / After", COLORS["blue"])
    items = _layout_items(slide_data, 5, ["人手・個別対応が中心", "判断基準が分散", "AIと人の確認で支援", "既存システムへ連携", "履歴を改善へ活用"])
    add_card(slide, "現状", items[0], 0.95, 1.82, 3.2, 2.2, COLORS["muted"], COLORS["white"])
    add_card(slide, "変化の要点", items[1], 5.02, 1.82, 3.2, 2.2, COLORS["blue"], COLORS["white"])
    add_card(slide, "目指す姿", items[2], 9.08, 1.82, 3.2, 2.2, COLORS["green"], COLORS["white"])
    add_shape(slide, MSO_SHAPE.CHEVRON, 4.34, 2.66, 0.42, 0.56, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    add_shape(slide, MSO_SHAPE.CHEVRON, 8.4, 2.66, 0.42, 0.56, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    add_insight_band(slide, "導入後の運用", f"{_trim(items[3], 46)} / {_trim(items[4], 46)}", 0.95, 5.38, 11.4, 0.82)
    add_footer(slide, slide_data.slide_no)


def render_strategic_options(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _add_layout_header(slide, slide_data, "推奨方針", COLORS["blue"])
    items = _layout_items(slide_data, 4, ["現状維持", "部分改善", "段階導入", "推奨案"])
    labels = ["選択肢A", "選択肢B", "推奨案", "判断理由"]
    for idx, item in enumerate(items[:4]):
        fill = COLORS["blue_light"] if idx == 2 else COLORS["white"]
        add_card(slide, labels[idx], item, 0.95 + (idx % 2) * 5.78, 1.72 + (idx // 2) * 1.82, 5.22, 1.28, SECTION_COLORS[idx], fill, number=str(idx + 1))
    add_insight_band(slide, "本提案の立ち位置", context.winning_strategy, 0.95, 5.72, 11.4, 0.58)
    add_footer(slide, slide_data.slide_no)


def render_value_proposition(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _add_layout_header(slide, slide_data, "提供価値", COLORS["green"])
    items = _layout_items(slide_data, 4, ["業務負荷を下げる", "品質を安定させる", "判断を早める", "継続改善につなげる"])
    widths = [8.8, 6.6, 4.4]
    labels = ["経営価値", "業務価値", "実行価値"]
    for idx, width in enumerate(widths):
        x = 6.65 - width / 2
        y = 1.82 + idx * 1.05
        add_shape(slide, MSO_SHAPE.TRAPEZOID, x, y, width, 0.76, fill=[COLORS["green"], COLORS["blue"], COLORS["teal"]][idx], line=COLORS["white"])
        add_text(slide, labels[idx], x + 0.2, y + 0.24, width - 0.4, 0.16, size=15, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
    for idx, item in enumerate(items[:4]):
        add_card(slide, f"価値 {idx + 1}", item, 0.95 + idx * 2.88, 4.78, 2.5, 1.0, SECTION_COLORS[idx], COLORS["white"])
    add_footer(slide, slide_data.slide_no)


def render_kpi_design_dashboard(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _add_layout_header(slide, slide_data, "KPI設計", COLORS["green"])
    items = _layout_items(slide_data, 4, ["候補正答率", "人手修正率", "確認時間", "処理品質"])
    headers = ["指標", "現状", "目標", "測定方法"]
    rows = []
    for idx, item in enumerate(items[:4]):
        label, body = split_label_body(item, f"評価指標 {idx + 1}")
        rows.append([label, "現状確認", "PoCで確定", body or "ログ・レビューで測定"])
    add_table(slide, headers, rows, 0.85, 1.64, 11.75, 3.42, [2.2, 2.1, 2.1, 5.35])
    add_insight_band(slide, "SMART設計", "目標値を断定せず、現状値・目標値・測定方法・判定タイミングをPoCで合意します。", 0.95, 5.58, 11.4, 0.72)
    add_footer(slide, slide_data.slide_no)


def render_roi_logic(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _add_layout_header(slide, slide_data, "ROIの考え方", COLORS["orange"])
    items = _layout_items(slide_data, 4, ["投資範囲を明確化", "削減・改善効果を測定", "運用負荷を確認", "本番判断へ接続"])
    labels = ["投資", "業務効果", "品質効果", "判断"]
    for idx, item in enumerate(items[:4]):
        x = 0.95 + idx * 2.88
        add_metric_card(slide, labels[idx], _trim(item, 18), x, 1.76, 2.5, 1.32, SECTION_COLORS[idx])
        add_shape(slide, MSO_SHAPE.RECTANGLE, x + 0.15, 3.48, 2.2, 0.12 + idx * 0.12, fill=SECTION_COLORS[idx], line=SECTION_COLORS[idx])
    add_insight_band(slide, "ROI説明方針", "根拠のない金額効果は置かず、PoCで測定できる業務時間・品質・運用負荷から投資判断します。", 0.95, 5.58, 11.4, 0.72)
    add_footer(slide, slide_data.slide_no)


def render_competitive_positioning(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _add_layout_header(slide, slide_data, "競合に対する位置取り", COLORS["purple"])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 6.52, 1.72, 0.03, 3.9, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 1.05, 3.68, 10.95, 0.03, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    add_text(slide, "提案適合度 高", 5.35, 1.42, 2.4, 0.22, size=12, color=COLORS["muted"], bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "導入負荷 低", 0.95, 3.42, 1.7, 0.22, size=12, color=COLORS["muted"], bold=True)
    add_card(slide, "推奨ポジション", context.winning_strategy, 7.0, 2.02, 3.8, 1.22, COLORS["blue"], COLORS["blue_light"])
    items = _layout_items(slide_data, 3, ["競合仮説", "差別化", "確認事項"])
    for idx, item in enumerate(items[:3]):
        add_card(slide, ["想定競合", "差別化", "確認事項"][idx], item, 1.05 + idx * 3.68, 5.0, 3.25, 0.92, SECTION_COLORS[idx], COLORS["white"])
    add_footer(slide, slide_data.slide_no)


def render_layered_architecture(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _add_layout_header(slide, slide_data, "構成案", COLORS["blue"])
    layers = ["入力", "AI判定", "人の確認", "連携", "改善運用"]
    labels = ["入力", "AI処理", "人の確認", "連携", "改善運用"]
    items = _layout_items(slide_data, 5, ["商品画像・データ", "候補判定", "修正・承認", "API/CSV連携", "ログ・再学習"])
    for idx, item in enumerate(items[:5]):
        x = 0.86 + idx * 2.45
        accent = SECTION_COLORS[idx % len(SECTION_COLORS)]
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.78, 2.12, 3.05, fill=COLORS["white"], line=COLORS["line"])
        add_text(slide, labels[idx], x + 0.18, 2.08, 1.76, 0.24, size=14, color=accent, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, _trim(item, 36), x + 0.18, 2.84, 1.76, 0.66, size=13, color=COLORS["text"], bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, layers[idx], x + 0.18, 4.16, 1.76, 0.18, size=9, color=COLORS["muted"], align=PP_ALIGN.CENTER)
        if idx < 4:
            add_shape(slide, MSO_SHAPE.CHEVRON, x + 2.02, 3.0, 0.28, 0.42, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    add_insight_band(slide, "運用改善ループ", "確認・修正履歴を次回改善に活用し、完全自動化ではなく人の判断を支援します。", 0.95, 5.62, 11.4, 0.62)
    add_footer(slide, slide_data.slide_no)


def render_workstream_roadmap(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _add_layout_header(slide, slide_data, "実行ロードマップ", COLORS["teal"])
    phases = _layout_items(slide_data, 5, ["要件・評価基準", "データ確認", "検証", "現場確認", "本番判断"])
    lanes = ["作業", "成果物", "判断"]
    for lane_idx, lane in enumerate(lanes):
        y = 1.82 + lane_idx * 1.18
        add_text(slide, lane, 0.92, y + 0.12, 1.1, 0.22, size=12, color=COLORS["muted"], bold=True)
        add_shape(slide, MSO_SHAPE.RECTANGLE, 2.05, y + 0.24, 9.98, 0.05, fill=COLORS["line"], line=COLORS["line"])
        for idx, phase in enumerate(phases[:5]):
            x = 2.1 + idx * 1.95
            add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 1.56, 0.48, fill=SECTION_COLORS[idx], line=SECTION_COLORS[idx])
            label = _trim(phase, 14) if lane_idx == 0 else ("確認" if lane_idx == 2 else "成果物")
            add_text(slide, label, x + 0.08, y + 0.15, 1.4, 0.12, size=9, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
    add_insight_band(slide, "意思決定地点", "各フェーズで成果物と判断ポイントを確認し、PoCから本番判断へ進めます。", 0.95, 5.7, 11.4, 0.58)
    add_footer(slide, slide_data.slide_no)


def render_risk_heatmap(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _add_layout_header(slide, slide_data, "リスクと対策", COLORS["red"])
    items = _layout_items(slide_data, 4, ["精度不足", "データ不足", "運用負荷", "連携条件"])
    colors = [COLORS["green_light"], COLORS["orange_light"], COLORS["red_light"], COLORS["blue_light"]]
    labels = ["低", "中", "高", "要確認"]
    for row in range(2):
        for col in range(2):
            idx = row * 2 + col
            x = 1.08 + col * 3.52
            y = 1.86 + row * 1.58
            add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 3.1, 1.16, fill=colors[idx], line=COLORS["white"])
            add_text(slide, labels[idx], x + 0.18, y + 0.18, 0.82, 0.2, size=13, color=SECTION_COLORS[idx], bold=True)
            add_text(slide, _trim(items[idx], 36), x + 1.05, y + 0.18, 1.75, 0.42, size=13, color=COLORS["text"], bold=True)
    add_side_panel(slide, "対策方針", ["事前評価", "PoC確認", "運用設計", "段階導入"], 8.45, 1.86, 3.5, 2.74, COLORS["teal"])
    add_insight_band(slide, "提出前の見方", "リスクを隠さず、確認方法と対策を同時に提示することで意思決定しやすくします。", 0.95, 5.72, 11.4, 0.58)
    add_footer(slide, slide_data.slide_no)


def render_governance_map(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _add_layout_header(slide, slide_data, "推進体制", COLORS["purple"])
    items = _layout_items(slide_data, 5, ["意思決定", "業務責任", "IT連携", "現場確認", "弊社支援"])
    add_shape(slide, MSO_SHAPE.OVAL, 5.58, 2.38, 2.0, 1.4, fill=COLORS["navy"], line=COLORS["navy"])
    add_text(slide, "共同推進", 5.9, 2.9, 1.35, 0.18, size=15, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
    positions = [(1.0, 1.62), (9.15, 1.62), (1.0, 4.45), (9.15, 4.45), (5.05, 4.9)]
    for idx, (x, y) in enumerate(positions):
        add_card(slide, ["決裁", "業務", "IT", "現場", "支援"][idx], items[idx], x, y, 3.0, 1.0, SECTION_COLORS[idx], COLORS["white"])
    add_footer(slide, slide_data.slide_no)


def render_cost_breakdown(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _add_layout_header(slide, slide_data, "費用構成", COLORS["orange"])
    columns = [
        ("必須", context.estimate.required or ["要件整理", "検証"], COLORS["blue"], COLORS["blue_light"]),
        ("推奨", context.estimate.recommended or ["効果測定", "運用支援"], COLORS["green"], COLORS["green_light"]),
        ("任意", context.estimate.optional or ["追加連携", "拡張支援"], COLORS["orange"], COLORS["orange_light"]),
    ]
    add_metric_card(slide, "概算費用", context.estimate.total_label, 0.95, 1.48, 3.2, 1.15, COLORS["orange"])
    add_metric_card(slide, "予算適合", context.estimate.budget_fit, 4.58, 1.48, 3.2, 1.15, COLORS["green"])
    add_card(slide, "説明方針", "必須範囲を先に確保し、推奨・任意を段階的に判断します。", 8.2, 1.48, 3.9, 1.15, COLORS["blue"], COLORS["white"])
    for idx, (title, items, accent, fill) in enumerate(columns):
        x = 0.92 + idx * 4.08
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 3.18, 3.58, 2.3, fill=fill, line=COLORS["white"])
        add_text(slide, title, x + 0.22, 3.48, 3.12, 0.28, size=17, color=accent, bold=True, align=PP_ALIGN.CENTER)
        add_bullet_list(slide, items, x + 0.32, 3.95, 2.92, 1.15, max_items=3, size=12)
    add_footer(slide, slide_data.slide_no)


def render_scope_definition(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _add_layout_header(slide, slide_data, "提案範囲", COLORS["blue"])
    items = _layout_items(slide_data, 6, ["今回対象", "今回対象外", "前提条件", "確認事項", "変更条件", "次回合意"])
    labels = ["対象", "対象外", "前提", "確認", "変更条件", "次回合意"]
    for idx, item in enumerate(items[:6]):
        x = 0.95 + (idx % 3) * 3.88
        y = 1.72 + (idx // 3) * 1.72
        add_card(slide, labels[idx], item, x, y, 3.35, 1.18, SECTION_COLORS[idx % len(SECTION_COLORS)], COLORS["white"])
    add_insight_band(slide, "範囲管理", "提案範囲と対象外を明確にし、見積・スケジュールの前提を揃えます。", 0.95, 5.72, 11.4, 0.58)
    add_footer(slide, slide_data.slide_no)


def render_section_divider(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT, fill=COLORS["navy"], line=COLORS["navy"])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 0.18, SLIDE_HEIGHT, fill=COLORS["teal"], line=COLORS["teal"])
    add_text(slide, "Section", 0.95, 1.55, 2.0, 0.24, size=13, color=COLORS["teal_light"], bold=True)
    add_text(slide, slide_data.title or "提案の章", 0.95, 2.18, 8.8, 0.82, size=38, color=COLORS["white"], bold=True)
    add_text(slide, _trim(chapter_message(slide_data.title, context), 86), 0.98, 3.42, 8.2, 0.48, size=17, color=COLORS["teal_light"], bold=True)
    add_footer(slide, slide_data.slide_no)


def _v31_text(value: str, limit: int = 38) -> str:
    return normalize_customer_facing_text(value, limit=limit)


def _v31_items(slide_data: PowerPointSlide, count: int, fallback: list[str] | None = None, *, limit: int = 38) -> list[str]:
    return [_v31_text(item, limit) for item in _layout_items(slide_data, count, fallback)]


def _v31_add_header(slide, slide_data: PowerPointSlide, section: str, accent: str) -> None:
    add_header(slide, _v31_text(slide_data.title or section, 38), section, accent=accent)


def _v31_add_value_card(slide, title: str, body: str, x: float, y: float, w: float, h: float, accent: str, *, fill: str | None = None, number: str | None = None) -> None:
    body_limit = max(10, min(32, int(w * 7)))
    title_limit = max(8, min(18, int(w * 5)))
    clean_title = _v31_text(title, title_limit)
    clean_body = _v31_text(body, body_limit)
    if h < 1.08:
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill or COLORS["white"], line=COLORS["line"])
        add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, 0.06, h, fill=accent, line=accent)
        add_text(slide, clean_title, x + 0.18, y + 0.16, w - 0.36, 0.16, size=10, color=accent, bold=True)
        if clean_body:
            add_text(slide, clean_body, x + 0.18, y + 0.44, w - 0.36, max(0.18, h - 0.52), size=10, color=COLORS["text"], bold=True)
        return
    add_card(slide, clean_title, clean_body, x, y, w, h, accent, fill or COLORS["white"], number=number)


def _v31_add_micro_icon(slide, label: str, x: float, y: float, accent: str) -> None:
    add_shape(slide, MSO_SHAPE.OVAL, x, y, 0.5, 0.5, fill=accent, line=accent)
    add_text(slide, _trim(label, 3), x + 0.05, y + 0.19, 0.4, 0.1, size=9, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)


def _v31_add_arrow(slide, x: float, y: float, accent: str = COLORS["line_dark"]) -> None:
    add_shape(slide, MSO_SHAPE.CHEVRON, x, y, 0.36, 0.48, fill=accent, line=accent)


def render_v31_title_only(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT, fill=COLORS["navy"], line=COLORS["navy"])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 0.16, SLIDE_HEIGHT, fill=COLORS["teal"], line=COLORS["teal"])
    add_section_label(slide, "要点", 0.95, 0.76, fill=COLORS["teal"], color=COLORS["white"])
    add_text(slide, _v31_text(slide_data.title, 34), 0.95, 2.18, 9.6, 0.78, size=38, color=COLORS["white"], bold=True)
    if slide_data.bullets:
        add_text(slide, _v31_text(slide_data.bullets[0], 52), 1.0, 3.3, 8.7, 0.42, size=18, color=COLORS["teal_light"], bold=True)
    add_footer(slide, slide_data.slide_no)


def render_v31_title_body(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _v31_add_header(slide, slide_data, "要点", COLORS["blue"])
    items = _v31_items(slide_data, 4, ["背景", "課題", "提案", "次の判断"], limit=36)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.95, 1.78, 7.25, 2.0, fill=COLORS["navy"], line=COLORS["navy"])
    add_text(slide, "このページの結論", 1.26, 2.06, 2.15, 0.24, size=13, color=COLORS["teal_light"], bold=True)
    add_text(slide, items[0], 1.26, 2.48, 6.32, 0.44, size=24, color=COLORS["white"], bold=True)
    for idx, item in enumerate(items[1:4]):
        _v31_add_value_card(slide, ["背景", "提案", "次の判断"][idx], item, 1.0 + idx * 3.78, 4.32, 3.2, 1.12, SECTION_COLORS[idx], number=str(idx + 1))
    add_visual_frame(slide, "図解で補足", 8.72, 1.78, 3.35, 2.0)
    add_footer(slide, slide_data.slide_no)


def render_v31_two_column(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _v31_add_header(slide, slide_data, "比較", COLORS["teal"])
    items = _v31_items(slide_data, 4, ["現状", "課題", "改善後", "判断材料"], limit=36)
    _v31_add_value_card(slide, "現状の見方", items[0], 0.95, 1.82, 5.32, 2.72, COLORS["blue"], number="1")
    _v31_add_value_card(slide, "目指す姿", items[2], 7.02, 1.82, 5.32, 2.72, COLORS["green"], fill=COLORS["green_light"], number="2")
    _v31_add_arrow(slide, 6.48, 2.9, COLORS["line_dark"])
    add_insight_band(slide, "判断材料", f"{items[1]} / {items[3]}", 0.95, 5.45, 11.4, 0.64)
    add_footer(slide, slide_data.slide_no)


def render_v31_three_column(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _v31_add_header(slide, slide_data, "重点", COLORS["purple"])
    items = _v31_items(slide_data, 3, ["課題", "打ち手", "効果"], limit=38)
    for idx, item in enumerate(items[:3]):
        x = 1.05 + idx * 3.85
        _v31_add_micro_icon(slide, ["課", "策", "効"][idx], x + 1.34, 1.9, SECTION_COLORS[idx])
        _v31_add_value_card(slide, ["課題", "提案", "効果"][idx], item, x, 2.62, 3.12, 2.15, SECTION_COLORS[idx], fill=COLORS["white"], number=str(idx + 1))
    add_footer(slide, slide_data.slide_no)


def render_v31_kpi_cards(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _v31_add_header(slide, slide_data, "KPI", COLORS["green"])
    items = _v31_items(slide_data, 4, ["現状値", "目標値", "測定方法", "判定基準"], limit=12)
    for idx, item in enumerate(items[:4]):
        x = 0.95 + idx * 2.88
        add_metric_card(slide, ["現状", "目標", "測定", "判定"][idx], _v31_text(item, 12), x, 1.72, 2.5, 1.28, SECTION_COLORS[idx])
        add_shape(slide, MSO_SHAPE.ARC, x + 0.58, 3.38, 1.36, 0.64, fill=COLORS["white"], line=SECTION_COLORS[idx])
        add_shape(slide, MSO_SHAPE.RECTANGLE, x + 0.42, 4.2, 1.68, 0.12, fill=COLORS["line"], line=COLORS["line"])
        add_shape(slide, MSO_SHAPE.RECTANGLE, x + 0.42, 4.2, 0.58 + idx * 0.28, 0.12, fill=SECTION_COLORS[idx], line=SECTION_COLORS[idx])
    add_insight_band(slide, "SMART設計", "現状・目標・測定方法・判定基準をPoCで合意します。", 0.95, 5.55, 11.4, 0.66)
    add_footer(slide, slide_data.slide_no)


def render_v31_comparison_cards(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _v31_add_header(slide, slide_data, "比較", COLORS["blue"])
    items = _v31_items(slide_data, 5, ["現状", "課題", "改善後", "判断軸", "確認事項"], limit=36)
    labels = ["現状", "課題", "改善後"]
    fills = [COLORS["canvas"], COLORS["orange_light"], COLORS["blue_light"]]
    for idx, item in enumerate(items[:3]):
        x = 0.95 + idx * 3.88
        _v31_add_value_card(slide, labels[idx], item, x, 1.86, 3.28, 2.18, SECTION_COLORS[idx], fill=fills[idx], number=str(idx + 1))
        if idx < 2:
            _v31_add_arrow(slide, x + 3.36, 2.62)
    add_insight_band(slide, "比較の見方", f"{items[3]} / {items[4]}", 0.95, 5.38, 11.4, 0.72)
    add_footer(slide, slide_data.slide_no)


def render_v31_timeline(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _v31_add_header(slide, slide_data, "進行計画", COLORS["teal"])
    items = _v31_items(slide_data, 5, ["要件整理", "設計", "検証", "導入", "改善"], limit=22)
    add_timeline(slide, items, 0.98, 2.1, 11.3)
    for idx, item in enumerate(items[:5]):
        _v31_add_value_card(slide, f"Phase {idx + 1}", item, 1.0 + idx * 2.25, 4.15, 1.9, 0.82, SECTION_COLORS[idx])
    add_footer(slide, slide_data.slide_no)


def render_v31_roadmap(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _v31_add_header(slide, slide_data, "ロードマップ", COLORS["orange"])
    items = _v31_items(slide_data, 5, ["準備", "検証", "改善", "導入", "定着"], limit=24)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.95, 2.0, 11.35, 0.34, fill=COLORS["blue_light"], line=COLORS["blue_light"])
    for idx, item in enumerate(items[:5]):
        x = 1.08 + idx * 2.22
        add_shape(slide, MSO_SHAPE.OVAL, x, 1.84, 0.64, 0.64, fill=SECTION_COLORS[idx], line=SECTION_COLORS[idx])
        add_text(slide, str(idx + 1), x, 2.04, 0.64, 0.12, size=10, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
        _v31_add_value_card(slide, f"Step {idx + 1}", item, x - 0.55, 3.08, 1.72, 1.0, SECTION_COLORS[idx])
    add_insight_band(slide, "意思決定", "PoCの結果を見て、本番化範囲と開始時期を合意します。", 0.95, 5.62, 11.4, 0.62)
    add_footer(slide, slide_data.slide_no)


def render_v31_flow(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _v31_add_header(slide, slide_data, "業務フロー", COLORS["purple"])
    items = _v31_items(slide_data, 5, ["入力", "AI判定", "確認", "連携", "改善"], limit=22)
    add_step_flow(slide, items[:4], 0.95, 2.05, 8.5, 1.58)
    add_side_panel(slide, "確認ポイント", items[1:5], 9.72, 1.88, 2.42, 2.7, COLORS["teal"])
    add_insight_band(slide, "運用設計", "AIは候補を出し、人が最終確認する流れを前提にします。", 0.95, 5.55, 11.4, 0.66)
    add_footer(slide, slide_data.slide_no)


def render_v31_matrix(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _v31_add_header(slide, slide_data, "優先順位", COLORS["orange"])
    items = _v31_items(slide_data, 4, ["すぐ着手", "効果大", "確認必要", "後続対応"], limit=32)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 6.48, 1.68, 0.03, 4.0, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 1.0, 3.65, 10.9, 0.03, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    positions = [(1.18, 1.92), (6.92, 1.92), (1.18, 4.05), (6.92, 4.05)]
    labels = ["高優先", "高効果", "要確認", "後続"]
    for idx, (x, y) in enumerate(positions):
        _v31_add_value_card(slide, labels[idx], items[idx], x, y, 4.78, 1.18, SECTION_COLORS[idx], fill=COLORS["white"])
    add_footer(slide, slide_data.slide_no)


def render_v31_quote(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _v31_add_header(slide, slide_data, "結論", COLORS["teal"])
    quote = _v31_text(slide_data.bullets[0] if slide_data.bullets else slide_data.title, 58)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 1.28, 2.05, 10.72, 2.3, fill=COLORS["canvas"], line=COLORS["line"])
    add_text(slide, quote, 1.78, 2.7, 9.72, 0.62, size=26, color=COLORS["navy"], bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, slide_data.slide_no)


def render_v31_image_left(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _v31_add_header(slide, slide_data, "図解", COLORS["blue"])
    add_visual_frame(slide, slide_data.visual_suggestion or "差し替え可能な図解", 0.95, 1.82, 4.72, 3.62)
    items = _v31_items(slide_data, 3, ["要点", "根拠", "次の判断"], limit=38)
    for idx, item in enumerate(items[:3]):
        _v31_add_value_card(slide, ["要点", "根拠", "判断"][idx], item, 6.2, 1.82 + idx * 1.24, 5.65, 0.92, SECTION_COLORS[idx])
    add_footer(slide, slide_data.slide_no)


def render_v31_image_right(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _v31_add_header(slide, slide_data, "図解", COLORS["blue"])
    items = _v31_items(slide_data, 3, ["要点", "根拠", "次の判断"], limit=38)
    for idx, item in enumerate(items[:3]):
        _v31_add_value_card(slide, ["要点", "根拠", "判断"][idx], item, 0.95, 1.82 + idx * 1.24, 5.65, 0.92, SECTION_COLORS[idx])
    add_visual_frame(slide, slide_data.visual_suggestion or "差し替え可能な図解", 7.1, 1.82, 4.72, 3.62)
    add_footer(slide, slide_data.slide_no)


def render_v31_large_number(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _v31_add_header(slide, slide_data, "主要指標", COLORS["green"])
    numbers = _layout_numbers(slide_data, 1)
    items = _v31_items(slide_data, 4, ["見るべき指標", "改善余地", "測定方法", "次の判断"], limit=34)
    value = numbers[0] if numbers else "PoCで確定"
    add_text(slide, value, 0.95, 2.02, 4.2, 0.9, size=46, color=COLORS["green"], bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, items[0], 0.95, 3.04, 4.2, 0.46, size=18, color=COLORS["navy"], bold=True, align=PP_ALIGN.CENTER)
    for idx, item in enumerate(items[1:4]):
        _v31_add_value_card(slide, ["改善余地", "測定", "判断"][idx], item, 5.72, 1.86 + idx * 1.18, 5.72, 0.88, SECTION_COLORS[idx])
    add_footer(slide, slide_data.slide_no)


def render_v31_checklist(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _v31_add_header(slide, slide_data, "確認事項", COLORS["teal"])
    items = _v31_items(slide_data, 5, ["目的", "範囲", "体制", "日程", "判断"], limit=34)
    for idx, item in enumerate(items[:5]):
        x = 1.0 + (idx % 3) * 3.75
        y = 1.8 + (idx // 3) * 1.55
        _v31_add_micro_icon(slide, str(idx + 1), x, y + 0.18, SECTION_COLORS[idx])
        add_text(slide, item, x + 0.68, y + 0.24, 2.65, 0.32, size=15, color=COLORS["text"], bold=True)
    add_insight_band(slide, "提出前チェック", "未確認事項を減らし、顧客説明に進められる状態にします。", 0.95, 5.58, 11.4, 0.64)
    add_footer(slide, slide_data.slide_no)


def render_v31_closing(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_section_label(slide, "次回打ち合わせ", 0.9, 0.72, fill=COLORS["navy"], color=COLORS["white"])
    add_title(slide, slide_data.title or "次回打ち合わせで合意すること", 1.02, 1.32, 10.8, 0.64, size=34, color=COLORS["navy"])
    items = _v31_items(slide_data, 4, ["対象範囲を確認", "評価基準を合意", "開始日を決定", "必要資料を共有"], limit=28)
    for idx, item in enumerate(items[:4]):
        _v31_add_value_card(slide, ["確認", "合意", "決定", "共有"][idx], item, 0.95 + idx * 2.88, 2.38, 2.5, 1.68, SECTION_COLORS[idx], number=str(idx + 1))
    add_insight_band(slide, "営業アクション", "次回は範囲・KPI・日程を合意し、PoC開始判断へ進めます。", 0.95, 5.46, 11.4, 0.72)
    add_footer(slide, slide_data.slide_no)


def render_v31_executive_message(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _v31_add_header(slide, slide_data, "経営判断", COLORS["blue"])
    items = _v31_items(slide_data, 6, [context.concept, "背景", "課題", "提案", "期待効果", "次アクション"], limit=28)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.95, 1.62, 11.35, 1.18, fill=COLORS["navy"], line=COLORS["navy"])
    add_text(slide, "提案の結論", 1.22, 1.94, 1.72, 0.2, size=13, color=COLORS["teal_light"], bold=True)
    add_text(slide, _v31_text(items[0], 46), 3.08, 1.84, 8.56, 0.38, size=25, color=COLORS["white"], bold=True)
    labels = ["背景", "課題", "提案", "効果", "ROI", "次回"]
    bodies = [items[1], items[2], items[3], items[4], "PoCで測定", items[5]]
    for idx, body in enumerate(bodies[:6]):
        x = 0.95 + (idx % 3) * 3.88
        y = 3.28 + (idx // 3) * 1.12
        _v31_add_value_card(slide, labels[idx], body, x, y, 3.28, 0.78, SECTION_COLORS[idx % len(SECTION_COLORS)])
    add_footer(slide, slide_data.slide_no)


def render_v31_current_state_map(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _v31_add_header(slide, slide_data, "現状整理", COLORS["teal"])
    items = _v31_items(slide_data, 5, ["現状", "課題", "影響", "改善機会", "確認事項"], limit=28)
    labels = ["現状", "課題", "影響", "機会", "確認"]
    for idx, item in enumerate(items[:5]):
        x = 0.95 + idx * 2.3
        _v31_add_micro_icon(slide, labels[idx], x + 0.7, 1.86, SECTION_COLORS[idx % len(SECTION_COLORS)])
        _v31_add_value_card(slide, labels[idx], item, x, 2.6, 1.9, 1.34, SECTION_COLORS[idx % len(SECTION_COLORS)])
        if idx < 4:
            _v31_add_arrow(slide, x + 1.92, 3.02)
    add_insight_band(slide, "提案へのつながり", "事実から課題、影響、改善機会へ順に整理します。", 0.95, 5.58, 11.4, 0.64)
    add_footer(slide, slide_data.slide_no)


def render_v31_problem_structure(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _v31_add_header(slide, slide_data, "課題構造", COLORS["red"])
    items = _v31_items(slide_data, 5, ["表面課題", "業務要因", "根本原因", "事業影響", "解決方向"], limit=30)
    _v31_add_value_card(slide, "表面課題", items[0], 0.95, 1.92, 3.0, 1.0, COLORS["orange"])
    _v31_add_value_card(slide, "業務要因", items[1], 0.95, 3.55, 3.0, 1.0, COLORS["purple"])
    _v31_add_arrow(slide, 4.2, 2.9)
    _v31_add_value_card(slide, "根本原因", items[2], 4.82, 2.56, 3.18, 1.24, COLORS["red"], fill=COLORS["red_light"], number="1")
    _v31_add_arrow(slide, 8.22, 2.9)
    _v31_add_value_card(slide, "事業影響", items[3], 8.95, 1.92, 3.0, 1.0, COLORS["blue"])
    _v31_add_value_card(slide, "解決方向", items[4], 8.95, 3.55, 3.0, 1.0, COLORS["green"])
    add_insight_band(slide, "課題認識", "症状ではなく原因に手を打ち、効果測定へつなげます。", 0.95, 5.6, 11.4, 0.62)
    add_footer(slide, slide_data.slide_no)


def render_v31_before_after(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _v31_add_header(slide, slide_data, "Before / After", COLORS["blue"])
    items = _v31_items(slide_data, 5, ["人手中心", "属人化", "AI候補提示", "連携", "改善学習"], limit=28)
    _v31_add_value_card(slide, "Before", items[0], 0.95, 1.92, 3.2, 2.05, COLORS["muted"], fill=COLORS["canvas"])
    _v31_add_value_card(slide, "Change", items[1], 5.02, 1.92, 3.2, 2.05, COLORS["blue"], fill=COLORS["blue_light"])
    _v31_add_value_card(slide, "After", items[2], 9.08, 1.92, 3.2, 2.05, COLORS["green"], fill=COLORS["green_light"])
    _v31_add_arrow(slide, 4.34, 2.68)
    _v31_add_arrow(slide, 8.4, 2.68)
    add_insight_band(slide, "運用の要点", f"{items[3]} / {items[4]}", 0.95, 5.38, 11.4, 0.72)
    add_footer(slide, slide_data.slide_no)


def render_v31_strategic_options(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _v31_add_header(slide, slide_data, "推奨方針", COLORS["blue"])
    items = _v31_items(slide_data, 4, ["現状維持", "部分改善", "段階導入", "推奨理由"], limit=32)
    labels = ["現状維持", "部分改善", "推奨案", "判断理由"]
    for idx, item in enumerate(items[:4]):
        fill = COLORS["blue_light"] if idx == 2 else COLORS["white"]
        _v31_add_value_card(slide, labels[idx], item, 0.95 + (idx % 2) * 5.78, 1.76 + (idx // 2) * 1.7, 5.22, 1.14, SECTION_COLORS[idx], fill=fill, number=str(idx + 1))
    add_insight_band(slide, "本提案の立ち位置", context.winning_strategy, 0.95, 5.6, 11.4, 0.62)
    add_footer(slide, slide_data.slide_no)


def render_v31_value_proposition(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _v31_add_header(slide, slide_data, "提供価値", COLORS["green"])
    items = _v31_items(slide_data, 4, ["工数削減", "品質安定", "判断迅速化", "継続改善"], limit=28)
    labels = ["経営価値", "業務価値", "実行価値"]
    widths = [8.7, 6.4, 4.1]
    for idx, width in enumerate(widths):
        x = 6.66 - width / 2
        y = 1.88 + idx * 0.98
        add_shape(slide, MSO_SHAPE.TRAPEZOID, x, y, width, 0.72, fill=[COLORS["green"], COLORS["blue"], COLORS["teal"]][idx], line=COLORS["white"])
        add_text(slide, labels[idx], x + 0.2, y + 0.24, width - 0.4, 0.16, size=15, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
    for idx, item in enumerate(items[:4]):
        _v31_add_value_card(slide, f"価値 {idx + 1}", item, 0.95 + idx * 2.88, 4.75, 2.5, 0.92, SECTION_COLORS[idx])
    add_footer(slide, slide_data.slide_no)


def render_v31_roi_logic(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _v31_add_header(slide, slide_data, "ROI", COLORS["orange"])
    items = _v31_items(slide_data, 4, ["投資範囲", "業務効果", "回収判断", "将来効果"], limit=12)
    labels = ["投資", "効果", "回収", "将来効果"]
    for idx, item in enumerate(items[:4]):
        x = 0.98 + idx * 2.88
        add_metric_card(slide, labels[idx], _v31_text(item, 12), x, 1.78, 2.5, 1.26, SECTION_COLORS[idx])
        add_shape(slide, MSO_SHAPE.DOWN_ARROW, x + 1.02, 3.28, 0.48, 0.52, fill=SECTION_COLORS[idx], line=SECTION_COLORS[idx])
        _v31_add_value_card(slide, ["範囲", "測定", "判断", "拡張"][idx], ["費用範囲を明確化", "効果をPoCで測定", "回収条件を合意", "改善を継続"][idx], x, 4.02, 2.5, 0.82, SECTION_COLORS[idx])
        if idx < 3:
            _v31_add_arrow(slide, x + 2.52, 2.26)
    add_insight_band(slide, "説明方針", "根拠のない効果額は置かず、測定可能な効果から判断します。", 0.95, 5.58, 11.4, 0.64)
    add_footer(slide, slide_data.slide_no)


def render_v31_competitive_positioning(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _v31_add_header(slide, slide_data, "競合に対する位置取り", COLORS["purple"])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 6.52, 1.78, 0.03, 3.6, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 1.05, 3.58, 10.95, 0.03, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    add_text(slide, "提案適合度 高", 5.35, 1.44, 2.4, 0.22, size=12, color=COLORS["muted"], bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "導入負荷 低", 0.95, 3.34, 1.7, 0.22, size=12, color=COLORS["muted"], bold=True)
    _v31_add_value_card(slide, "推奨ポジション", context.winning_strategy, 7.0, 2.04, 3.82, 1.12, COLORS["blue"], fill=COLORS["blue_light"])
    items = _v31_items(slide_data, 3, ["想定競合", "差別化", "確認事項"], limit=30)
    for idx, item in enumerate(items[:3]):
        _v31_add_value_card(slide, ["想定競合", "差別化", "確認事項"][idx], item, 1.05 + idx * 3.68, 5.02, 3.22, 0.82, SECTION_COLORS[idx])
    add_footer(slide, slide_data.slide_no)


def render_v31_layered_architecture(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _v31_add_header(slide, slide_data, "システム構成", COLORS["blue"])
    items = _v31_items(slide_data, 5, ["入力", "AI処理", "人の確認", "連携", "改善運用"], limit=22)
    labels = ["入力", "AI", "確認", "連携", "改善"]
    for idx, item in enumerate(items[:5]):
        x = 0.86 + idx * 2.45
        accent = SECTION_COLORS[idx % len(SECTION_COLORS)]
        _v31_add_micro_icon(slide, labels[idx], x + 0.74, 1.86, accent)
        _v31_add_value_card(slide, labels[idx], item, x, 2.62, 2.02, 1.42, accent)
        if idx < 4:
            _v31_add_arrow(slide, x + 2.02, 3.0)
    add_shape(slide, MSO_SHAPE.CIRCULAR_ARROW, 9.25, 4.48, 1.1, 0.72, fill=COLORS["teal_light"], line=COLORS["teal"])
    add_insight_band(slide, "改善ループ", "確認・修正履歴を活用し、人の判断を支援します。", 0.95, 5.58, 11.4, 0.64)
    add_footer(slide, slide_data.slide_no)


def render_v31_workstream_roadmap(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _v31_add_header(slide, slide_data, "実行ロードマップ", COLORS["teal"])
    phases = _v31_items(slide_data, 5, ["要件整理", "データ確認", "検証", "現場確認", "本番判断"], limit=20)
    lanes = ["作業", "成果物", "判断"]
    for lane_idx, lane in enumerate(lanes):
        y = 1.82 + lane_idx * 1.12
        add_text(slide, lane, 0.92, y + 0.14, 1.1, 0.22, size=12, color=COLORS["muted"], bold=True)
        add_shape(slide, MSO_SHAPE.RECTANGLE, 2.05, y + 0.26, 9.98, 0.04, fill=COLORS["line"], line=COLORS["line"])
        for idx, phase in enumerate(phases[:5]):
            x = 2.1 + idx * 1.95
            label = phase if lane_idx == 0 else ("成果物" if lane_idx == 1 else "判断")
            add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 1.56, 0.46, fill=SECTION_COLORS[idx], line=SECTION_COLORS[idx])
            add_text(slide, _trim(label, 8), x + 0.08, y + 0.15, 1.4, 0.12, size=9, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
    add_insight_band(slide, "意思決定地点", "各フェーズで成果物と判断ポイントを確認します。", 0.95, 5.65, 11.4, 0.62)
    add_footer(slide, slide_data.slide_no)


def render_v31_risk_heatmap(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _v31_add_header(slide, slide_data, "リスクと対策", COLORS["red"])
    items = _v31_items(slide_data, 4, ["精度不足", "データ不足", "運用負荷", "連携条件"], limit=28)
    colors = [COLORS["green_light"], COLORS["orange_light"], COLORS["red_light"], COLORS["blue_light"]]
    labels = ["低", "中", "高", "要確認"]
    for row in range(2):
        for col in range(2):
            idx = row * 2 + col
            x = 1.08 + col * 3.52
            y = 1.86 + row * 1.54
            add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 3.1, 1.08, fill=colors[idx], line=COLORS["white"])
            add_text(slide, labels[idx], x + 0.2, y + 0.18, 0.82, 0.2, size=13, color=SECTION_COLORS[idx], bold=True)
            add_text(slide, items[idx], x + 1.02, y + 0.2, 1.78, 0.34, size=13, color=COLORS["text"], bold=True)
    add_side_panel(slide, "対策方針", ["事前評価", "PoC確認", "運用設計", "段階導入"], 8.45, 1.86, 3.5, 2.66, COLORS["teal"])
    add_insight_band(slide, "提出前の見方", "リスクは隠さず、確認方法と対策を同時に示します。", 0.95, 5.65, 11.4, 0.62)
    add_footer(slide, slide_data.slide_no)


def render_v31_governance_map(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _v31_add_header(slide, slide_data, "推進体制", COLORS["purple"])
    items = _v31_items(slide_data, 5, ["意思決定", "業務責任", "IT連携", "現場確認", "支援"], limit=24)
    add_shape(slide, MSO_SHAPE.OVAL, 5.56, 2.38, 2.04, 1.36, fill=COLORS["navy"], line=COLORS["navy"])
    add_text(slide, "共同推進", 5.9, 2.88, 1.36, 0.18, size=15, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
    positions = [(1.0, 1.62), (9.15, 1.62), (1.0, 4.45), (9.15, 4.45), (5.05, 4.88)]
    for idx, (x, y) in enumerate(positions):
        _v31_add_value_card(slide, ["決裁", "業務", "IT", "現場", "支援"][idx], items[idx], x, y, 3.0, 0.9, SECTION_COLORS[idx])
    add_footer(slide, slide_data.slide_no)


def render_v31_cost_breakdown(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _v31_add_header(slide, slide_data, "概算見積", COLORS["orange"])
    add_metric_card(slide, "概算費用", _v31_text(context.estimate.total_label, 14), 0.95, 1.52, 3.2, 1.14, COLORS["orange"])
    add_metric_card(slide, "予算適合", _v31_text(context.estimate.budget_fit, 14), 4.58, 1.52, 3.2, 1.14, COLORS["green"])
    _v31_add_value_card(slide, "ROI説明", "必須範囲を確保し、推奨・任意を段階判断します。", 8.2, 1.52, 3.9, 1.14, COLORS["blue"])
    columns = [
        ("必須", context.estimate.required or ["要件整理", "検証"], COLORS["blue"], COLORS["blue_light"]),
        ("推奨", context.estimate.recommended or ["効果測定", "運用支援"], COLORS["green"], COLORS["green_light"]),
        ("任意", context.estimate.optional or ["追加連携", "拡張支援"], COLORS["orange"], COLORS["orange_light"]),
    ]
    for idx, (title, items, accent, fill) in enumerate(columns):
        x = 0.92 + idx * 4.08
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 3.24, 3.58, 1.82, fill=fill, line=COLORS["white"])
        add_text(slide, title, x + 0.22, 3.52, 3.12, 0.24, size=17, color=accent, bold=True, align=PP_ALIGN.CENTER)
        add_bullet_list(slide, [_v31_text(item, 20) for item in items], x + 0.38, 3.98, 2.72, 0.68, max_items=2, size=12)
    add_footer(slide, slide_data.slide_no)


def render_v31_scope_definition(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    _v31_add_header(slide, slide_data, "提案範囲", COLORS["blue"])
    items = _v31_items(slide_data, 6, ["対象", "対象外", "前提", "確認", "変更条件", "次回合意"], limit=26)
    labels = ["対象", "対象外", "前提", "確認", "変更", "合意"]
    for idx, item in enumerate(items[:6]):
        x = 0.95 + (idx % 3) * 3.88
        y = 1.76 + (idx // 3) * 1.52
        _v31_add_value_card(slide, labels[idx], item, x, y, 3.35, 1.0, SECTION_COLORS[idx % len(SECTION_COLORS)])
    add_insight_band(slide, "範囲管理", "対象と対象外を明確にし、見積・日程の前提を揃えます。", 0.95, 5.58, 11.4, 0.64)
    add_footer(slide, slide_data.slide_no)


def render_v31_section_divider(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT, fill=COLORS["navy"], line=COLORS["navy"])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 0.18, SLIDE_HEIGHT, fill=COLORS["teal"], line=COLORS["teal"])
    add_text(slide, "章立て", 0.95, 1.55, 2.0, 0.24, size=13, color=COLORS["teal_light"], bold=True)
    add_text(slide, _v31_text(slide_data.title or "提案の章", 30), 0.95, 2.16, 8.8, 0.82, size=38, color=COLORS["white"], bold=True)
    add_text(slide, _v31_text(chapter_message(slide_data.title, context), 58), 0.98, 3.34, 8.2, 0.48, size=17, color=COLORS["teal_light"], bold=True)
    add_footer(slide, slide_data.slide_no)


def _v4_title(value: str, limit: int = 34) -> str:
    return _v31_text(value, limit)


def _v4_items(slide_data: PowerPointSlide, count: int, fallback: list[str], *, limit: int = 26) -> list[str]:
    return [_v31_text(item, limit) for item in _layout_items(slide_data, count, fallback)]


def _v4_frame(slide, slide_data: PowerPointSlide, section: str, accent: str = COLORS["blue"]) -> None:
    set_background(slide)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 0.12, SLIDE_HEIGHT, fill=accent, line=accent)
    add_section_label(slide, section, 0.9, 0.55, fill=accent, color=COLORS["white"])
    add_title(slide, _v4_title(slide_data.title or section), 0.9, 0.96, 11.4, 0.56, size=35, color=COLORS["navy"])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0.9, 1.6, 11.48, 0.018, fill=COLORS["line"], line=COLORS["line"])


def _v4_soft_panel(slide, x: float, y: float, w: float, h: float, *, fill: str = COLORS["white"], line: str = COLORS["line"]):
    return add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, line=line)


def _v4_chip(slide, label: str, x: float, y: float, w: float, accent: str) -> None:
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, 0.34, fill=accent, line=accent)
    add_text(slide, _v31_text(label, 16), x + 0.12, y + 0.1, w - 0.24, 0.11, size=8, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)


def _v4_micro_text(slide, title: str, body: str, x: float, y: float, w: float, accent: str) -> None:
    add_shape(slide, MSO_SHAPE.RECTANGLE, x, y + 0.03, 0.05, 0.62, fill=accent, line=accent)
    add_text(slide, _v31_text(title, 12), x + 0.18, y, w - 0.18, 0.18, size=11, color=accent, bold=True)
    add_text(slide, _v31_text(body, 26), x + 0.18, y + 0.28, w - 0.18, 0.28, size=12, color=COLORS["text"], bold=True)


def _v4_connector(slide, x: float, y: float, w: float, accent: str = COLORS["line_dark"]) -> None:
    add_shape(slide, MSO_SHAPE.RECTANGLE, x, y + 0.17, w, 0.035, fill=accent, line=accent)
    add_shape(slide, MSO_SHAPE.CHEVRON, x + w - 0.08, y, 0.32, 0.38, fill=accent, line=accent)


def render_v4_hero_cover(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    theme = resolve_template_colors(context.design_template)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT, fill="081526", line="081526")
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 0.14, SLIDE_HEIGHT, fill=theme["accent"], line=theme["accent"])
    add_shape(slide, MSO_SHAPE.OVAL, 8.1, 0.72, 3.1, 3.1, fill="173B66", line="173B66")
    add_shape(slide, MSO_SHAPE.OVAL, 10.25, 3.35, 1.9, 1.9, fill="0E7490", line="0E7490")
    add_shape(slide, MSO_SHAPE.OVAL, 8.7, 4.96, 1.35, 1.35, fill="1D4ED8", line="1D4ED8")
    for idx, label in enumerate(icon_labels_for_category(context.proposal_category)[:3]):
        add_icon_badge(slide, label, 8.78 + idx * 0.92, 2.02 + idx * 0.86, SECTION_COLORS[idx], size=0.74)
    section_label = "Web制作提案書" if context.proposal_category == "web" else f"{context.proposal_label} 提案書"
    add_section_label(slide, normalize_customer_facing_text(section_label, limit=18), 0.82, 0.76, fill=theme["accent"], color=COLORS["white"])
    add_title(slide, _v4_title(slide_data.title or data.deck_title, 30), 0.82, 1.56, 7.4, 1.2, size=46, color=COLORS["white"])
    add_text(slide, f"{context.client_name} 御中", 0.86, 3.18, 6.6, 0.34, size=18, color=COLORS["teal_light"], bold=True)
    add_text(slide, _v31_text(f"{context.concept}の意思決定資料", 32), 0.86, 3.74, 6.7, 0.38, size=20, color=COLORS["white"], bold=True)
    add_text(slide, f"提案日 {date.today().strftime('%Y.%m.%d')}", 0.86, 5.86, 3.0, 0.2, size=11, color=COLORS["teal_light"])
    add_text(slide, "ProposalPilot / AI営業秘書", 0.86, 6.48, 4.0, 0.22, size=11, color=COLORS["teal_light"])


def render_v4_executive_brief(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    _v4_frame(slide, slide_data, "経営判断", COLORS["blue"])
    items = _v4_items(slide_data, 6, [context.concept, "背景", "課題", "提案", "効果", "次回"], limit=24)
    _v4_soft_panel(slide, 0.95, 1.92, 11.35, 1.35, fill=COLORS["navy"], line=COLORS["navy"])
    add_text(slide, "結論", 1.24, 2.16, 0.8, 0.18, size=12, color=COLORS["teal_light"], bold=True)
    add_text(slide, _v31_text(items[0], 42), 2.2, 2.07, 8.95, 0.34, size=24, color=COLORS["white"], bold=True)
    for idx, (label, body) in enumerate(zip(["背景", "課題", "提案", "効果", "ROI", "次回"], [*items[1:5], "PoCで測定", items[5]])):
        x = 1.0 + (idx % 3) * 3.85
        y = 3.88 + (idx // 3) * 0.86
        _v4_micro_text(slide, label, body, x, y, 3.2, SECTION_COLORS[idx % len(SECTION_COLORS)])
    add_footer(slide, slide_data.slide_no)


def render_v4_issue_tree(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    _v4_frame(slide, slide_data, "課題構造", COLORS["red"])
    items = _v4_items(slide_data, 5, ["表面課題", "業務要因", "根本原因", "事業影響", "解決方向"], limit=22)
    add_shape(slide, MSO_SHAPE.OVAL, 4.9, 2.38, 3.1, 1.38, fill=COLORS["red_light"], line=COLORS["red"])
    add_text(slide, "根本原因", 5.38, 2.72, 2.08, 0.2, size=14, color=COLORS["red"], bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, items[2], 5.18, 3.05, 2.5, 0.22, size=13, color=COLORS["text"], bold=True, align=PP_ALIGN.CENTER)
    branches = [("表面課題", items[0], 1.0, 1.95, COLORS["orange"]), ("業務要因", items[1], 1.0, 4.15, COLORS["purple"]), ("事業影響", items[3], 9.18, 1.95, COLORS["blue"]), ("解決方向", items[4], 9.18, 4.15, COLORS["green"])]
    for title, body, x, y, accent in branches:
        _v4_micro_text(slide, title, body, x, y, 2.7, accent)
        _v4_connector(slide, x + (2.85 if x < 5 else -1.15), y + 0.16, 1.02, accent)
    add_insight_band(slide, "見方", "表面課題ではなく原因に手を打つことで、効果測定へつなげます。", 0.95, 5.72, 11.4, 0.58)
    add_footer(slide, slide_data.slide_no)


def render_v4_before_after(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    _v4_frame(slide, slide_data, "変化設計", COLORS["blue"])
    items = _v4_items(slide_data, 5, ["人手中心", "属人化", "AI候補提示", "連携", "改善学習"], limit=24)
    panels = [("Before", items[0], 0.95, COLORS["canvas"], COLORS["muted"]), ("Shift", items[1], 5.0, COLORS["blue_light"], COLORS["blue"]), ("After", items[2], 9.05, COLORS["green_light"], COLORS["green"])]
    for idx, (label, body, x, fill, accent) in enumerate(panels):
        _v4_soft_panel(slide, x, 2.02, 3.18, 2.4, fill=fill, line=accent)
        _v4_chip(slide, label, x + 0.52, 2.38, 2.05, accent)
        add_text(slide, body, x + 0.44, 3.16, 2.3, 0.36, size=19, color=COLORS["navy"], bold=True, align=PP_ALIGN.CENTER)
        if idx < 2:
            _v4_connector(slide, x + 3.32, 3.02, 0.58)
    add_insight_band(slide, "運用の要点", f"{items[3]} / {items[4]}", 0.95, 5.54, 11.4, 0.62)
    add_footer(slide, slide_data.slide_no)


def render_v4_value_pyramid(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    _v4_frame(slide, slide_data, "提供価値", COLORS["green"])
    items = _v4_items(slide_data, 4, ["経営価値", "業務価値", "実行価値", "継続改善"], limit=22)
    layers = [(3.0, 2.02, 7.3, "経営価値", COLORS["green"], items[0]), (3.75, 3.02, 5.8, "業務価値", COLORS["blue"], items[1]), (4.55, 4.02, 4.2, "実行価値", COLORS["teal"], items[2])]
    for x, y, w, title, accent, body in layers:
        add_shape(slide, MSO_SHAPE.TRAPEZOID, x, y, w, 0.78, fill=accent, line=COLORS["white"])
        add_text(slide, f"{title}: {_v31_text(body, 16)}", x + 0.22, y + 0.24, w - 0.44, 0.16, size=14, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
    _v4_micro_text(slide, "継続改善", items[3], 1.0, 5.24, 4.0, COLORS["purple"])
    add_footer(slide, slide_data.slide_no)


def render_v4_competitive_positioning(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    _v4_frame(slide, slide_data, "競合差別化", COLORS["purple"])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 6.56, 1.96, 0.025, 3.5, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 1.16, 3.7, 10.8, 0.025, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    add_text(slide, "提案適合度 高", 5.35, 1.7, 2.4, 0.18, size=11, color=COLORS["muted"], bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "導入負荷 低", 1.1, 3.42, 1.6, 0.18, size=11, color=COLORS["muted"], bold=True)
    add_shape(slide, MSO_SHAPE.OVAL, 7.48, 2.48, 1.42, 1.0, fill=COLORS["blue_light"], line=COLORS["blue"])
    add_text(slide, "推奨\nポジション", 7.64, 2.75, 1.1, 0.36, size=13, color=COLORS["blue"], bold=True, align=PP_ALIGN.CENTER)
    items = _v4_items(slide_data, 3, ["想定競合", "差別化", "確認事項"], limit=24)
    for idx, item in enumerate(items[:3]):
        _v4_micro_text(slide, ["想定競合", "差別化", "確認事項"][idx], item, 1.05 + idx * 3.68, 5.56, 3.0, SECTION_COLORS[idx])
    add_footer(slide, slide_data.slide_no)


def render_v4_kpi_dashboard(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    _v4_frame(slide, slide_data, "KPI設計", COLORS["green"])
    items = _v4_items(slide_data, 4, ["現状値", "目標値", "測定方法", "判定基準"], limit=14)
    for idx, item in enumerate(items[:4]):
        x = 0.98 + idx * 2.88
        add_shape(slide, MSO_SHAPE.OVAL, x + 0.52, 2.0, 1.42, 1.42, fill=COLORS["white"], line=SECTION_COLORS[idx])
        add_text(slide, ["現状", "目標", "測定", "判定"][idx], x + 0.72, 2.42, 1.0, 0.16, size=13, color=SECTION_COLORS[idx], bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, item, x + 0.08, 3.7, 2.3, 0.3, size=15, color=COLORS["navy"], bold=True, align=PP_ALIGN.CENTER)
        add_shape(slide, MSO_SHAPE.RECTANGLE, x + 0.28, 4.4, 1.95, 0.1, fill=COLORS["line"], line=COLORS["line"])
        add_shape(slide, MSO_SHAPE.RECTANGLE, x + 0.28, 4.4, 0.62 + idx * 0.28, 0.1, fill=SECTION_COLORS[idx], line=SECTION_COLORS[idx])
    add_insight_band(slide, "SMART設計", "現状・目標・測定方法・判定基準をPoCで合意します。", 0.95, 5.64, 11.4, 0.58)
    add_footer(slide, slide_data.slide_no)


def render_v4_roi_dashboard(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    _v4_frame(slide, slide_data, "ROI設計", COLORS["orange"])
    items = _v4_items(slide_data, 4, ["投資範囲", "業務効果", "回収判断", "将来効果"], limit=16)
    labels = ["投資", "効果", "回収", "将来"]
    for idx, (label, item) in enumerate(zip(labels, items)):
        x = 1.0 + idx * 2.86
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.0 + idx * 0.18, 2.26, 0.86 + idx * 0.18, fill=COLORS["white"], line=SECTION_COLORS[idx])
        add_text(slide, label, x + 0.18, 2.18 + idx * 0.18, 1.88, 0.18, size=13, color=SECTION_COLORS[idx], bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, item, x + 0.18, 2.56 + idx * 0.18, 1.88, 0.22, size=13, color=COLORS["text"], bold=True, align=PP_ALIGN.CENTER)
        if idx < 3:
            _v4_connector(slide, x + 2.35, 2.5 + idx * 0.18, 0.45)
    add_insight_band(slide, "判断方針", "効果額を断定せず、測定可能な効果から回収条件を合意します。", 0.95, 5.58, 11.4, 0.62)
    add_footer(slide, slide_data.slide_no)


def render_v4_architecture_stack(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    _v4_frame(slide, slide_data, "構成設計", COLORS["blue"])
    items = _v4_items(slide_data, 5, ["入力", "AI処理", "人の確認", "連携", "改善運用"], limit=18)
    for idx, item in enumerate(items[:5]):
        y = 1.9 + idx * 0.68
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 2.1 + idx * 0.46, y, 8.5 - idx * 0.92, 0.48, fill=[COLORS["blue_light"], COLORS["teal_light"], COLORS["green_light"], COLORS["purple_light"], COLORS["orange_light"]][idx], line=COLORS["white"])
        add_text(slide, f"{['Input','AI','Review','Integration','Learning'][idx]}  {item}", 2.38 + idx * 0.46, y + 0.16, 7.7 - idx * 0.92, 0.1, size=11, color=COLORS["navy"], bold=True, align=PP_ALIGN.CENTER)
    add_shape(slide, MSO_SHAPE.CIRCULAR_ARROW, 9.6, 4.75, 0.98, 0.62, fill=COLORS["teal_light"], line=COLORS["teal"])
    add_insight_band(slide, "運用改善", "確認・修正履歴を次回改善に活用します。", 0.95, 5.7, 11.4, 0.56)
    add_footer(slide, slide_data.slide_no)


def render_v4_swimlane_roadmap(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    _v4_frame(slide, slide_data, "実行計画", COLORS["teal"])
    phases = _v4_items(slide_data, 5, ["要件整理", "設計", "検証", "導入", "定着"], limit=14)
    lanes = ["作業", "成果物", "判断"]
    for lane_idx, lane in enumerate(lanes):
        y = 2.02 + lane_idx * 1.05
        add_text(slide, lane, 0.95, y + 0.14, 1.0, 0.18, size=12, color=COLORS["muted"], bold=True)
        add_shape(slide, MSO_SHAPE.RECTANGLE, 2.0, y + 0.27, 10.0, 0.025, fill=COLORS["line"], line=COLORS["line"])
        for idx, phase in enumerate(phases[:5]):
            x = 2.1 + idx * 1.94
            add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 1.42, 0.42, fill=SECTION_COLORS[idx], line=SECTION_COLORS[idx])
            label = phase if lane_idx == 0 else ("成果物" if lane_idx == 1 else "判断")
            add_text(slide, _trim(label, 8), x + 0.05, y + 0.14, 1.32, 0.1, size=8, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
    add_insight_band(slide, "意思決定地点", "各フェーズで成果物と判断ポイントを確認します。", 0.95, 5.64, 11.4, 0.58)
    add_footer(slide, slide_data.slide_no)


def render_v4_risk_radar(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    _v4_frame(slide, slide_data, "リスク対策", COLORS["red"])
    items = _v4_items(slide_data, 4, ["精度", "データ", "運用", "連携"], limit=16)
    center = (6.35, 3.35)
    rings = [(2.8, COLORS["line"]), (1.9, COLORS["line"]), (1.0, COLORS["line"])]
    for size, line in rings:
        add_shape(slide, MSO_SHAPE.OVAL, center[0] - size / 2, center[1] - size / 2, size, size, fill=COLORS["white"], line=line)
    positions = [(4.95, 2.05), (7.52, 2.15), (4.85, 4.28), (7.6, 4.26)]
    for idx, (x, y) in enumerate(positions):
        add_shape(slide, MSO_SHAPE.OVAL, x, y, 0.58, 0.58, fill=SECTION_COLORS[idx], line=SECTION_COLORS[idx])
        add_text(slide, items[idx], x - 0.35, y + 0.74, 1.28, 0.16, size=11, color=COLORS["text"], bold=True, align=PP_ALIGN.CENTER)
    add_insight_band(slide, "対策方針", "リスクは隠さず、確認方法と対策を同時に示します。", 0.95, 5.64, 11.4, 0.58)
    add_footer(slide, slide_data.slide_no)


def render_v4_investment_mix(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    _v4_frame(slide, slide_data, "投資判断", COLORS["orange"])
    add_text(slide, _v31_text(context.estimate.total_label, 18), 1.05, 1.95, 3.2, 0.38, size=24, color=COLORS["orange"], bold=True)
    add_text(slide, _v31_text(context.estimate.budget_fit, 18), 4.55, 1.95, 3.0, 0.38, size=22, color=COLORS["green"], bold=True)
    columns = [("必須", context.estimate.required or ["要件整理", "検証"], COLORS["blue"]), ("推奨", context.estimate.recommended or ["効果測定", "運用支援"], COLORS["green"]), ("任意", context.estimate.optional or ["追加連携", "拡張支援"], COLORS["orange"])]
    x = 1.0
    for title, items, accent in columns:
        w = 2.7 if title != "推奨" else 3.4
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 3.15, w, 0.72, fill=accent, line=accent)
        add_text(slide, title, x + 0.2, 3.4, w - 0.4, 0.12, size=12, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, " / ".join(_v31_text(item, 8) for item in items[:2]), x + 0.2, 4.18, w - 0.4, 0.2, size=11, color=COLORS["text"], bold=True, align=PP_ALIGN.CENTER)
        x += w + 0.28
    add_insight_band(slide, "説明方針", "必須範囲を確保し、推奨・任意は段階的に判断します。", 0.95, 5.64, 11.4, 0.58)
    add_footer(slide, slide_data.slide_no)


def render_v4_next_meeting(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    _v4_frame(slide, slide_data, "次回打ち合わせ", COLORS["blue"])
    items = _v4_items(slide_data, 4, ["範囲を確認", "KPIを合意", "日程を決定", "資料を共有"], limit=18)
    add_text(slide, "次回は、開始判断に必要な3点を合意します", 1.05, 2.0, 9.6, 0.42, size=25, color=COLORS["navy"], bold=True)
    for idx, item in enumerate(items[:4]):
        x = 1.05 + idx * 2.82
        add_shape(slide, MSO_SHAPE.OVAL, x + 0.78, 3.04, 0.72, 0.72, fill=SECTION_COLORS[idx], line=SECTION_COLORS[idx])
        add_text(slide, str(idx + 1), x + 0.78, 3.26, 0.72, 0.12, size=10, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, item, x + 0.08, 4.04, 2.1, 0.22, size=15, color=COLORS["text"], bold=True, align=PP_ALIGN.CENTER)
    add_insight_band(slide, "営業アクション", "範囲・KPI・日程を合意し、PoC開始判断へ進めます。", 0.95, 5.64, 11.4, 0.58)
    add_footer(slide, slide_data.slide_no)

def _v5_items(slide_data: PowerPointSlide, count: int, fallback: list[str] | None = None, *, limit: int = 15) -> list[str]:
    base = fallback or ["背景", "課題", "提案", "効果", "次の一手"]
    items = ensure_items(unique_items(slide_data.bullets, count), base, count)
    return [_trim(normalize_customer_facing_text(item), limit) for item in items]


def _v5_title(slide_data: PowerPointSlide, *, limit: int = 24) -> str:
    return _trim(normalize_customer_facing_text(slide_data.title or "提案の要点"), limit)


def _v5_bg(slide, *, dark: bool, accent: str) -> tuple[str, str]:
    if dark:
        add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT, fill="07111F", line="07111F")
        add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 0.08, SLIDE_HEIGHT, fill=accent, line=accent)
        return COLORS["white"], "D0D5DD"
    set_background(slide)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 0.08, SLIDE_HEIGHT, fill=accent, line=accent)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0.78, 1.36, 11.75, 0.025, fill=COLORS["line"], line=COLORS["line"])
    return COLORS["navy"], COLORS["muted"]


def _v5_header(slide, slide_data: PowerPointSlide, label: str, accent: str, *, dark: bool = False, y: float = 0.48) -> None:
    title_color, muted = _v5_bg(slide, dark=dark, accent=accent)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.78, y, 1.38, 0.3, fill=accent, line=accent)
    add_text(slide, _trim(label, 10), 0.92, y + 0.08, 1.1, 0.1, size=8, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, _v5_title(slide_data), 0.78, y + 0.52, 8.8, 0.52, size=35, color=title_color, bold=True)
    add_text(slide, f"{slide_data.slide_no:02}", 11.8, 6.88, 0.5, 0.16, size=10, color=muted, align=PP_ALIGN.RIGHT)
    add_text(slide, "ProposalPilot / AI営業秘書", 0.78, 6.88, 2.4, 0.16, size=8, color=muted)


def _v5_caption(slide, text: str, x: float, y: float, w: float, *, color: str = COLORS["muted"], align: PP_ALIGN | None = PP_ALIGN.CENTER) -> None:
    add_text(slide, _trim(normalize_customer_facing_text(text), 18), x, y, w, 0.18, size=10, color=color, bold=True, align=align)


def _v5_has(text: str, terms: tuple[str, ...]) -> bool:
    return any(term in text for term in terms)


def _v5_render_cover(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT, fill="06111F", line="06111F")
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 0.1, SLIDE_HEIGHT, fill=COLORS["blue"], line=COLORS["blue"])
    add_shape(slide, MSO_SHAPE.OVAL, 8.2, 0.8, 2.2, 2.2, fill="102A43", line="102A43")
    add_shape(slide, MSO_SHAPE.OVAL, 9.74, 2.38, 1.28, 1.28, fill=COLORS["teal"], line=COLORS["teal"])
    add_shape(slide, MSO_SHAPE.OVAL, 7.62, 4.62, 1.05, 1.05, fill=COLORS["blue"], line=COLORS["blue"])
    add_shape(slide, MSO_SHAPE.OVAL, 8.72, 1.72, 0.48, 0.48, fill=COLORS["blue"], line=COLORS["blue"])
    add_shape(slide, MSO_SHAPE.OVAL, 10.12, 3.02, 0.36, 0.36, fill=COLORS["green"], line=COLORS["green"])
    add_text(slide, _trim(slide_data.title or data.deck_title, 26), 0.86, 1.34, 6.4, 0.72, size=42, color=COLORS["white"], bold=True)
    add_text(slide, _trim(context.concept or "成果につながる提案", 28), 0.9, 2.58, 5.8, 0.34, size=17, color="D0E2FF", bold=True)
    add_text(slide, _trim(data.client_name or context.customer_name, 30), 0.9, 3.08, 5.6, 0.3, size=14, color="AEBFD6")
    add_text(slide, f"ISSUED {date.today():%Y.%m.%d}", 0.9, 5.68, 2.0, 0.16, size=8, color="AEBFD6")
    add_text(slide, "ProposalPilot / AI営業秘書", 0.9, 6.25, 2.5, 0.16, size=8, color="AEBFD6")


def _v5_render_orbit(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext, *, accent: str, dark: bool = False) -> None:
    slide = blank_slide(prs)
    _v5_header(slide, slide_data, "INSIGHT", accent, dark=dark)
    title_color = COLORS["white"] if dark else COLORS["text"]
    items = _v5_items(slide_data, 4, ["背景", "課題", "提案", "効果"], limit=14)
    cx, cy = 6.52, 3.68
    add_shape(slide, MSO_SHAPE.OVAL, cx - 1.0, cy - 0.72, 2.0, 1.18, fill="FFF1F0", line="FADBD8")
    add_text(slide, _trim(items[0], 14), cx - 0.72, cy - 0.36, 1.44, 0.18, size=12, color=COLORS["red"], bold=True, align=PP_ALIGN.CENTER)
    positions = [(2.1, 2.45), (9.5, 2.05), (2.4, 5.1), (9.55, 4.95)]
    for idx, (x, y) in enumerate(positions):
        col = SECTION_COLORS[idx % len(SECTION_COLORS)]
        add_shape(slide, MSO_SHAPE.RIGHT_ARROW, x + (1.18 if x < 6 else -1.2), y + 0.16, 1.1, 0.28, fill=col, line=col)
        add_shape(slide, MSO_SHAPE.OVAL, x, y, 0.38, 0.38, fill=col, line=col)
        _v5_caption(slide, items[idx], x + 0.48, y + 0.1, 1.95, color=title_color, align=PP_ALIGN.LEFT)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.95, 6.02, 11.3, 0.46, fill=COLORS["navy"] if not dark else "12233B", line=COLORS["navy"] if not dark else "12233B")
    add_text(slide, _trim(context.winning_strategy or context.concept, 50), 1.24, 6.16, 10.7, 0.12, size=12, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)


def _v5_render_chain(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext, *, accent: str) -> None:
    slide = blank_slide(prs)
    _v5_header(slide, slide_data, "FLOW", accent)
    items = _v5_items(slide_data, 5, ["現状", "課題", "解決策", "導入", "効果"], limit=13)
    y = 3.0
    for idx, item in enumerate(items[:5]):
        x = 1.0 + idx * 2.3
        col = SECTION_COLORS[idx % len(SECTION_COLORS)]
        add_shape(slide, MSO_SHAPE.CHEVRON, x, y, 1.86, 0.72, fill=col, line=col)
        add_text(slide, item, x + 0.18, y + 0.25, 1.25, 0.14, size=12, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
        add_shape(slide, MSO_SHAPE.RECTANGLE, x + 0.52, y + 1.08, 0.82, 0.045, fill=col, line=col)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 1.0, 5.62, 10.9, 0.46, fill=COLORS["navy"], line=COLORS["navy"])
    add_text(slide, _trim(context.concept or "導入判断に必要な流れを整理します", 48), 1.26, 5.76, 10.35, 0.12, size=12, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)


def _v5_render_swimlane(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext, *, accent: str) -> None:
    slide = blank_slide(prs)
    _v5_header(slide, slide_data, "ROADMAP", accent)
    phases = _v5_items(slide_data, 5, ["要件", "設計", "検証", "導入", "改善"], limit=12)
    lanes = ["業務", "システム", "判断"]
    for row, lane in enumerate(lanes):
        y = 2.08 + row * 1.18
        add_text(slide, lane, 1.05, y + 0.22, 0.7, 0.12, size=11, color=COLORS["muted"], bold=True)
        add_shape(slide, MSO_SHAPE.RECTANGLE, 1.88, y + 0.39, 9.85, 0.02, fill=COLORS["line"], line=COLORS["line"])
        for idx, phase in enumerate(phases[:5]):
            x = 2.05 + idx * 1.9
            col = SECTION_COLORS[(idx + row) % len(SECTION_COLORS)]
            add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y + 0.12, 1.35, 0.38, fill=col, line=col)
            add_text(slide, phase, x + 0.08, y + 0.25, 1.18, 0.08, size=8, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 1.0, 5.8, 10.96, 0.42, fill=COLORS["navy"], line=COLORS["navy"])
    add_text(slide, "各フェーズで成果物と判断ポイントを確認します", 1.28, 5.93, 10.4, 0.12, size=12, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)


def _v5_render_matrix(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext, *, accent: str) -> None:
    slide = blank_slide(prs)
    _v5_header(slide, slide_data, "POSITION", accent)
    items = _v5_items(slide_data, 4, ["現状", "競合", "提案", "優位性"], limit=13)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 2.3, 2.08, 7.35, 0.03, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 5.95, 1.45, 0.03, 4.4, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    add_text(slide, "事業影響 大", 5.25, 1.2, 1.4, 0.16, size=10, color=COLORS["muted"], bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "実行難度 高", 9.2, 4.02, 1.4, 0.16, size=10, color=COLORS["muted"], bold=True, align=PP_ALIGN.CENTER)
    points = [(3.0, 4.9), (7.95, 4.42), (7.5, 2.0), (4.0, 2.75)]
    for idx, (x, y) in enumerate(points):
        col = SECTION_COLORS[idx % len(SECTION_COLORS)]
        add_shape(slide, MSO_SHAPE.OVAL, x, y, 0.5, 0.5, fill=col, line=col)
        _v5_caption(slide, items[idx], x - 0.5, y + 0.62, 1.5, color=COLORS["text"])


def _v5_render_dashboard(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext, *, accent: str) -> None:
    slide = blank_slide(prs)
    _v5_header(slide, slide_data, "DASHBOARD", accent)
    items = _v5_items(slide_data, 4, ["現状", "目標", "測定", "判定"], limit=10)
    for idx, item in enumerate(items[:4]):
        x = 1.15 + idx * 2.85
        col = SECTION_COLORS[idx % len(SECTION_COLORS)]
        add_shape(slide, MSO_SHAPE.OVAL, x, 2.16, 1.06, 1.06, fill="FFFFFF", line=col)
        add_shape(slide, MSO_SHAPE.RECTANGLE, x + 0.2, 3.74, 0.7 + idx * 0.14, 0.08, fill=col, line=col)
        add_text(slide, item, x + 0.08, 2.58, 0.9, 0.12, size=9, color=col, bold=True, align=PP_ALIGN.CENTER)
        _v5_caption(slide, item, x - 0.42, 3.5, 1.95, color=COLORS["text"])
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 1.0, 5.72, 11.0, 0.44, fill=COLORS["navy"], line=COLORS["navy"])
    add_text(slide, "SMART形式で現状・目標・測定・判定を確認します", 1.22, 5.86, 10.55, 0.12, size=12, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)


def _v5_render_investment(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext, *, accent: str) -> None:
    slide = blank_slide(prs)
    _v5_header(slide, slide_data, "INVESTMENT", accent)
    add_text(slide, _trim(context.estimate.total_label, 18), 1.05, 1.82, 2.7, 0.32, size=25, color=COLORS["orange"], bold=True)
    add_text(slide, _trim(context.estimate.budget_fit, 22), 4.15, 1.86, 3.2, 0.28, size=18, color=COLORS["green"], bold=True)
    tiers = [("必須", COLORS["blue"], 2.8), ("推奨", COLORS["green"], 3.4), ("任意", COLORS["orange"], 2.1)]
    x = 1.1
    for label, col, width in tiers:
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 3.1, width, 0.72, fill=col, line=col)
        add_text(slide, label, x + 0.14, 3.36, width - 0.28, 0.12, size=13, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
        add_shape(slide, MSO_SHAPE.RECTANGLE, x + 0.2, 4.24, width - 0.4, 0.08, fill=col, line=col)
        x += width + 0.34
    add_shape(slide, MSO_SHAPE.RIGHT_ARROW, 1.05, 5.45, 10.8, 0.42, fill="D0D5DD", line="D0D5DD")
    add_text(slide, "投資 → 効果 → 回収 → 将来効果", 2.2, 5.58, 8.0, 0.12, size=14, color=COLORS["navy"], bold=True, align=PP_ALIGN.CENTER)


def _v5_render_radar(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext, *, accent: str) -> None:
    slide = blank_slide(prs)
    _v5_header(slide, slide_data, "RISK", accent)
    items = _v5_items(slide_data, 5, ["価格", "納期", "ROI", "体制", "運用"], limit=11)
    cx, cy = 6.5, 3.55
    for radius in [0.7, 1.15, 1.6]:
        add_shape(slide, MSO_SHAPE.OVAL, cx - radius, cy - radius, radius * 2, radius * 2, fill="FFFFFF", line=COLORS["line"])
    dots = [(5.0, 2.2), (7.7, 2.2), (8.2, 4.55), (6.4, 5.25), (4.7, 4.45)]
    for idx, (x, y) in enumerate(dots):
        col = SECTION_COLORS[idx % len(SECTION_COLORS)]
        add_shape(slide, MSO_SHAPE.OVAL, x, y, 0.38, 0.38, fill=col, line=col)
        _v5_caption(slide, items[idx], x - 0.58, y + 0.5, 1.55, color=COLORS["text"])
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 1.05, 5.88, 11.0, 0.42, fill=COLORS["navy"], line=COLORS["navy"])
    add_text(slide, "懸念は先回りして説明し、判断リスクを下げます", 1.32, 6.02, 10.48, 0.12, size=12, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)


def _v5_render_pyramid(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext, *, accent: str) -> None:
    slide = blank_slide(prs)
    _v5_header(slide, slide_data, "VALUE", accent)
    items = _v5_items(slide_data, 4, ["経営効果", "業務効果", "実行施策", "確認事項"], limit=18)
    widths = [5.4, 6.6, 7.8, 9.0]
    y = 2.1
    for idx, width in enumerate(widths):
        x = (SLIDE_WIDTH - width) / 2
        col = SECTION_COLORS[idx % len(SECTION_COLORS)]
        add_shape(slide, MSO_SHAPE.TRAPEZOID, x, y + idx * 0.76, width, 0.54, fill=col, line=col)
        add_text(slide, items[idx], x + 0.26, y + idx * 0.76 + 0.18, width - 0.52, 0.1, size=11, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, _trim(context.winning_strategy or context.concept, 42), 2.08, 5.58, 9.0, 0.24, size=14, color=COLORS["navy"], bold=True, align=PP_ALIGN.CENTER)


def _v5_render_journey(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext, *, accent: str) -> None:
    slide = blank_slide(prs)
    _v5_header(slide, slide_data, "JOURNEY", accent)
    items = _v5_items(slide_data, 5, ["認知", "理解", "比較", "判断", "行動"], limit=12)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 1.25, 3.16, 9.2, 0.04, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    for idx, item in enumerate(items[:5]):
        x = 1.36 + idx * 2.15
        col = SECTION_COLORS[idx % len(SECTION_COLORS)]
        add_shape(slide, MSO_SHAPE.OVAL, x, 2.74, 0.72, 0.72, fill=col, line=col)
        add_text(slide, str(idx + 1), x, 2.98, 0.72, 0.1, size=10, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
        _v5_caption(slide, item, x - 0.48, 3.72, 1.72, color=COLORS["text"])
    add_side_panel(slide, "判断ポイント", _v5_items(slide_data, 3, ["不安を解消", "比較材料を提示", "行動を短縮"], limit=13), 10.72, 2.3, 1.58, 2.6, accent)


def _v5_render_architecture(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext, *, accent: str) -> None:
    slide = blank_slide(prs)
    _v5_header(slide, slide_data, "STRUCTURE", accent)
    items = _v5_items(slide_data, 5, ["Input", "AI/業務", "確認", "連携", "運用"], limit=14)
    fills = ["EAF2FF", "EAFBFF", "E7F7EF", "F0EEFF", "FFF7E6"]
    for idx, item in enumerate(items[:5]):
        y = 1.86 + idx * 0.72
        col = SECTION_COLORS[idx % len(SECTION_COLORS)]
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 2.4 + idx * 0.22, y, 6.8 - idx * 0.44, 0.45, fill=fills[idx], line=COLORS["line"])
        add_shape(slide, MSO_SHAPE.RECTANGLE, 2.55 + idx * 0.22, y + 0.16, 0.44, 0.05, fill=col, line=col)
        add_text(slide, item, 3.16 + idx * 0.22, y + 0.13, 5.0 - idx * 0.44, 0.1, size=11, color=COLORS["navy"], bold=True, align=PP_ALIGN.CENTER)
    add_shape(slide, MSO_SHAPE.CIRCULAR_ARROW, 8.78, 4.46, 1.25, 1.0, fill=COLORS["teal_light"], line=COLORS["teal"])
    _v5_caption(slide, "改善ループ", 8.55, 5.55, 1.7, color=COLORS["teal"])


def _v5_render_issue_tree(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext, *, accent: str) -> None:
    slide = blank_slide(prs)
    _v5_header(slide, slide_data, "ISSUE", accent)
    items = _v5_items(slide_data, 5, ["症状", "原因", "影響", "打ち手", "成果"], limit=14)
    add_shape(slide, MSO_SHAPE.OVAL, 5.34, 2.76, 2.22, 0.95, fill="FFF1F0", line="FADBD8")
    add_text(slide, items[1], 5.55, 3.08, 1.8, 0.12, size=12, color=COLORS["red"], bold=True, align=PP_ALIGN.CENTER)
    nodes = [(1.45, 2.05), (1.45, 4.72), (9.6, 2.05), (9.6, 4.72)]
    for idx, (x, y) in enumerate(nodes):
        col = SECTION_COLORS[idx % len(SECTION_COLORS)]
        add_shape(slide, MSO_SHAPE.RIGHT_ARROW, x + (1.15 if x < 6 else -1.1), y + 0.25, 1.2, 0.28, fill=col, line=col)
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 1.55, 0.72, fill="FFFFFF", line=col)
        label = items[idx if idx < len(items) else -1]
        add_text(slide, label, x + 0.14, y + 0.25, 1.27, 0.1, size=10, color=COLORS["navy"], bold=True, align=PP_ALIGN.CENTER)


def _v5_render_split(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext, *, accent: str) -> None:
    slide = blank_slide(prs)
    _v5_header(slide, slide_data, "SHIFT", accent)
    items = _v5_items(slide_data, 4, ["現状", "課題", "改善後", "効果"], limit=14)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0.96, 2.0, 4.65, 3.45, fill="F8FAFC", line=COLORS["line"])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 7.32, 2.0, 4.65, 3.45, fill="EAFBFF", line=COLORS["teal"])
    add_shape(slide, MSO_SHAPE.RIGHT_ARROW, 5.78, 3.35, 1.25, 0.52, fill=accent, line=accent)
    add_text(slide, items[0], 1.35, 2.45, 3.85, 0.18, size=16, color=COLORS["muted"], bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, items[2], 7.68, 2.45, 3.85, 0.18, size=16, color=COLORS["navy"], bold=True, align=PP_ALIGN.CENTER)
    _v5_caption(slide, items[1], 1.3, 4.38, 3.85, color=COLORS["red"])
    _v5_caption(slide, items[3], 7.66, 4.38, 3.85, color=COLORS["green"])


def _v5_render_closing(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext, *, accent: str) -> None:
    slide = blank_slide(prs)
    _v5_header(slide, slide_data, "NEXT", accent, dark=True, y=0.55)
    items = _v5_items(slide_data, 4, ["範囲合意", "KPI確認", "日程決定", "開始判断"], limit=12)
    add_text(slide, "次回打ち合わせで、実施範囲と判断条件を合意します", 1.05, 2.03, 9.6, 0.44, size=28, color=COLORS["white"], bold=True)
    for idx, item in enumerate(items[:4]):
        x = 1.08 + idx * 2.68
        col = SECTION_COLORS[idx % len(SECTION_COLORS)]
        add_shape(slide, MSO_SHAPE.OVAL, x, 3.55, 0.72, 0.72, fill=col, line=col)
        add_text(slide, str(idx + 1), x, 3.78, 0.72, 0.12, size=10, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, item, x - 0.2, 4.58, 1.25, 0.18, size=13, color="D0E2FF", bold=True, align=PP_ALIGN.CENTER)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 1.05, 5.72, 10.7, 0.5, fill=accent, line=accent)
    add_text(slide, "提案方針・KPI・開始条件を確認", 1.35, 5.88, 10.1, 0.12, size=14, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)


def _v5_render_executive_canvas(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext, *, accent: str) -> None:
    slide = blank_slide(prs)
    _v5_header(slide, slide_data, "EXECUTIVE", accent, dark=True)
    items = _v5_items(slide_data, 6, ["Background", "Issue", "Answer", "Impact", "ROI", "Next"], limit=16)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.95, 2.0, 4.3, 2.45, fill="101F33", line="223856")
    add_text(slide, _trim(items[0], 26), 1.25, 2.28, 3.7, 0.55, size=24, color=COLORS["white"], bold=True)
    add_text(slide, _trim(context.winning_strategy or context.concept, 42), 1.28, 3.42, 3.62, 0.34, size=15, color="BFD7FF", bold=True)
    positions = [(6.05, 1.92), (8.65, 2.58), (6.15, 4.08), (8.85, 4.78)]
    for idx, (x, y) in enumerate(positions):
        col = SECTION_COLORS[idx % len(SECTION_COLORS)]
        add_shape(slide, MSO_SHAPE.OVAL, x, y, 0.76, 0.76, fill=col, line=col)
        add_text(slide, str(idx + 1), x, y + 0.25, 0.76, 0.1, size=10, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
        add_shape(slide, MSO_SHAPE.RECTANGLE, x + 0.88, y + 0.36, 1.55, 0.035, fill=col, line=col)
        _v5_caption(slide, items[idx + 1], x + 0.88, y + 0.58, 1.8, color="D0D5DD", align=PP_ALIGN.LEFT)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.95, 5.88, 10.95, 0.42, fill=accent, line=accent)
    add_text(slide, _trim(items[-1], 48), 1.24, 6.02, 10.4, 0.1, size=12, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)


def _v5_render_decision_lens(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext, *, accent: str) -> None:
    slide = blank_slide(prs)
    _v5_header(slide, slide_data, "DECISION", accent)
    items = _v5_items(slide_data, 5, ["Now", "Why", "How", "Value", "Proof"], limit=14)
    add_shape(slide, MSO_SHAPE.OVAL, 4.85, 1.92, 3.45, 3.45, fill="EAF2FF", line=COLORS["blue"])
    add_shape(slide, MSO_SHAPE.OVAL, 5.55, 2.62, 2.05, 2.05, fill=COLORS["white"], line=accent)
    add_text(slide, _trim(items[0], 16), 5.84, 3.3, 1.45, 0.22, size=18, color=COLORS["navy"], bold=True, align=PP_ALIGN.CENTER)
    spokes = [(2.05, 2.15), (9.0, 2.1), (2.28, 4.75), (8.9, 4.75)]
    for idx, (x, y) in enumerate(spokes):
        col = SECTION_COLORS[idx % len(SECTION_COLORS)]
        add_shape(slide, MSO_SHAPE.RECTANGLE, 4.05 if x < 5 else 8.1, y + 0.28, 0.78, 0.04, fill=col, line=col)
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 1.95, 0.72, fill=COLORS["white"], line=col)
        add_text(slide, _trim(items[idx + 1], 14), x + 0.15, y + 0.24, 1.65, 0.1, size=11, color=COLORS["navy"], bold=True, align=PP_ALIGN.CENTER)


def _v5_render_staircase(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext, *, accent: str) -> None:
    slide = blank_slide(prs)
    _v5_header(slide, slide_data, "STORYLINE", accent)
    items = _v5_items(slide_data, 5, ["Current", "Issue", "Shift", "Action", "Outcome"], limit=13)
    for idx, item in enumerate(items[:5]):
        x = 1.02 + idx * 2.05
        y = 4.85 - idx * 0.58
        col = SECTION_COLORS[idx % len(SECTION_COLORS)]
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 1.72, 0.48 + idx * 0.16, fill=col, line=col)
        add_text(slide, item, x + 0.12, y + 0.18, 1.48, 0.1, size=10, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
        if idx < 4:
            add_shape(slide, MSO_SHAPE.RIGHT_ARROW, x + 1.78, y - 0.18, 0.52, 0.32, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    add_text(slide, _trim(context.concept, 46), 1.05, 5.88, 10.5, 0.22, size=14, color=COLORS["navy"], bold=True, align=PP_ALIGN.CENTER)


def _v5_render_funnel(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext, *, accent: str) -> None:
    slide = blank_slide(prs)
    _v5_header(slide, slide_data, "SELECTION", accent)
    items = _v5_items(slide_data, 5, ["Criteria", "Risk", "Fit", "Value", "Decision"], limit=14)
    widths = [9.6, 7.7, 5.9, 4.15, 2.55]
    for idx, width in enumerate(widths):
        x = (SLIDE_WIDTH - width) / 2
        y = 1.9 + idx * 0.78
        col = SECTION_COLORS[idx % len(SECTION_COLORS)]
        add_shape(slide, MSO_SHAPE.TRAPEZOID, x, y, width, 0.55, fill=col, line=COLORS["white"])
        add_text(slide, _trim(items[idx], 18), x + 0.28, y + 0.18, width - 0.56, 0.1, size=10, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
    add_shape(slide, MSO_SHAPE.OVAL, 5.82, 5.75, 1.0, 0.52, fill=accent, line=accent)
    add_text(slide, "GO", 5.82, 5.91, 1.0, 0.1, size=12, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)


def _v5_render_meter_wall(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext, *, accent: str) -> None:
    slide = blank_slide(prs)
    _v5_header(slide, slide_data, "MEASURE", accent)
    items = _v5_items(slide_data, 4, ["Baseline", "Target", "Method", "Owner"], limit=12)
    for idx, item in enumerate(items[:4]):
        x = 1.2 + idx * 2.75
        col = SECTION_COLORS[idx % len(SECTION_COLORS)]
        add_shape(slide, MSO_SHAPE.ARC, x + 0.28, 2.1, 1.7, 1.15, fill=COLORS["white"], line=col)
        add_shape(slide, MSO_SHAPE.OVAL, x + 0.9, 2.76, 0.36, 0.36, fill=col, line=col)
        add_text(slide, item, x + 0.05, 3.58, 2.2, 0.16, size=13, color=COLORS["navy"], bold=True, align=PP_ALIGN.CENTER)
        add_shape(slide, MSO_SHAPE.RECTANGLE, x + 0.35, 4.28, 1.55, 0.08, fill=COLORS["line"], line=COLORS["line"])
        add_shape(slide, MSO_SHAPE.RECTANGLE, x + 0.35, 4.28, 0.55 + idx * 0.26, 0.08, fill=col, line=col)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 1.05, 5.75, 10.95, 0.38, fill="F8FAFC", line=COLORS["line"])
    add_text(slide, "Measured in pilot before final scope", 1.3, 5.88, 10.4, 0.1, size=11, color=COLORS["muted"], bold=True, align=PP_ALIGN.CENTER)


def _v5_render_waterfall_roi(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext, *, accent: str) -> None:
    slide = blank_slide(prs)
    _v5_header(slide, slide_data, "ROI PATH", accent)
    labels = _v5_items(slide_data, 4, ["Investment", "Efficiency", "Payback", "Future"], limit=13)
    heights = [0.82, 1.18, 1.55, 2.0]
    base_y = 5.12
    for idx, height in enumerate(heights):
        x = 1.2 + idx * 2.65
        col = SECTION_COLORS[idx % len(SECTION_COLORS)]
        add_shape(slide, MSO_SHAPE.RECTANGLE, x, base_y - height, 1.75, height, fill=col, line=col)
        add_text(slide, labels[idx], x - 0.2, base_y + 0.24, 2.15, 0.15, size=11, color=COLORS["navy"], bold=True, align=PP_ALIGN.CENTER)
        if idx < 3:
            add_shape(slide, MSO_SHAPE.RIGHT_ARROW, x + 1.88, base_y - height - 0.06, 0.62, 0.32, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    add_text(slide, _trim(context.estimate.total_label, 22), 1.2, 1.82, 3.1, 0.34, size=24, color=COLORS["orange"], bold=True)
    add_text(slide, _trim(context.estimate.budget_fit, 24), 4.55, 1.88, 4.2, 0.26, size=17, color=COLORS["green"], bold=True)


def _v5_render_compass(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext, *, accent: str) -> None:
    slide = blank_slide(prs)
    _v5_header(slide, slide_data, "CONTROL", accent)
    items = _v5_items(slide_data, 4, ["Price", "Schedule", "ROI", "Operation"], limit=13)
    cx, cy = 6.48, 3.6
    add_shape(slide, MSO_SHAPE.OVAL, cx - 1.55, cy - 1.55, 3.1, 3.1, fill="FFFFFF", line=COLORS["line"])
    add_shape(slide, MSO_SHAPE.RECTANGLE, cx - 0.02, cy - 1.55, 0.04, 3.1, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    add_shape(slide, MSO_SHAPE.RECTANGLE, cx - 1.55, cy - 0.02, 3.1, 0.04, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    positions = [(6.12, 1.72), (8.28, 3.3), (6.12, 5.25), (3.85, 3.3)]
    for idx, (x, y) in enumerate(positions):
        col = SECTION_COLORS[idx % len(SECTION_COLORS)]
        add_shape(slide, MSO_SHAPE.OVAL, x, y, 0.72, 0.72, fill=col, line=col)
        _v5_caption(slide, items[idx], x - 0.65, y + 0.84, 2.0, color=COLORS["text"])
    add_text(slide, "Risk", cx - 0.5, cy - 0.12, 1.0, 0.12, size=11, color=accent, bold=True, align=PP_ALIGN.CENTER)


def _v5_render_value_chain(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext, *, accent: str) -> None:
    slide = blank_slide(prs)
    _v5_header(slide, slide_data, "VALUE CHAIN", accent)
    items = _v5_items(slide_data, 5, ["Input", "Design", "Build", "Launch", "Learn"], limit=12)
    for idx, item in enumerate(items[:5]):
        x = 1.0 + idx * 2.18
        col = SECTION_COLORS[idx % len(SECTION_COLORS)]
        add_shape(slide, MSO_SHAPE.RECTANGLE, x, 2.35, 1.45, 2.22, fill="F8FAFC", line=COLORS["line"])
        add_shape(slide, MSO_SHAPE.RECTANGLE, x, 4.18 - idx * 0.28, 1.45, 0.38 + idx * 0.28, fill=col, line=col)
        add_text(slide, item, x - 0.08, 4.88, 1.62, 0.14, size=10, color=COLORS["navy"], bold=True, align=PP_ALIGN.CENTER)
        if idx < 4:
            add_shape(slide, MSO_SHAPE.CHEVRON, x + 1.56, 3.16, 0.48, 0.42, fill=COLORS["line_dark"], line=COLORS["line_dark"])


def _v5_render_objection_wall(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext, *, accent: str) -> None:
    slide = blank_slide(prs)
    _v5_header(slide, slide_data, "OBJECTIONS", accent)
    items = _v5_items(slide_data, 6, ["Price", "ROI", "Schedule", "System", "Operation", "Security"], limit=12)
    for idx, item in enumerate(items[:6]):
        x = 1.1 + (idx % 3) * 3.48
        y = 2.0 + (idx // 3) * 1.42
        col = SECTION_COLORS[idx % len(SECTION_COLORS)]
        add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, 2.66, 0.72, fill=COLORS["white"], line=col)
        add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, 0.14, 0.72, fill=col, line=col)
        add_text(slide, item, x + 0.3, y + 0.25, 2.1, 0.1, size=11, color=COLORS["navy"], bold=True, align=PP_ALIGN.CENTER)
    add_shape(slide, MSO_SHAPE.RIGHT_ARROW, 5.74, 5.2, 1.12, 0.46, fill=accent, line=accent)
    add_text(slide, "answer before asked", 4.7, 5.88, 3.2, 0.12, size=11, color=COLORS["muted"], bold=True, align=PP_ALIGN.CENTER)


def _v5_render_fishbone(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext, *, accent: str) -> None:
    slide = blank_slide(prs)
    _v5_header(slide, slide_data, "ROOT CAUSE", accent)
    items = _v5_items(slide_data, 5, ["People", "Process", "Content", "Data", "Decision"], limit=13)
    add_shape(slide, MSO_SHAPE.RIGHT_ARROW, 2.0, 3.55, 7.85, 0.24, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    add_shape(slide, MSO_SHAPE.OVAL, 9.5, 3.14, 1.35, 1.05, fill="FFF1F0", line=COLORS["red"])
    add_text(slide, _trim(items[-1], 13), 9.68, 3.48, 1.0, 0.1, size=10, color=COLORS["red"], bold=True, align=PP_ALIGN.CENTER)
    bones = [(3.0, 2.42, 0.9), (4.35, 4.42, -0.9), (5.7, 2.42, 0.9), (7.05, 4.42, -0.9)]
    for idx, (x, y, direction) in enumerate(bones):
        col = SECTION_COLORS[idx % len(SECTION_COLORS)]
        add_shape(slide, MSO_SHAPE.RECTANGLE, x, 3.55 if direction > 0 else 3.72, 0.04, direction, fill=col, line=col)
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x - 0.68, y, 1.45, 0.44, fill=COLORS["white"], line=col)
        add_text(slide, items[idx], x - 0.58, y + 0.15, 1.25, 0.08, size=9, color=COLORS["navy"], bold=True, align=PP_ALIGN.CENTER)


def _v5_render_market_map(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext, *, accent: str) -> None:
    slide = blank_slide(prs)
    _v5_header(slide, slide_data, "MARKET MAP", accent)
    items = _v5_items(slide_data, 4, ["Current", "Competitor", "Our edge", "Decision"], limit=13)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 1.12, 2.02, 10.2, 3.55, fill="FFFFFF", line=COLORS["line"])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 1.12, 3.76, 10.2, 0.035, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 6.18, 2.02, 0.035, 3.55, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    coords = [(2.15, 4.42), (3.8, 2.66), (7.55, 2.56), (8.85, 4.42)]
    for idx, (x, y) in enumerate(coords):
        col = SECTION_COLORS[idx % len(SECTION_COLORS)]
        add_shape(slide, MSO_SHAPE.OVAL, x, y, 0.54, 0.54, fill=col, line=col)
        _v5_caption(slide, items[idx], x - 0.5, y + 0.68, 1.55, color=COLORS["text"])
    add_text(slide, "Impact", 1.2, 1.78, 1.2, 0.16, size=10, color=COLORS["muted"], bold=True)
    add_text(slide, "Feasibility", 9.35, 5.76, 1.7, 0.16, size=10, color=COLORS["muted"], bold=True)


def _v5_render_battlecard(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext, *, accent: str) -> None:
    slide = blank_slide(prs)
    _v5_header(slide, slide_data, "DIFFERENCE", accent)
    items = _v5_items(slide_data, 4, ["Market", "Competitor", "Our edge", "Check"], limit=14)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0.95, 2.05, 5.2, 3.6, fill="F8FAFC", line=COLORS["line"])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 6.42, 2.05, 5.2, 3.6, fill="EAFBFF", line=COLORS["teal"])
    add_text(slide, items[1], 1.28, 2.55, 4.5, 0.24, size=18, color=COLORS["muted"], bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, items[2], 6.75, 2.55, 4.5, 0.24, size=18, color=COLORS["navy"], bold=True, align=PP_ALIGN.CENTER)
    add_shape(slide, MSO_SHAPE.CHEVRON, 5.72, 3.54, 0.85, 0.55, fill=accent, line=accent)
    add_text(slide, "win", 5.72, 3.72, 0.85, 0.1, size=10, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, _trim(context.winning_strategy or items[3], 52), 2.0, 5.96, 8.8, 0.16, size=12, color=COLORS["text"], bold=True, align=PP_ALIGN.CENTER)


def _v5_render_audience_map(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext, *, accent: str) -> None:
    slide = blank_slide(prs)
    _v5_header(slide, slide_data, "AUDIENCE", accent)
    items = _v5_items(slide_data, 5, ["Executive", "Manager", "User", "IT", "Decision"], limit=12)
    center_x, center_y = 6.2, 3.55
    add_shape(slide, MSO_SHAPE.OVAL, center_x - 0.72, center_y - 0.72, 1.44, 1.44, fill=accent, line=accent)
    add_text(slide, items[-1], center_x - 0.55, center_y - 0.05, 1.1, 0.1, size=10, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
    positions = [(2.0, 2.05), (9.0, 2.0), (2.32, 5.0), (9.0, 5.02)]
    for idx, (x, y) in enumerate(positions):
        col = SECTION_COLORS[idx % len(SECTION_COLORS)]
        add_shape(slide, MSO_SHAPE.RECTANGLE, x + (1.55 if x < 6 else -1.2), y + 0.24, 2.0, 0.035, fill=col, line=col)
        add_shape(slide, MSO_SHAPE.OVAL, x, y, 0.68, 0.68, fill=col, line=col)
        _v5_caption(slide, items[idx], x + 0.8 if x < 6 else x - 1.65, y + 0.22, 1.55, color=COLORS["text"], align=PP_ALIGN.LEFT if x < 6 else PP_ALIGN.RIGHT)


def _v5_render_strategy_tree(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext, *, accent: str) -> None:
    slide = blank_slide(prs)
    _v5_header(slide, slide_data, "STRATEGY TREE", accent)
    items = _v5_items(slide_data, 7, ["Goal", "Acquire", "Convert", "Trust", "Proof", "Content", "CTA"], limit=12)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 5.15, 1.88, 2.08, 0.62, fill=accent, line=accent)
    add_text(slide, items[0], 5.35, 2.1, 1.68, 0.1, size=10, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
    for idx in range(3):
        x = 2.15 + idx * 3.36
        col = SECTION_COLORS[idx % len(SECTION_COLORS)]
        add_shape(slide, MSO_SHAPE.RECTANGLE, 6.16, 2.5, 0.04, 0.5, fill=COLORS["line_dark"], line=COLORS["line_dark"])
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 3.02, 2.15, 0.52, fill=COLORS["white"], line=col)
        add_text(slide, items[idx + 1], x + 0.12, 3.2, 1.9, 0.08, size=9, color=COLORS["navy"], bold=True, align=PP_ALIGN.CENTER)
        for child in range(2):
            y = 4.1 + child * 0.72
            add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x + 0.22, y, 1.72, 0.36, fill="F8FAFC", line=COLORS["line"])
            add_text(slide, items[4 + ((idx + child) % 3)], x + 0.34, y + 0.12, 1.48, 0.08, size=8, color=COLORS["text"], bold=True, align=PP_ALIGN.CENTER)


def _v5_render_sitemap_topology(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext, *, accent: str) -> None:
    slide = blank_slide(prs)
    _v5_header(slide, slide_data, "TOPOLOGY", accent)
    items = _v5_items(slide_data, 6, ["Home", "Service", "Case", "FAQ", "Contact", "Recruit"], limit=10)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 5.3, 1.92, 2.1, 0.5, fill=accent, line=accent)
    add_text(slide, items[0], 5.48, 2.08, 1.72, 0.1, size=10, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
    for idx, item in enumerate(items[1:6]):
        x = 1.35 + idx * 2.14
        y = 3.45 + (idx % 2) * 0.66
        col = SECTION_COLORS[idx % len(SECTION_COLORS)]
        add_shape(slide, MSO_SHAPE.RECTANGLE, 6.32, 2.42, 0.035, y - 2.42, fill=COLORS["line"], line=COLORS["line"])
        add_shape(slide, MSO_SHAPE.RECTANGLE, min(x + 0.78, 6.32), y - 0.05, abs(6.32 - (x + 0.78)), 0.035, fill=COLORS["line"], line=COLORS["line"])
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 1.55, 0.42, fill="FFFFFF", line=col)
        add_text(slide, item, x + 0.12, y + 0.14, 1.3, 0.08, size=8, color=COLORS["navy"], bold=True, align=PP_ALIGN.CENTER)


def _v5_render_content_architecture(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext, *, accent: str) -> None:
    slide = blank_slide(prs)
    _v5_header(slide, slide_data, "CONTENT SYSTEM", accent)
    items = _v5_items(slide_data, 5, ["Lead", "Proof", "FAQ", "Case", "CTA"], limit=12)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 1.0, 1.94, 3.7, 3.9, fill="07111F", line="07111F")
    add_text(slide, _trim(items[0], 16), 1.35, 2.42, 2.95, 0.4, size=24, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
    for idx, item in enumerate(items[1:5]):
        x = 5.25 + (idx % 2) * 3.05
        y = 2.0 + (idx // 2) * 1.72
        col = SECTION_COLORS[idx % len(SECTION_COLORS)]
        add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, 2.45, 1.08, fill="F8FAFC", line=COLORS["line"])
        add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, 2.45, 0.08, fill=col, line=col)
        add_text(slide, item, x + 0.22, y + 0.44, 2.0, 0.12, size=12, color=COLORS["navy"], bold=True, align=PP_ALIGN.CENTER)


def _v5_render_spotlight_summary(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext, *, accent: str) -> None:
    slide = blank_slide(prs)
    _v5_header(slide, slide_data, "SUMMARY", accent, dark=True)
    items = _v5_items(slide_data, 4, ["Purpose", "Approach", "Proof", "Action"], limit=14)
    add_shape(slide, MSO_SHAPE.OVAL, 1.15, 1.82, 3.25, 3.25, fill="101F33", line="223856")
    add_shape(slide, MSO_SHAPE.OVAL, 1.95, 2.62, 1.65, 1.65, fill=accent, line=accent)
    add_text(slide, "ONE", 2.23, 3.19, 1.1, 0.12, size=12, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, _trim(context.concept, 30), 1.35, 5.18, 2.85, 0.24, size=15, color="D0E2FF", bold=True, align=PP_ALIGN.CENTER)
    for idx, item in enumerate(items[:4]):
        y = 1.92 + idx * 0.92
        col = SECTION_COLORS[idx % len(SECTION_COLORS)]
        add_shape(slide, MSO_SHAPE.RECTANGLE, 5.28, y + 0.25, 4.7, 0.035, fill=col, line=col)
        add_shape(slide, MSO_SHAPE.OVAL, 10.18, y, 0.58, 0.58, fill=col, line=col)
        add_text(slide, item, 5.18, y + 0.06, 4.55, 0.16, size=14, color=COLORS["white"], bold=True, align=PP_ALIGN.RIGHT)


def _v5_render_landscape_bands(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext, *, accent: str) -> None:
    slide = blank_slide(prs)
    _v5_header(slide, slide_data, "LANDSCAPE", accent)
    items = _v5_items(slide_data, 5, ["Market", "Customer", "Competitor", "Edge", "Move"], limit=12)
    band_specs = [
        (1.0, 1.98, 10.6, 0.62, COLORS["blue"], 0),
        (1.45, 2.82, 9.7, 0.62, COLORS["teal"], 1),
        (1.9, 3.66, 8.8, 0.62, COLORS["green"], 2),
        (2.35, 4.5, 7.9, 0.62, COLORS["purple"], 3),
        (2.8, 5.34, 7.0, 0.62, COLORS["orange"], 4),
    ]
    for x, y, w, h, col, idx in band_specs:
        add_shape(slide, MSO_SHAPE.PARALLELOGRAM, x, y, w, h, fill=col, line=COLORS["white"])
        add_text(slide, items[idx], x + 0.42, y + 0.22, w - 0.84, 0.1, size=11, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, _trim(context.winning_strategy or context.concept, 48), 2.4, 6.12, 8.0, 0.16, size=12, color=COLORS["navy"], bold=True, align=PP_ALIGN.CENTER)


def render_v5_masterpiece_slide(
    prs: Presentation,
    slide_data: PowerPointSlide,
    data: PowerPointData,
    context: PptxContext,
    index: int,
) -> None:
    title = slide_data.title or ""
    accent = SECTION_COLORS[index % len(SECTION_COLORS)]
    if index == 0 or slide_data.layout == "title":
        _v5_render_cover(prs, slide_data, data, context)
    elif index == 1:
        _v5_render_executive_canvas(prs, slide_data, data, context, accent=accent)
    elif index == 2:
        _v5_render_decision_lens(prs, slide_data, data, context, accent=accent)
    elif index == 3:
        _v5_render_staircase(prs, slide_data, data, context, accent=accent)
    elif index == 4:
        _v5_render_funnel(prs, slide_data, data, context, accent=accent)
    elif index == 5:
        _v5_render_meter_wall(prs, slide_data, data, context, accent=accent)
    elif index == 6:
        _v5_render_waterfall_roi(prs, slide_data, data, context, accent=accent)
    elif index == 7:
        _v5_render_compass(prs, slide_data, data, context, accent=accent)
    elif index == 8:
        _v5_render_objection_wall(prs, slide_data, data, context, accent=accent)
    elif index == 9:
        _v5_render_pyramid(prs, slide_data, data, context, accent=accent)
    elif index == 10:
        _v5_render_value_chain(prs, slide_data, data, context, accent=accent)
    elif index == 11:
        _v5_render_radar(prs, slide_data, data, context, accent=accent)
    elif index == 12:
        _v5_render_investment(prs, slide_data, data, context, accent=accent)
    elif index == 13:
        _v5_render_swimlane(prs, slide_data, data, context, accent=accent)
    elif index == 14:
        _v5_render_spotlight_summary(prs, slide_data, data, context, accent=accent)
    elif index == 15:
        _v5_render_market_map(prs, slide_data, data, context, accent=accent)
    elif index == 16:
        _v5_render_fishbone(prs, slide_data, data, context, accent=accent)
    elif index == 17:
        _v5_render_landscape_bands(prs, slide_data, data, context, accent=accent)
    elif index == 18:
        _v5_render_battlecard(prs, slide_data, data, context, accent=accent)
    elif index == 19:
        _v5_render_audience_map(prs, slide_data, data, context, accent=accent)
    elif index == 20:
        _v5_render_journey(prs, slide_data, data, context, accent=accent)
    elif index == 21:
        _v5_render_strategy_tree(prs, slide_data, data, context, accent=accent)
    elif index == 22:
        _v5_render_sitemap_topology(prs, slide_data, data, context, accent=accent)
    elif index == 23:
        _v5_render_content_architecture(prs, slide_data, data, context, accent=accent)
    elif index >= 24:
        _v5_render_closing(prs, slide_data, data, context, accent=accent)
    elif _v5_has(title, ("\u4eca\u5f8c", "\u6b21\u56de", "\u30a2\u30af\u30b7\u30e7\u30f3")):
        _v5_render_closing(prs, slide_data, data, context, accent=accent)
    elif _v5_has(title, ("\u8cbb\u7528", "\u898b\u7a4d", "\u4e88\u7b97", "\u6982\u7b97", "\u5fc5\u9808", "\u63a8\u5968", "\u30aa\u30d7\u30b7\u30e7\u30f3", "\u4e07\u5186")):
        _v5_render_investment(prs, slide_data, data, context, accent=accent)
    elif "ROI" in title.upper() or _v5_has(title, ("\u6295\u8cc7", "\u56de\u53ce", "\u8cbb\u7528\u5bfe\u52b9\u679c")):
        _v5_render_investment(prs, slide_data, data, context, accent=accent)
    elif "KPI" in title.upper() or _v5_has(title, ("\u6307\u6a19", "\u6e2c\u5b9a", "\u76ee\u6a19")):
        _v5_render_dashboard(prs, slide_data, data, context, accent=accent)
    elif _v5_has(title, ("\u30ed\u30fc\u30c9\u30de\u30c3\u30d7", "\u30b9\u30b1\u30b8\u30e5\u30fc\u30eb", "\u30d5\u30a7\u30fc\u30ba", "\u5de5\u7a0b", "\u5c0e\u5165")):
        _v5_render_swimlane(prs, slide_data, data, context, accent=accent)
    elif "FAQ" in title.upper() or _v5_has(title, ("\u30ea\u30b9\u30af", "\u61f8\u5ff5", "\u56de\u7b54", "\u30bb\u30ad\u30e5\u30ea\u30c6\u30a3")):
        _v5_render_radar(prs, slide_data, data, context, accent=accent)
    elif _v5_has(title, ("\u7af6\u5408", "\u5e02\u5834", "\u5dee\u5225\u5316", "\u52dd\u3061\u7b4b", "\u6bd4\u8f03")):
        _v5_render_matrix(prs, slide_data, data, context, accent=accent)
    elif _v5_has(title, ("\u8ab2\u984c", "\u539f\u56e0", "\u8981\u56e0", "\u30dc\u30c8\u30eb\u30cd\u30c3\u30af")):
        _v5_render_issue_tree(prs, slide_data, data, context, accent=accent)
    elif _v5_has(title, ("\u30b8\u30e3\u30fc\u30cb\u30fc", "\u5c0e\u7dda", "\u6d41\u308c")):
        _v5_render_journey(prs, slide_data, data, context, accent=accent)
    elif any(term in title.upper() for term in ("API", "AI", "OCR")) or _v5_has(title, ("\u30b5\u30a4\u30c8\u30de\u30c3\u30d7", "\u69cb\u6210", "\u9023\u643a", "\u30b7\u30b9\u30c6\u30e0", "\u30a2\u30fc\u30ad\u30c6\u30af\u30c1\u30e3")):
        _v5_render_architecture(prs, slide_data, data, context, accent=accent)
    elif "BEFORE" in title.upper() or "AFTER" in title.upper() or _v5_has(title, ("\u6539\u5584\u5f8c", "\u73fe\u72b6")):
        _v5_render_split(prs, slide_data, data, context, accent=accent)
    elif _v5_has(title, ("\u6226\u7565", "\u65b9\u91dd", "\u4fa1\u5024", "\u30bf\u30fc\u30b2\u30c3\u30c8", "\u30e6\u30fc\u30b6\u30fc", "\u30b3\u30f3\u30c6\u30f3\u30c4")):
        _v5_render_pyramid(prs, slide_data, data, context, accent=accent)
    else:
        variants = [
            _v5_render_orbit,
            _v5_render_chain,
            _v5_render_split,
            _v5_render_pyramid,
            _v5_render_journey,
            _v5_render_architecture,
        ]
        renderer = variants[index % len(variants)]
        if renderer is _v5_render_orbit:
            renderer(prs, slide_data, data, context, accent=accent, dark=index % 4 == 0)
        else:
            renderer(prs, slide_data, data, context, accent=accent)


LAYOUT_RENDERER_REGISTRY = {
    "LAYOUT-001": render_v31_title_only,
    "LAYOUT-002": render_v31_title_body,
    "LAYOUT-003": render_v31_two_column,
    "LAYOUT-004": render_v31_three_column,
    "LAYOUT-005": render_v4_hero_cover,
    "LAYOUT-006": render_v31_kpi_cards,
    "LAYOUT-007": render_v31_comparison_cards,
    "LAYOUT-008": render_v31_timeline,
    "LAYOUT-009": render_v31_roadmap,
    "LAYOUT-010": render_v31_flow,
    "LAYOUT-011": render_v31_matrix,
    "LAYOUT-012": render_v31_quote,
    "LAYOUT-013": render_v31_image_left,
    "LAYOUT-014": render_v31_image_right,
    "LAYOUT-015": render_v31_large_number,
    "LAYOUT-016": render_v31_checklist,
    "LAYOUT-017": render_v31_closing,
    "LAYOUT-018": render_v31_executive_message,
    "LAYOUT-019": render_v31_current_state_map,
    "LAYOUT-020": render_v31_problem_structure,
    "LAYOUT-021": render_v31_before_after,
    "LAYOUT-022": render_v31_strategic_options,
    "LAYOUT-023": render_v31_value_proposition,
    "LAYOUT-024": render_v31_kpi_cards,
    "LAYOUT-025": render_v31_roi_logic,
    "LAYOUT-026": render_v31_competitive_positioning,
    "LAYOUT-027": render_v31_layered_architecture,
    "LAYOUT-028": render_v31_workstream_roadmap,
    "LAYOUT-029": render_v31_risk_heatmap,
    "LAYOUT-030": render_v31_governance_map,
    "LAYOUT-031": render_v31_cost_breakdown,
    "LAYOUT-032": render_v31_scope_definition,
    "LAYOUT-033": render_v31_section_divider,
}


def _register_layout_range(start: int, end: int, renderer: Callable[[Presentation, PowerPointSlide, PowerPointData, PptxContext], None]) -> None:
    for layout_index in range(start, end + 1):
        LAYOUT_RENDERER_REGISTRY[f"LAYOUT-{layout_index:03d}"] = renderer


_register_layout_range(34, 39, render_v4_executive_brief)
_register_layout_range(40, 43, render_v4_issue_tree)
_register_layout_range(44, 46, render_v4_before_after)
_register_layout_range(47, 50, render_v4_value_pyramid)
_register_layout_range(51, 53, render_v4_competitive_positioning)
_register_layout_range(54, 56, render_v4_kpi_dashboard)
_register_layout_range(57, 59, render_v4_roi_dashboard)
_register_layout_range(60, 64, render_v4_swimlane_roadmap)
_register_layout_range(65, 69, render_v31_flow)
_register_layout_range(70, 76, render_v4_architecture_stack)
_register_layout_range(77, 79, render_v31_governance_map)
_register_layout_range(80, 82, render_v4_risk_radar)
_register_layout_range(83, 85, render_v4_investment_mix)
_register_layout_range(86, 87, render_v31_scope_definition)
_register_layout_range(88, 90, render_v4_next_meeting)
_register_layout_range(91, 93, render_v31_quote)
_register_layout_range(94, 99, render_v31_checklist)
_register_layout_range(100, 106, render_v31_matrix)
_register_layout_range(107, 112, render_v4_before_after)
_register_layout_range(113, 122, render_v4_swimlane_roadmap)


def add_designed_slide(
    prs: Presentation,
    slide_data: PowerPointSlide,
    data: PowerPointData,
    index: int,
    context: PptxContext,
) -> None:
    render_v5_masterpiece_slide(prs, slide_data, data, context, index)
    return

    layout_id = layout_id_from_layout_key(slide_data.layout)
    renderer = LAYOUT_RENDERER_REGISTRY.get(layout_id or "")
    if renderer is not None:
        renderer(prs, slide_data, data, context)
        return

    kind = resolve_slide_kind(slide_data, index)
    if kind == "cover":
        add_cover_slide(prs, slide_data, data, context)
    elif kind == "proposal_summary":
        add_proposal_summary_slide(prs, slide_data, context)
    elif kind == "proposal_concept":
        add_concept_slide(prs, slide_data, context)
    elif kind == "current_understanding":
        add_current_understanding_slide(prs, slide_data, context)
    elif kind == "competitor":
        add_competitor_slide(prs, slide_data, context)
    elif kind == "target_user":
        add_target_user_slide(prs, slide_data, context)
    elif kind == "customer_journey":
        add_customer_journey_slide(prs, slide_data, context)
    elif kind == "web_strategy":
        add_web_strategy_slide(prs, slide_data, context)
    elif kind == "sitemap":
        add_sitemap_slide(prs, slide_data, context)
    elif kind == "content":
        add_content_design_slide(prs, slide_data, context)
    elif kind == "kpi":
        add_kpi_slide(prs, slide_data, context)
    elif kind == "understanding":
        add_understanding_slide(prs, slide_data)
    elif kind == "issues":
        add_issues_slide(prs, slide_data, context)
    elif kind == "solution":
        add_solution_slide(prs, slide_data, context)
    elif kind == "process":
        add_process_slide(prs, slide_data, context)
    elif kind == "schedule":
        add_schedule_slide(prs, slide_data, context)
    elif kind == "case_studies":
        add_case_studies_slide(prs, slide_data, context)
    elif kind == "team":
        add_team_slide(prs, slide_data, context)
    elif kind == "estimate":
        add_estimate_slide(prs, slide_data, context)
    elif kind == "budget_fit":
        add_budget_fit_slide(prs, slide_data, context)
    elif kind == "estimate_priority":
        add_estimate_priority_slide(prs, slide_data, context)
    elif kind == "cost":
        add_cost_slide(prs, slide_data, context)
    elif kind == "win_probability":
        if context.win_probability is not None:
            add_win_probability_slide(prs, context.win_probability, slide_data.slide_no)
        else:
            add_generic_slide(prs, slide_data)
    elif kind == "next_steps":
        add_next_steps_slide(prs, slide_data, context)
    elif kind == "summary":
        add_summary_slide(prs, slide_data, context)
    elif kind == "quality_comparison":
        add_quality_comparison_slide(prs, slide_data, context)
    elif kind == "quality_timeline":
        add_quality_timeline_slide(prs, slide_data, context)
    elif kind == "quality_roadmap":
        add_quality_timeline_slide(prs, slide_data, context, section="ROADMAP")
    elif kind == "quality_kpi":
        add_quality_kpi_slide(prs, slide_data, context)
    elif kind == "quality_flow":
        add_quality_flow_slide(prs, slide_data, context)
    elif kind == "quality_matrix":
        add_quality_matrix_slide(prs, slide_data, context)
    else:
        add_generic_slide(prs, slide_data)


def resolve_slide_kind(slide_data: PowerPointSlide, index: int) -> str:
    title = slide_data.title
    if slide_data.layout in {
        "quality_comparison",
        "quality_timeline",
        "quality_roadmap",
        "quality_kpi",
        "quality_flow",
        "quality_matrix",
    }:
        return slide_data.layout
    if index == 0 or slide_data.layout == "title":
        return "cover"
    if "提案サマリー" in title:
        return "proposal_summary"
    if "提案コンセプト" in title:
        return "proposal_concept"
    if "現状理解" in title:
        return "current_understanding"
    if "市場" in title or "競合" in title:
        return "competitor"
    if "ターゲット" in title or "ユーザー分析" in title:
        return "target_user"
    if "カスタマージャーニー" in title or "ユーザー行動" in title:
        return "customer_journey"
    if "Web戦略" in title or "WEB戦略" in title or "導入戦略" in title:
        return "web_strategy"
    if "サイトマップ" in title or "サイト構成" in title or "導入構成" in title:
        return "sitemap"
    if "コンテンツ" in title or "施策設計" in title:
        return "content"
    if "KPI" in title or "指標" in title:
        return "kpi"
    if "貴社理解" in title or "企業理解" in title:
        return "understanding"
    if "課題" in title and "解決" not in title:
        return "issues"
    if "解決" in title:
        return "solution"
    if "制作方針" in title or "方針" in title:
        return "process"
    if "スケジュール" in title:
        return "schedule"
    if "実績" in title:
        return "case_studies"
    if "体制" in title:
        return "team"
    if "予算適合" in title:
        return "budget_fit"
    if "必須" in title and "推奨" in title and "オプション" in title:
        return "estimate_priority"
    if "見積" in title:
        return "estimate"
    if "費用" in title or "概算" in title:
        return "cost"
    if "受注確率" in title or "案件ランク" in title:
        return "win_probability"
    if "今後" in title or "進め方" in title or "まとめ" in title:
        return "next_steps"

    order_map = {
        1: "proposal_summary",
        2: "current_understanding",
        3: "issues",
        4: "competitor",
        5: "target_user",
        6: "customer_journey",
        7: "web_strategy",
        8: "sitemap",
        9: "content",
        10: "kpi",
        11: "process",
        12: "schedule",
        13: "team",
        14: "cost",
        15: "next_steps",
    }
    return order_map.get(index, "generic")


def add_cover_slide(prs: Presentation, slide_data: PowerPointSlide, data: PowerPointData, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    theme = resolve_template_colors(context.design_template)

    add_shape(slide, MSO_SHAPE.RECTANGLE, 0.02, 0.02, SLIDE_WIDTH - 0.04, SLIDE_HEIGHT - 0.04, fill=theme["background"], line=theme["background"])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0.02, 0.02, 0.14, SLIDE_HEIGHT - 0.04, fill=theme["accent"], line=theme["accent"])
    add_shape(slide, MSO_SHAPE.RIGHT_TRIANGLE, 8.55, 0.02, 4.72, 7.46, fill=theme["secondary"], line=theme["secondary"])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 9.15, 0.02, 4.12, SLIDE_HEIGHT - 0.04, fill=theme["surface"], line=theme["surface"])
    add_shape(slide, MSO_SHAPE.OVAL, 9.65, 0.92, 3.05, 3.05, fill=COLORS["teal_light"], line=COLORS["teal_light"])
    add_shape(slide, MSO_SHAPE.OVAL, 10.82, 3.55, 1.72, 1.72, fill=COLORS["blue_light"], line=COLORS["blue_light"])
    cover_badges = icon_labels_for_category(context.proposal_category)
    add_icon_badge(slide, cover_badges[0], 9.82, 1.68, theme["accent"])
    add_icon_badge(slide, cover_badges[1], 11.05, 2.78, theme["primary"])
    add_icon_badge(slide, cover_badges[2], 9.9, 4.75, theme["support"])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 9.75, 5.92, 2.7, 0.13, fill=theme["accent"], line=theme["accent"])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 9.75, 6.2, 1.9, 0.13, fill=theme["support"], line=theme["support"])

    section_label = "Web制作提案書" if context.proposal_category == "web" else f"{context.proposal_label} 提案書"
    add_section_label(slide, normalize_customer_facing_text(section_label, limit=18), 0.88, 0.8, fill=theme["accent"], color=theme["text_on_dark"])
    add_title(slide, slide_data.title or data.deck_title, 0.88, 1.55, 7.35, 1.24, size=42, color=theme["text_on_dark"])
    add_text(slide, f"{context.client_name} 御中", 0.92, 3.02, 6.9, 0.36, size=18, color=COLORS["teal_light"], bold=True)
    add_text(
        slide,
        "成果につながるWebサイト制作・改善のご提案" if context.proposal_category == "web" else f"{context.concept}のご提案",
        0.92,
        3.68,
        7.25,
        0.5,
        size=20,
        color=theme["text_on_dark"],
        bold=True,
    )
    add_text(slide, f"提案日 {date.today().strftime('%Y.%m.%d')}", 0.92, 5.95, 3.5, 0.26, size=12, color=COLORS["teal_light"])
    add_text(slide, "ProposalPilot / AI営業秘書", 0.92, 6.58, 4.0, 0.24, size=11, color=COLORS["teal_light"])
    add_text(slide, f"{slide_data.slide_no:02}", 11.58, 6.86, 0.76, 0.22, size=10, color=COLORS["muted"], align=PP_ALIGN.RIGHT)

    add_text(slide, "戦略", 9.45, 1.0, 2.6, 0.32, size=16, color=theme["text_on_light"], bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "設計", 10.28, 3.2, 1.9, 0.3, size=14, color=theme["text_on_light"], bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "成果", 9.95, 5.18, 2.2, 0.3, size=14, color=theme["text_on_light"], bold=True, align=PP_ALIGN.CENTER)


def add_chapter_title_slide(prs: Presentation, slide_data: PowerPointSlide, chapter_no: int, display_slide_no: int, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    accent = SECTION_COLORS[(chapter_no - 1) % len(SECTION_COLORS)]
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT, fill=COLORS["navy"], line=COLORS["navy"])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0.0, 0, 0.18, SLIDE_HEIGHT, fill=accent, line=accent)
    add_shape(slide, MSO_SHAPE.RIGHT_TRIANGLE, 8.4, 0, 4.94, 7.5, fill=COLORS["navy_2"], line=COLORS["navy_2"])
    add_text(slide, f"Chapter {chapter_no:02}", 0.92, 1.65, 2.4, 0.28, size=14, color=COLORS["teal_light"], bold=True)
    add_text(slide, slide_data.title, 0.92, 2.14, 7.9, 0.72, size=34, color=COLORS["white"], bold=True)
    add_text(slide, _trim(chapter_message(slide_data.title, context), 86), 0.95, 3.22, 7.55, 0.56, size=17, color=COLORS["teal_light"], bold=True)
    add_icon_badge(slide, chapter_icon(slide_data.title), 10.18, 2.38, accent, size=1.3)
    add_text(slide, context.concept, 9.42, 4.1, 2.85, 0.32, size=17, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, display_slide_no)


def add_proposal_summary_slide(prs: Presentation, slide_data: PowerPointSlide, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_header(slide, slide_data.title or "提案サマリー", "SUMMARY", accent=COLORS["blue"])
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.92, 1.62, 11.48, 1.26, fill=COLORS["navy"], line=COLORS["navy"])
    add_text(slide, "提案コンセプト", 1.25, 1.88, 1.7, 0.24, size=12, color=COLORS["teal_light"], bold=True)
    add_text(slide, context.concept, 3.0, 1.8, 3.2, 0.36, size=24, color=COLORS["white"], bold=True)
    add_text(slide, _trim(concept_statement(context.concept), 62), 6.35, 1.78, 5.25, 0.48, size=15, color=COLORS["white"], bold=True)

    summary_items = ensure_items(
        slide_data.bullets + context.web_strategy_items,
        ["現状課題を実行方針へ落とし込む", "導入範囲と施策を再構成", "KPIを定義して導入後改善につなげる"],
        3,
    )
    titles = ["解決する課題", "主要施策", "期待成果"]
    for idx, item in enumerate(summary_items[:3]):
        add_card(slide, titles[idx], item, 0.95 + idx * 4.02, 3.36, 3.35, 1.68, SECTION_COLORS[idx], COLORS["white"], number=str(idx + 1))
    conclusion = "顧客理解、競合比較、KPI設計を起点に、公開後も成果を追えるWebサイトへ改善します。"
    if context.proposal_category != "web":
        conclusion = "顧客理解、比較軸、KPI設計を起点に、導入後も成果を追える実行計画へ整理します。"
    add_insight_band(slide, "提案の結論", conclusion, 0.92, 5.82, 11.4, 0.54)
    add_footer(slide, slide_data.slide_no)


def add_current_understanding_slide(prs: Presentation, slide_data: PowerPointSlide, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_header(slide, slide_data.title or "現状理解と事業課題", "顧客理解")
    understanding = merge_understanding_items(context.current_understanding, slide_data.bullets)
    items = [
        ("現状", understanding["現状"], COLORS["teal"]),
        ("課題", understanding["課題"], COLORS["red"]),
        ("機会", understanding["機会"], COLORS["blue"]),
        ("目指す状態", understanding["目指す状態"], COLORS["green"]),
    ]
    positions = [(0.9, 1.72), (6.78, 1.72), (0.9, 4.08), (6.78, 4.08)]

    for idx, ((title, body, accent), (x, y)) in enumerate(zip(items, positions), start=1):
        add_card(slide, title, body, x, y, 5.25, 1.64, accent, COLORS["white"], number=str(idx))
    add_footer(slide, slide_data.slide_no)


def add_competitor_slide(prs: Presentation, slide_data: PowerPointSlide, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_header(slide, slide_data.title or "市場・競合分析", "COMPETITOR", accent=COLORS["purple"])
    competitor_name = context.competitor_company_name or "競合サイト"
    target_label = "競合あり" if has_competitor_context(context) else "競合情報を整理"
    target_text = f"比較対象: {competitor_name}"
    if context.competitor_site_url:
        target_text = f"{target_text} / {context.competitor_site_url}"
    add_text(slide, target_label, 0.9, 1.16, 1.25, 0.28, size=11, color=COLORS["purple"], bold=True)
    add_text(slide, _trim(target_text, 70), 2.05, 1.13, 7.6, 0.32, size=12, color=COLORS["muted"], bold=True)
    add_text(slide, "5段階評価", 10.2, 1.18, 1.8, 0.22, size=12, color=COLORS["muted"], bold=True, align=PP_ALIGN.RIGHT)
    add_table(
        slide,
        headers=["比較項目", "現状仮説", competitor_name[:12], "提案後"],
        rows=context.competitor_rows,
        x=0.85,
        y=1.64,
        w=11.65,
        h=3.55,
        col_widths=[2.35, 2.8, 2.65, 3.85],
    )
    add_insight_band(
        slide,
        "競合に対する勝ち筋",
        f"{context.winning_strategy}を軸に、競合が強い領域を踏まえて重点改善領域を明確化します。",
        0.92,
        5.48,
        11.4,
        0.72,
    )
    add_footer(slide, slide_data.slide_no)


def add_target_user_slide(prs: Presentation, slide_data: PowerPointSlide, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_header(slide, slide_data.title or "ターゲットユーザー分析", "TARGET", accent=COLORS["teal"])
    rows = context.target_user_rows[:4]
    for idx, (label, body) in enumerate(rows):
        x = 0.92 + (idx % 2) * 5.88
        y = 1.72 + (idx // 2) * 2.05
        add_card(slide, label, body, x, y, 5.24, 1.48, SECTION_COLORS[idx], COLORS["white"], number=str(idx + 1))
    target_note = "ターゲットの不安をFAQ・実績・サービス詳細で解消し、問い合わせ前の比較検討を支援します。"
    if context.proposal_category != "web":
        target_note = "利用者、意思決定者、運用担当の不安を整理し、導入判断と定着を支援します。"
    add_insight_band(slide, "設計方針", target_note, 0.92, 5.92, 11.4, 0.5)
    add_footer(slide, slide_data.slide_no)


def add_web_strategy_slide(prs: Presentation, slide_data: PowerPointSlide, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_header(slide, slide_data.title or ("Web戦略" if context.proposal_category == "web" else "導入戦略"), "STRATEGY", accent=COLORS["blue"])
    items = ensure_items(context.web_strategy_items + slide_data.bullets, ["現状整理", "導入設計", "小規模検証", "運用改善"], 4)
    labels = ["集客", "比較検討", "問い合わせ", "運用改善"] if context.proposal_category == "web" else ["現状", "設計", "検証", "運用"]
    badges = ["SEO", "INFO", "CV", "PDCA"] if context.proposal_category == "web" else ["ASIS", "PLAN", "TEST", "OPS"]
    for idx, item in enumerate(items[:4]):
        x = 0.86 + idx * 3.05
        accent = SECTION_COLORS[idx]
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.82, 2.58, 3.0, fill=COLORS["white"], line=COLORS["line"])
        add_icon_badge(slide, badges[idx], x + 0.75, 2.12, accent, size=0.72)
        add_text(slide, labels[idx], x + 0.26, 3.14, 2.04, 0.28, size=16, color=accent, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, _trim(item, 40), x + 0.28, 3.72, 2.02, 0.54, size=13, color=COLORS["text"], bold=True, align=PP_ALIGN.CENTER)
        if idx < 3:
            add_shape(slide, MSO_SHAPE.CHEVRON, x + 2.5, 3.0, 0.36, 0.5, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    strategy_note = f"{context.concept}を軸に、入口設計からCV改善、公開後運用まで一貫して設計します。"
    if context.proposal_category != "web":
        strategy_note = f"{context.concept}を軸に、現状整理、導入設計、検証、運用改善まで一貫して設計します。"
    add_insight_band(slide, "戦略の要点", strategy_note, 0.92, 5.62, 11.4, 0.58)
    add_footer(slide, slide_data.slide_no)


def add_content_design_slide(prs: Presentation, slide_data: PowerPointSlide, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_header(slide, slide_data.title or "コンテンツ設計", "CONTENT", accent=COLORS["green"])
    fallback_content = ["サービス詳細", "実績・事例", "FAQ", "お問い合わせ"] if context.proposal_category == "web" else ["課題整理", "導入範囲", "効果測定", "運用設計"]
    items = ensure_items(context.content_items + slide_data.bullets, fallback_content, 4)
    for idx, item in enumerate(items[:4]):
        x = 0.92 + (idx % 2) * 5.88
        y = 1.68 + (idx // 2) * 2.0
        add_card(slide, content_title(item, idx), item, x, y, 5.24, 1.45, SECTION_COLORS[idx], COLORS["white"], number=str(idx + 1))
    content_note = "認知、比較検討、問い合わせ前の不安解消まで、各コンテンツに明確な役割を持たせます。"
    if context.proposal_category != "web":
        content_note = "各施策に役割を持たせ、導入判断、実行、効果測定、運用定着を支援します。"
    add_insight_band(slide, "施策の役割", content_note, 0.92, 5.88, 11.4, 0.54)
    add_footer(slide, slide_data.slide_no)


def add_concept_slide(prs: Presentation, slide_data: PowerPointSlide, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_header(slide, slide_data.title or "本提案の方向性", "提案方針", accent=COLORS["blue"])

    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.92, 1.72, 11.48, 1.88, fill=COLORS["navy"], line=COLORS["navy"])
    add_text(slide, "提案コンセプト", 1.28, 2.0, 2.6, 0.28, size=13, color=COLORS["teal_light"], bold=True)
    add_text(slide, context.concept, 1.28, 2.36, 5.6, 0.58, size=32, color=COLORS["white"], bold=True)
    add_text(slide, _trim(concept_statement(context.concept), 62), 7.0, 2.06, 4.65, 0.78, size=16, color=COLORS["white"], bold=True)

    focus_items = ensure_items(
        context.solution_points + slide_data.bullets,
        ["情報設計と導線を再設計", "成果指標に沿って改善", "公開後の運用まで見据える"],
        3,
    )
    for idx, item in enumerate(focus_items[:3]):
        x = 0.95 + idx * 4.0
        add_card(slide, ["重点施策", "実行方針", "成果設計"][idx], item, x, 4.26, 3.35, 1.34, SECTION_COLORS[idx], COLORS["white"], number=str(idx + 1))
    add_footer(slide, slide_data.slide_no)


def add_customer_journey_slide(prs: Presentation, slide_data: PowerPointSlide, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_header(slide, slide_data.title or "カスタマージャーニー", "ユーザー行動", accent=COLORS["purple"])
    stages = context.journey_points[:3]
    stage_width = 3.48
    stage_y = 2.0

    for idx, (stage, description) in enumerate(stages):
        x = 0.9 + idx * 4.04
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, stage_y, stage_width, 2.25, fill=COLORS["white"], line=COLORS["line"])
        add_shape(slide, MSO_SHAPE.OVAL, x + 0.26, stage_y + 0.28, 0.52, 0.52, fill=SECTION_COLORS[idx], line=SECTION_COLORS[idx])
        add_text(slide, str(idx + 1), x + 0.26, stage_y + 0.4, 0.52, 0.16, size=11, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, stage, x + 0.92, stage_y + 0.34, 2.2, 0.3, size=20, color=COLORS["navy"], bold=True)
        add_text(slide, _trim(description, 48), x + 0.34, stage_y + 1.0, stage_width - 0.68, 0.62, size=14, color=COLORS["text"], bold=True)
        add_text(slide, _trim(journey_action(stage, context.concept), 34), x + 0.34, stage_y + 1.78, stage_width - 0.68, 0.32, size=12, color=COLORS["muted"])
        if idx < len(stages) - 1:
            add_shape(slide, MSO_SHAPE.CHEVRON, x + stage_width + 0.18, stage_y + 0.88, 0.38, 0.46, fill=COLORS["line_dark"], line=COLORS["line_dark"])

    journey_note = "認知から問い合わせまでの離脱点を減らし、情報設計・CTA・実績訴求を一連の導線として改善します。"
    if context.proposal_category != "web":
        journey_note = "導入前後の流れを整理し、例外処理、連携、運用改善まで一連の業務として設計します。"
    add_insight_band(slide, "提案との関係", journey_note, 0.92, 5.58, 11.4, 0.58)
    add_footer(slide, slide_data.slide_no)


def add_sitemap_slide(prs: Presentation, slide_data: PowerPointSlide, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_header(slide, slide_data.title or ("推奨サイト構成" if context.proposal_category == "web" else "導入構成"), "サイトマップ" if context.proposal_category == "web" else "構成案", accent=COLORS["green"])
    fallback_structure = ["TOP", "会社案内", "サービス", "実績", "お知らせ", "FAQ", "お問い合わせ"] if context.proposal_category == "web" else ["対象業務", "課題", "導入範囲", "連携先", "運用体制", "効果測定"]
    items = ensure_items(context.sitemap_items, fallback_structure, 8)
    if context.proposal_category != "web":
        add_architecture_diagram(slide, architecture_nodes_for_category(context.proposal_category), 0.9, 2.04, 11.35, 1.86)
        add_insight_band(
            slide,
            "構成の考え方",
            "入力、AI判定、人による確認、既存業務への連携、運用改善までを一連の流れとして設計します。",
            0.92,
            5.55,
            11.4,
            0.62,
        )
        add_footer(slide, slide_data.slide_no)
        return
    top_label = items[0] if items else "TOP"
    children = [item for item in items[1:8] if item != top_label]

    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 5.12, 1.62, 3.1, 0.68, fill=COLORS["navy"], line=COLORS["navy"])
    add_text(slide, top_label, 5.36, 1.82, 2.62, 0.22, size=18, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 6.62, 2.3, 0.04, 0.46, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 1.6, 2.76, 10.08, 0.04, fill=COLORS["line_dark"], line=COLORS["line_dark"])

    for idx, item in enumerate(children[:7]):
        row = idx // 4
        col = idx % 4
        x = 0.92 + col * 3.08 + (0.78 if row else 0)
        y = 3.1 + row * 1.24
        accent = SECTION_COLORS[idx % len(SECTION_COLORS)]
        add_shape(slide, MSO_SHAPE.RECTANGLE, x + 1.24, 2.78, 0.04, y - 2.78, fill=COLORS["line_dark"], line=COLORS["line_dark"])
        add_card(slide, item, sitemap_note(item), x, y, 2.54, 0.88, accent, COLORS["white"])
    structure_note = "サービス・実績・FAQで比較検討を支え、CMS更新領域から最新情報を継続発信します。"
    if context.proposal_category != "web":
        structure_note = "対象業務、連携先、運用体制、効果測定を分け、導入後の迷いを減らします。"
    add_insight_band(slide, "構成意図", structure_note, 0.92, 5.92, 11.4, 0.5)
    add_footer(slide, slide_data.slide_no)


def add_kpi_slide(prs: Presentation, slide_data: PowerPointSlide, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_header(slide, slide_data.title or "KPI設計", "成果設計", accent=COLORS["orange"])
    if context.proposal_category == "web":
        metrics = [
            ("問い合わせ数", f"月{int(context.kpi_targets['inquiries'])}件", context.kpi_targets["inquiries"] / 30, COLORS["teal"]),
            ("CV率", f"{context.kpi_targets['cv_rate']}%", context.kpi_targets["cv_rate"] / 5, COLORS["blue"]),
            ("自然検索流入", f"月{int(context.kpi_targets['organic'])}セッション", context.kpi_targets["organic"] / 6000, COLORS["green"]),
            ("資料DL数", f"月{int(context.kpi_targets['downloads'])}件", context.kpi_targets["downloads"] / 40, COLORS["orange"]),
        ]
    else:
        metric_lines = ensure_items(
            [f"{label}: {value}" for label, value in context.kpi_rows],
            ["業務時間削減: 測定方法を合意", "処理品質向上: 評価基準を設定", "運用定着率: 目標値を設定", "改善サイクル: 運用後に測定"],
            4,
        )
        metrics = [
            (label, value, min(1.0, 0.48 + idx * 0.12), SECTION_COLORS[idx])
            for idx, metric in enumerate(metric_lines)
            for label, value in [_split_metric_text(metric)]
        ]
    for idx, (label, value, ratio, accent) in enumerate(metrics):
        x = 0.9 + idx * 3.02
        add_metric_card(slide, label, value, x, 1.62, 2.55, 1.28, accent)
        y = 3.28 + idx * 0.34
        add_shape(slide, MSO_SHAPE.RECTANGLE, 1.08, y + 0.08, 10.9, 0.08, fill=COLORS["line"], line=COLORS["line"])
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 1.08, y, max(0.7, min(10.9, 10.9 * ratio)), 0.24, fill=accent, line=accent)

    add_table(
        slide,
        headers=["設計観点", "見る指標", "改善アクション"],
        rows=[
            ["入口", kpi_metric_for(context.concept, "集客"), "対象範囲と入力条件を整理"],
            ["行動", kpi_metric_for(context.concept, "行動"), "施策と運用フローを強化"],
            ["成果", kpi_metric_for(context.concept, "成果"), "KPIと改善サイクルを設計"],
        ],
        x=0.92,
        y=5.12,
        w=11.4,
        h=1.05,
        col_widths=[2.0, 4.0, 5.4],
    )
    add_footer(slide, slide_data.slide_no)


def add_understanding_slide(prs: Presentation, slide_data: PowerPointSlide) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_header(slide, slide_data.title, "企業理解")
    points = ensure_items(slide_data.bullets, ["事業内容とターゲットを整理", "提案に期待される役割を仮説化", "確認すべき前提条件を明確化"], 3)

    for idx, item in enumerate(points[:3]):
        x = MARGIN_X + idx * 4.1
        add_card(
            slide,
            title=f"理解ポイント {idx + 1}",
            body=item,
            x=x,
            y=2.0,
            w=3.55,
            h=2.6,
            accent=SECTION_COLORS[idx],
            fill=COLORS["white"],
            number=str(idx + 1),
        )
    add_insight_band(slide, "提案の前提", slide_data.visual_suggestion or "事業・ターゲット・提案範囲を整理した図", 0.9, 5.25, 11.5, 0.72)
    add_footer(slide, slide_data.slide_no)


def add_issues_slide(prs: Presentation, slide_data: PowerPointSlide, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_header(slide, slide_data.title, "課題整理")
    issues = ensure_items(context.project_points + slide_data.bullets, ["目的とKPIの明確化", "ターゲット別導線の整理", "公開後の改善運用の検討"], 4)

    for idx, issue in enumerate(issues[:4]):
        y = 1.64 + idx * 1.08
        add_card(
            slide,
            title=f"Priority {idx + 1}",
            body=issue,
            x=0.9,
            y=y,
            w=7.6,
            h=0.82,
            accent=SECTION_COLORS[idx % len(SECTION_COLORS)],
            fill=COLORS["white"],
            number=str(idx + 1),
        )

    add_side_panel(slide, "整理観点", ["影響度", "緊急度", "提案適合度"], 9.05, 1.64, 3.25, 4.35, COLORS["navy"])
    add_footer(slide, slide_data.slide_no)


def add_solution_slide(prs: Presentation, slide_data: PowerPointSlide, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_header(slide, slide_data.title, "解決策")
    items = ensure_items(context.solution_points + slide_data.bullets, ["要件整理の再整理", "導入範囲と施策方針の策定", "効果測定につながる運用設計"], 4)
    rows = build_solution_rows(items)
    add_table(
        slide,
        headers=["区分", "想定課題", "提案する解決策"],
        rows=rows,
        x=0.85,
        y=1.68,
        w=11.65,
        h=4.25,
        col_widths=[1.35, 4.4, 5.9],
    )
    add_footer(slide, slide_data.slide_no)


def add_process_slide(prs: Presentation, slide_data: PowerPointSlide, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_header(slide, slide_data.title, "制作方針" if context.proposal_category == "web" else "実行方針")
    steps = ensure_items(context.service_points + slide_data.bullets, ["初期設計を重視した実行プロセス", "確認しやすい要件整理", "導入後の改善を見据えた設計"], 4)
    add_step_flow(slide, steps[:4], 0.82, 2.12, 11.7, 1.85)
    add_insight_band(slide, "品質担保の考え方", "要件整理・設計・制作・検証を段階的に進め、認識齟齬を抑えます。", 0.92, 5.18, 11.4, 0.8)
    add_footer(slide, slide_data.slide_no)


def add_schedule_slide(prs: Presentation, slide_data: PowerPointSlide, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_header(slide, slide_data.title, "スケジュール")
    phases = ensure_items(context.schedule_points + slide_data.bullets, ["要件整理", "設計・計画", "実装・検証", "運用開始・改善"], 4)
    week_labels = ["1-2週", "3-4週", "5-7週", "8週"]

    add_text(slide, "工程", 0.92, 1.48, 1.5, 0.25, size=13, color=COLORS["muted"], bold=True)
    for idx, label in enumerate(week_labels):
        add_text(slide, label, 3.15 + idx * 2.15, 1.48, 1.25, 0.25, size=13, color=COLORS["muted"], bold=True, align=PP_ALIGN.CENTER)

    for idx, phase in enumerate(phases[:4]):
        y = 2.0 + idx * 0.88
        add_text(slide, _trim(phase, 18), 0.92, y + 0.12, 2.0, 0.25, size=14, color=COLORS["text"], bold=True)
        add_shape(slide, MSO_SHAPE.RECTANGLE, 3.02, y + 0.22, 8.85, 0.08, fill=COLORS["line"], line=COLORS["line"])
        add_shape(
            slide,
            MSO_SHAPE.ROUNDED_RECTANGLE,
            3.05 + idx * 1.55,
            y,
            2.25,
            0.46,
            fill=SECTION_COLORS[idx % len(SECTION_COLORS)],
            line=SECTION_COLORS[idx % len(SECTION_COLORS)],
        )
    add_insight_band(slide, "進行イメージ", "詳細スケジュールは要件・素材準備・確認体制により調整します。", 0.92, 5.78, 11.4, 0.55)
    add_footer(slide, slide_data.slide_no)


def add_case_studies_slide(prs: Presentation, slide_data: PowerPointSlide, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_header(slide, slide_data.title, "実績紹介")
    case_triplets = context.case_triplets or build_case_triplets_from_items(context.case_studies + slide_data.bullets)
    if not case_triplets:
        case_triplets = [
            {
                "title": "関連実績",
                "current": "近しい課題の成功事例を提案時に差し替え",
                "action": "情報設計、導線改善、運用支援の実績を整理",
                "result": "成果につながった進め方を提案へ反映",
            }
        ]

    for idx, case in enumerate(case_triplets[:3]):
        x = 0.9 + idx * 4.05
        accent = SECTION_COLORS[idx]
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.72, 3.48, 4.15, fill=COLORS["white"], line=COLORS["line"])
        add_shape(slide, MSO_SHAPE.RECTANGLE, x, 1.72, 3.48, 0.08, fill=accent, line=accent)
        add_text(slide, _trim(display_case_title(case["title"], idx + 1), 24), x + 0.26, 1.98, 2.94, 0.28, size=14, color=accent, bold=True)
        add_case_line(slide, "現状", case["current"], x + 0.26, 2.62, accent)
        add_case_line(slide, "施策", case["action"], x + 0.26, 3.62, accent)
        add_case_line(slide, "成果", case["result"], x + 0.26, 4.62, accent)
    add_footer(slide, slide_data.slide_no)


def add_team_slide(prs: Presentation, slide_data: PowerPointSlide, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_header(slide, slide_data.title, "体制紹介")
    roles = ["PM/ディレクター", "設計担当", "実装担当", "運用・改善支援"] if context.proposal_category != "web" else ["PM/ディレクター", "デザイナー", "エンジニア", "運用・改善支援"]
    details = ensure_items(context.team_points + slide_data.bullets, ["進行管理", "要件・構成設計", "実装・検証", "導入後の改善相談"], 4)

    center_x, center_y = 6.1, 3.55
    add_shape(slide, MSO_SHAPE.OVAL, center_x - 0.82, center_y - 0.62, 1.65, 1.25, fill=COLORS["navy"], line=COLORS["navy"])
    add_text(slide, "Project\nCore", center_x - 0.55, center_y - 0.32, 1.1, 0.55, size=14, color=COLORS["white"], bold=True, align=PP_ALIGN.CENTER)

    positions = [(1.0, 1.85), (8.85, 1.85), (1.0, 4.65), (8.85, 4.65)]
    for idx, (x, y) in enumerate(positions):
        add_card(slide, roles[idx], details[idx], x, y, 3.25, 1.35, SECTION_COLORS[idx], COLORS["white"])
        add_shape(slide, MSO_SHAPE.RECTANGLE, 4.35 if idx % 2 == 0 else 8.15, y + 0.62, 1.05, 0.04, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    add_footer(slide, slide_data.slide_no)


def add_cost_slide(prs: Presentation, slide_data: PowerPointSlide, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_header(slide, slide_data.title, "費用概算")
    items = ensure_items(
        [
            f"合計概算: {context.estimate.total_label}",
            f"予算適合: {context.estimate.budget_fit}",
            f"必須対応: {'、'.join(context.estimate.required[:3])}",
            *context.cost_points,
            *slide_data.bullets,
        ],
        ["必須範囲とオプションを分離", "規模と機能要件に応じて調整", "詳細見積はヒアリング後に提示"],
        4,
    )
    scope_label = context.estimate.scope_label
    required_detail = "要件整理・設計・制作・検証" if context.proposal_category == "web" else "要件整理・設計・実装・検証"
    adjustable_detail = "CMS・SEO・特殊機能・撮影原稿" if context.proposal_category == "web" else "連携・学習・運用支援・追加検証"
    rows = [
        ["合計概算", context.estimate.total_label, scope_label],
        ["予算適合", context.estimate.budget_fit, f"予算感: {context.estimate.budget_label}"],
        ["必須対応", _trim("、".join(context.estimate.required[:4]), 36), required_detail],
        ["調整範囲", _trim(items[3] if len(items) > 3 else "推奨・オプションを段階化", 36), adjustable_detail],
    ]
    add_table(
        slide,
        headers=["区分", "考え方", "主な内容"],
        rows=rows,
        x=0.85,
        y=1.55,
        w=11.65,
        h=4.0,
        col_widths=[2.0, 4.3, 5.35],
    )
    add_insight_band(slide, "見積方針", "必須範囲を先に確保し、推奨・オプション対応を予算と納期に合わせて段階提案します。", 0.92, 5.84, 11.4, 0.55)
    add_footer(slide, slide_data.slide_no)


def add_estimate_slide(prs: Presentation, slide_data: PowerPointSlide, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_header(slide, slide_data.title or "概算見積", "ESTIMATE", accent=COLORS["teal"])
    estimate_scope = context.estimate.scope_label
    add_text(slide, f"合計概算 {context.estimate.total_label}", 3.35, 0.98, 3.15, 0.28, size=16, color=COLORS["teal"], bold=True, align=PP_ALIGN.RIGHT)
    add_text(slide, f"{estimate_scope} / 予算感: {context.estimate.budget_label}", 6.72, 0.98, 5.5, 0.28, size=12, color=COLORS["muted"], bold=True, align=PP_ALIGN.RIGHT)
    rows = [
        [
            str(line["name"]),
            f"{line['min']}万〜{line['max']}万円" if bool(line["enabled"]) else "対象外",
            str(line["priority"]),
        ]
        for line in context.estimate.lines
    ]
    add_table(
        slide,
        headers=["見積項目", "金額レンジ", "分類"],
        rows=rows,
        x=0.76,
        y=1.62,
        w=11.8,
        h=4.75,
        col_widths=[5.1, 3.2, 3.5],
    )
    add_footer(slide, slide_data.slide_no)


def add_budget_fit_slide(prs: Presentation, slide_data: PowerPointSlide, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_header(slide, slide_data.title or "予算適合判定", "BUDGET FIT", accent=COLORS["orange"])
    fit_color = COLORS["green"] if context.estimate.budget_fit == "予算内" else COLORS["orange"] if context.estimate.budget_fit == "やや調整必要" else COLORS["red"]
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.95, 1.62, 4.0, 3.2, fill=COLORS["white"], line=COLORS["line"])
    add_text(slide, "判定", 1.25, 2.05, 3.4, 0.3, size=16, color=COLORS["muted"], bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, context.estimate.budget_fit, 1.2, 2.55, 3.5, 0.55, size=25, color=fit_color, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, f"予算感: {context.estimate.budget_label}", 1.25, 3.46, 3.4, 0.28, size=14, color=COLORS["text"], bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, f"概算見積: {context.estimate.total_label}", 1.25, 3.9, 3.4, 0.28, size=14, color=COLORS["text"], bold=True, align=PP_ALIGN.CENTER)

    cards = [
        ("差分の見方", "上限予算と概算上限を比較し、調整要否を判断"),
        ("調整方針", "必須対応を優先し、推奨・オプションを段階化"),
        ("次回確認", "規模、導入範囲、連携条件、運用体制を確定" if context.proposal_category != "web" else "ページ数、CMS範囲、特殊機能、素材準備を確定"),
    ]
    for idx, (title, body) in enumerate(cards):
        add_card(slide, title, body, 5.38, 1.55 + idx * 1.38, 6.9, 1.08, SECTION_COLORS[idx], COLORS["white"], number=str(idx + 1))
    add_insight_band(slide, "営業判断", "予算と見積の差分を早期に共有し、提出前に必須範囲とオプション範囲を合意します。", 0.92, 5.72, 11.4, 0.58)
    add_footer(slide, slide_data.slide_no)


def add_estimate_priority_slide(prs: Presentation, slide_data: PowerPointSlide, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_header(slide, slide_data.title or "必須・推奨・オプション対応", "SCOPE", accent=COLORS["blue"])
    columns = [
        ("必須対応", context.estimate.required, COLORS["teal"], COLORS["teal_light"]),
        ("推奨対応", context.estimate.recommended or ["次回確認"], COLORS["blue"], COLORS["blue_light"]),
        ("オプション対応", context.estimate.optional or ["次回確認"], COLORS["orange"], COLORS["orange_light"]),
    ]
    for idx, (title, items, accent, fill) in enumerate(columns):
        x = 0.86 + idx * 4.1
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.62, 3.65, 4.4, fill=fill, line=COLORS["white"])
        add_text(slide, title, x + 0.25, 1.98, 3.15, 0.3, size=17, color=accent, bold=True, align=PP_ALIGN.CENTER)
        add_bullet_list(slide, items, x + 0.34, 2.62, 2.95, 2.7, max_items=5, size=13)
    add_footer(slide, slide_data.slide_no)


def add_summary_slide(prs: Presentation, slide_data: PowerPointSlide, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_header(slide, slide_data.title, "まとめ")
    confirmation = f"次回確認事項: {'・'.join(context.confirmation_items[:3])}" if context.confirmation_items else "次回確認事項を整理"
    values = ensure_items(unique_items(slide_data.bullets, 2) + [confirmation], ["課題仮説に基づく実行方針", "成果につながる導入設計", "次回確認事項の整理"], 3)

    for idx, item in enumerate(values[:3]):
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.95 + idx * 4.02, 2.08, 3.35, 2.55, fill=[COLORS["teal_light"], COLORS["blue_light"], COLORS["orange_light"]][idx], line=COLORS["white"])
        add_text(slide, f"提供価値 {idx + 1}", 1.2 + idx * 4.02, 2.38, 2.85, 0.32, size=16, color=SECTION_COLORS[idx], bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, _trim(item, 42), 1.25 + idx * 4.02, 3.15, 2.75, 0.72, size=18, color=COLORS["text"], bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, slide_data.slide_no)


def add_next_steps_slide(prs: Presentation, slide_data: PowerPointSlide, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_header(slide, slide_data.title or "今後の進め方", "次のステップ", accent=COLORS["blue"])
    steps = [
        "1. 目的・KPI・優先範囲の合意",
        "2. 提案範囲・必要データの確定" if context.proposal_category != "web" else "2. サイト構成・必要コンテンツの確定",
        "3. 見積・スケジュール・体制の最終化",
        "4. キックオフ・データ準備・実行開始" if context.proposal_category != "web" else "4. キックオフ・素材準備・制作開始",
    ]
    add_next_action_cards(slide, steps, 0.82, 1.9, 11.7)
    confirmations = ensure_items(context.confirmation_items + slide_data.bullets, ["予算内で優先する必須範囲", "希望納期に対する準備状況", "運用担当と確認フロー"], 4)
    for idx, item in enumerate(confirmations[:4]):
        x = 0.92 + (idx % 2) * 5.88
        y = 4.08 + (idx // 2) * 1.04
        add_card(slide, f"確認事項 {idx + 1}", item, x, y, 5.24, 0.9, SECTION_COLORS[idx], COLORS["white"])
    add_footer(slide, slide_data.slide_no)


def add_quality_comparison_slide(prs: Presentation, slide_data: PowerPointSlide, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_header(slide, slide_data.title, "COMPARISON", accent=COLORS["blue"])
    rows = []
    for idx, item in enumerate(unique_items(slide_data.bullets, 4), start=1):
        rows.append([f"観点 {idx}", item, "評価基準を合意"])
    if not rows:
        rows = [["観点", "入力内容を整理してください", "評価基準を合意"]]
    add_table(slide, ["比較軸", "入力情報", "確認ポイント"], rows, 0.9, 1.72, 11.55, 3.35, [1.8, 7.0, 2.75])
    add_insight_band(slide, "品質確認", "比較・競合・Before/Afterの内容を編集可能な表として整理しました。", 0.95, 5.55, 11.45, 0.72)
    add_footer(slide, slide_data.slide_no)


def add_quality_timeline_slide(prs: Presentation, slide_data: PowerPointSlide, context: PptxContext, *, section: str = "工程") -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_header(slide, slide_data.title, section, accent=COLORS["teal"])
    phases = unique_items(slide_data.bullets, 5)
    add_timeline(slide, phases, 0.88, 1.72, 11.6)
    for idx, item in enumerate(phases[:5]):
        add_card(slide, f"工程 {idx + 1}", item, 0.95 + idx * 2.3, 4.15, 2.05, 1.05, SECTION_COLORS[idx % len(SECTION_COLORS)], COLORS["white"])
    add_footer(slide, slide_data.slide_no)


def add_quality_kpi_slide(prs: Presentation, slide_data: PowerPointSlide, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_header(slide, slide_data.title, "KPI", accent=COLORS["green"])
    body = "\n".join(slide_data.bullets)
    numbers = unique_items(extract_numbers(body), 4)
    if numbers:
        metric_titles = [f"数値 {idx + 1}" for idx in range(len(numbers[:4]))]
        metric_values = numbers[:4]
    else:
        metric_titles = ["現状値", "目標値", "測定方法", "判定基準"]
        metric_values = ["入力情報を整理", "顧客と合意", "ログで測定", "PoCで確定"]
    for idx, value in enumerate(metric_values[:4]):
        add_metric_card(slide, metric_titles[idx], value, 0.95 + idx * 2.88, 1.72, 2.5, 1.45, SECTION_COLORS[idx % len(SECTION_COLORS)])
    add_bullet_list(slide, slide_data.bullets, 1.0, 3.7, 11.2, 1.82, max_items=5, size=13)
    add_insight_band(slide, "数値の確認", "本文中の数値をそのまま保持し、存在しない実績値は追加していません。", 0.95, 5.72, 11.45, 0.58)
    add_footer(slide, slide_data.slide_no)


def add_quality_flow_slide(prs: Presentation, slide_data: PowerPointSlide, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_header(slide, slide_data.title, "FLOW", accent=COLORS["purple"])
    steps = unique_items(slide_data.bullets, 5)
    add_step_flow(slide, steps, 0.95, 2.0, 11.35, 1.7)
    add_bullet_list(slide, steps, 1.0, 4.4, 11.1, 1.15, max_items=5, size=13)
    add_footer(slide, slide_data.slide_no)


def add_quality_matrix_slide(prs: Presentation, slide_data: PowerPointSlide, context: PptxContext) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_header(slide, slide_data.title, "MATRIX", accent=COLORS["orange"])
    items = unique_items(slide_data.bullets, 4)
    fallback_labels = ["評価対象を整理", "期待効果を整理", "検証方法を定義", "優先度を合意"]
    labels = (items + fallback_labels)[:4]
    positions = [(0.95, 1.78), (6.72, 1.78), (0.95, 4.15), (6.72, 4.15)]
    for idx, (x, y) in enumerate(positions):
        add_card(slide, ["高優先", "高効果", "要検証", "低優先"][idx], labels[idx], x, y, 5.32, 1.45, SECTION_COLORS[idx], COLORS["white"])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 6.48, 1.58, 0.03, 4.28, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0.9, 3.78, 11.52, 0.03, fill=COLORS["line_dark"], line=COLORS["line_dark"])
    add_footer(slide, slide_data.slide_no)


def add_generic_slide(prs: Presentation, slide_data: PowerPointSlide) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_header(slide, slide_data.title, "提案内容")
    add_bullet_list(slide, slide_data.bullets, 0.9, 1.78, 7.55, 3.9, max_items=5)
    add_visual_frame(slide, slide_data.visual_suggestion or "図表・画面イメージ・実績画像を配置", 9.0, 1.78, 3.25, 3.9)
    add_footer(slide, slide_data.slide_no)


def add_win_probability_slide(prs: Presentation, win: WinProbability, slide_no: int) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    rank_color = rank_color_for(win.rank)
    add_header(slide, "受注確率判定", "商談判断", accent=rank_color)
    probability = win.probability or rank_probability_for(win.rank)
    risk_score = max(1, min(5, win.risk_score or risk_score_for_probability(probability, len(win.risk_factors))))
    risk_label = win.risk_label or risk_label_for(risk_score)
    projected = win.projected_probability_after_actions or projected_probability_for(
        probability,
        risk_score,
        len(win.improvement_actions or win.recommended_next_actions),
    )

    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.95, 1.46, 3.28, 3.1, fill=rank_light_color_for(win.rank), line=rank_color)
    add_text(slide, "受注確率", 1.22, 1.76, 2.72, 0.24, size=13, color=COLORS["muted"], bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, f"{probability}%", 1.18, 2.1, 2.82, 0.68, size=42, color=rank_color, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, f"{win.rank}ランク", 1.32, 2.88, 2.48, 0.3, size=19, color=rank_color, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, f"受注リスク {risk_label}", 1.16, 3.38, 2.86, 0.24, size=12, color=COLORS["red"], bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, f"向上予測 {probability}% → {projected}%", 1.12, 3.8, 2.96, 0.28, size=14, color=COLORS["blue"], bold=True, align=PP_ALIGN.CENTER)

    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 4.62, 1.58, 7.55, 1.58, fill=COLORS["white"], line=COLORS["line"])
    add_text(slide, "判定理由", 4.96, 1.84, 1.3, 0.26, size=13, color=rank_color, bold=True)
    add_text(slide, _trim(win.reason, 82), 6.18, 1.84, 5.55, 0.62, size=13, color=COLORS["text"], bold=True)

    add_factor_column(slide, "リスク要因", win.risk_factors, 0.95, 5.02, COLORS["red"])
    add_factor_column(slide, "改善アクション", win.improvement_actions or win.recommended_next_actions, 4.68, 5.02, COLORS["blue"])
    add_factor_column(slide, "加点要因", win.positive_factors, 8.41, 5.02, COLORS["teal"])
    add_footer(slide, slide_no)
