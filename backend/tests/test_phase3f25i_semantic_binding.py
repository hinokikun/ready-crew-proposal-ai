from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from pptx import Presentation

from app.services.pptx_parts.native_trace_content_adapter import (
    FailureReason,
    SAMPLE_STRINGS,
    render_native_role_dry_run,
)
from app.services.pptx_parts.native_trace_renderers import dispatch_approved_native_slide
from app.services.pptx_parts.native_trace_renderers import (
    _iter_all_shapes,
    _semantic_expanded_bounds,
    _semantic_point_in_bounds,
    _semantic_shape_center,
)


ROLE_CASES = (
    ("CURRENT_STATE", "現状理解", "S03", {
        "current_state": "FAJで確認された現状情報を整理します。",
        "workflow": "FAJで確認された業務フローを整理します。",
        "issue": "FAJで確認された課題を整理します。",
        "priority": "確認済みの改善優先テーマを整理します。",
    }),
    ("PROBLEM_ANALYSIS", "主要課題", "S04", {
        "problem_summary": "FAJで確認された主要課題を整理します。",
        "priority": "確認済みの優先テーマを整理します。",
        "evidence": "確認済み情報を判断材料として整理します。",
        "next_step": "不足情報は確認後に確定します。",
    }),
    ("SOLUTION_CONCEPT", "提案コンセプト", "S05", {
        "concept": "FAJの確認済み方針を整理します。",
        "measure": "確認済みの施策候補を整理します。",
        "operation": "実行条件を確認します。",
        "outcome": "確認済みの成果条件を整理します。",
    }),
    ("SOLUTION_APPROACH", "導入戦略", "S06", {
        "approach": "FAJで確認された導入方針を整理します。",
        "step": "確認済みの導入ステップを整理します。",
        "condition": "実行条件は確認後に確定します。",
        "outcome": "確認済みの成果条件を整理します。",
    }),
)

PARTIAL_FALLBACK_CASES = (
    ("CURRENT_STATE", "現状理解", "S03", "current_state", "FAJで確認された受注業務の現状を整理します。"),
    ("PROBLEM_ANALYSIS", "主要課題", "S04", "problem_summary", "FAJで確認された主要課題を整理します。"),
    ("SOLUTION_CONCEPT", "提案コンセプト", "S05", "concept", "FAJで確認された課題に対する方針を整理します。"),
    ("SOLUTION_APPROACH", "導入戦略", "S06", "approach", "FAJで確認された導入方針を整理します。"),
)


def _slide(title: str) -> SimpleNamespace:
    return SimpleNamespace(title=title, layout="", bullets=[], slide_no=1)


def _context(role: str, values: dict[str, str]) -> SimpleNamespace:
    candidates = [
        {
            "semantic_type": f"{role.lower()}.{group}",
            "value": value,
            "source_type": "customer_input",
            "source_field": f"faj.{group}",
            "source_reference": f"faj://verified/{group}",
            "authority": "USER_EXPLICIT",
            "review_state": "CONFIRMED",
            "admissible_as_evidence": True,
        }
        for group, value in values.items()
    ]
    return SimpleNamespace(semantic_candidates={"candidates": candidates})


def _package_text(pptx_bytes: bytes) -> str:
    with ZipFile(BytesIO(pptx_bytes)) as package:
        return "\n".join(
            package.read(name).decode("utf-8", errors="ignore")
            for name in package.namelist()
            if name.endswith(".xml")
        )


@pytest.mark.parametrize("role,title,slide_id,values", ROLE_CASES)
def test_faj_semantic_binding_for_slides_05_to_08_is_explicit_and_sample_free(role, title, slide_id, values) -> None:
    result = render_native_role_dry_run(
        role,
        data=SimpleNamespace(),
        context=_context(role, values),
        slide=_slide(title),
        surface="summary",
        slide_id=slide_id,
    )
    assert result.success is True
    assert result.validation["injected"]["sample_content_leak_free"] is True
    text = _package_text(result.pptx_bytes or b"")
    assert all(sample not in text for sample in SAMPLE_STRINGS[role])
    assert all(value in text for value in values.values())
    assert "ProposalPilot" not in text
    assert "AI営業秘書" not in text


@pytest.mark.parametrize("role,title,slide_id,values", ROLE_CASES)
def test_missing_or_generated_semantic_evidence_fails_closed(role, title, slide_id, values) -> None:
    missing = render_native_role_dry_run(
        role,
        data=SimpleNamespace(),
        context=SimpleNamespace(),
        slide=_slide(title),
        surface="summary",
        slide_id=slide_id,
    )
    assert missing.success is False
    assert missing.failure_reason == FailureReason.EVIDENCE_REQUIRED.value

    generated_values = dict(values)
    first_group = next(iter(generated_values))
    generated = _context(role, generated_values)
    generated.semantic_candidates["candidates"][0]["authority"] = "AI_PROPOSED"
    generated.semantic_candidates["candidates"][0]["admissible_as_evidence"] = False
    generated.semantic_candidates["candidates"][0]["inferred"] = True
    rejected = render_native_role_dry_run(
        role,
        data=SimpleNamespace(),
        context=generated,
        slide=_slide(title),
        surface="summary",
        slide_id=slide_id,
    )
    assert rejected.success is False
    assert rejected.failure_reason == FailureReason.EVIDENCE_REQUIRED.value


@pytest.mark.parametrize("role,title,slide_id,values", ROLE_CASES)
def test_missing_evidence_uses_proposal_master_fallback_without_generic_copy(role, title, slide_id, values) -> None:
    prs = Presentation()
    trace = dispatch_approved_native_slide(
        prs,
        _slide(title),
        SimpleNamespace(),
        SimpleNamespace(),
        int(slide_id[1:]) - 1,
        role=role,
        surface="summary",
        visual_fallback=True,
    )
    assert trace["NATIVE_RENDERED"] is False
    assert trace["FALLBACK_USED"] is True
    assert trace["FALLBACK_RENDERED"] is True
    assert trace["FAILURE_REASON"] == FailureReason.EVIDENCE_REQUIRED.value
    text = "\n".join(shape.text for shape in prs.slides[-1].shapes if getattr(shape, "has_text_frame", False))
    assert all(sample not in text for sample in SAMPLE_STRINGS[role])
    assert "ProposalPilot" not in text
    assert "AI営業秘書" not in text
    assert "確認後に確定" in text


@pytest.mark.parametrize("role,title,slide_id,group,value", PARTIAL_FALLBACK_CASES)
def test_partial_fallback_uses_role_specific_verified_content_without_repeated_labels(
    role: str,
    title: str,
    slide_id: str,
    group: str,
    value: str,
) -> None:
    prs = Presentation()
    trace = dispatch_approved_native_slide(
        prs,
        _slide(title),
        SimpleNamespace(),
        _context(role, {group: value}),
        int(slide_id[1:]) - 1,
        role=role,
        surface="summary",
        visual_fallback=True,
    )
    assert trace["NATIVE_RENDERED"] is False
    assert trace["FALLBACK_RENDERED"] is True
    text = "\n".join(shape.text for shape in prs.slides[-1].shapes if getattr(shape, "has_text_frame", False))
    assert value in text
    assert "確認済み情報" not in text
    assert "未確認項目" not in text
    assert text.count("確認後に確定") <= 3
    assert all(sample not in text for sample in SAMPLE_STRINGS[role])


@pytest.mark.parametrize(
    "role,title,slide_id",
    (
        ("CURRENT_STATE", "現状理解", "S03"),
        ("PROBLEM_ANALYSIS", "主要課題", "S04"),
        ("SOLUTION_CONCEPT", "提案コンセプト", "S05"),
        ("SOLUTION_APPROACH", "導入戦略", "S06"),
    ),
)
def test_semantic_fallback_removes_empty_body_containers(role: str, title: str, slide_id: str) -> None:
    prs = Presentation()
    dispatch_approved_native_slide(
        prs,
        _slide(title),
        SimpleNamespace(),
        SimpleNamespace(),
        int(slide_id[1:]) - 1,
        role=role,
        surface="summary",
        visual_fallback=True,
    )
    prefix = f"trace:{slide_id}:"
    shapes = list(_iter_all_shapes(prs.slides[-1].shapes))
    body_shapes = [
        shape
        for shape in shapes
        if str(getattr(shape, "name", "") or "").startswith(prefix)
        and float(shape.top) >= 1.45 * 914400
        and float(shape.top) < 6.85 * 914400
    ]
    active_text = [
        shape
        for shape in body_shapes
        if getattr(shape, "has_text_frame", False)
        and str(getattr(shape, "text", "") or "").strip()
    ]
    empty_containers = [
        shape
        for shape in body_shapes
        if float(shape.width) * float(shape.height) >= (0.65 * 914400) ** 2
        and not any(
            _semantic_point_in_bounds(
                _semantic_shape_center(text_shape), _semantic_expanded_bounds(shape)
            )
            for text_shape in active_text
        )
    ]
    assert empty_containers == []


def test_solution_concept_insight_bar_stays_above_footer() -> None:
    prs = Presentation()
    dispatch_approved_native_slide(
        prs,
        _slide("提案コンセプト"),
        SimpleNamespace(),
        SimpleNamespace(),
        6,
        role="SOLUTION_CONCEPT",
        surface="summary",
        visual_fallback=True,
    )
    slide = prs.slides[-1]
    insight = next(
        shape
        for shape in slide.shapes
        if shape.name == "trace:S05:content.auto.100"
    )
    footer = next(shape for shape in slide.shapes if shape.name == "trace:S05:footer.brand")
    assert insight.top + insight.height < footer.top
