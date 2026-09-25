from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

from pptx import Presentation

from app.services.pptx_parts.native_trace_content_adapter import (
    FailureReason,
    SAMPLE_STRINGS,
    render_native_role_dry_run,
)
from app.services.pptx_parts.native_trace_renderers import dispatch_approved_native_slide


def _slide(title: str) -> SimpleNamespace:
    return SimpleNamespace(title=title, layout="", bullets=[], slide_no=1)


def _candidate(role: str, group: str, value: str | list[str], *, authority: str = "USER_EXPLICIT", review_state: str = "CONFIRMED") -> dict[str, object]:
    return {
        "semantic_type": f"{role.lower()}.{group}",
        "value": value,
        "source_type": "customer_input",
        "source_field": f"faj.{group}",
        "source_reference": f"faj://verified/{group}",
        "authority": authority,
        "review_state": review_state,
        "admissible_as_evidence": authority == "USER_EXPLICIT",
        "inferred": authority == "AI_PROPOSED",
    }


def _summary_context(role: str) -> SimpleNamespace:
    groups = {
        "EXECUTIVE_SUMMARY": {
            "summary_context": "FAJ案件の確認済み論点を整理し、判断材料を揃えます。",
            "background": "FAJで確認済みの業務課題を優先して整理します。",
            "current_state": "FAJの現状データで確認された分散課題を整理します。",
            "conclusion": "確認済みのFAJ課題に対する実行可能な方針を選定します。",
            "expected_effects": ["確認済みの工数課題を可視化", "判断材料を整理", "運用確認を継続"],
            "decision_items": ["対象範囲を確認", "測定方法を確認", "費用条件を確認", "推進体制を確認"],
            "insight": "確認済みの範囲から段階的に判断します。",
        },
        "DECISION_AND_EXPECTED_EFFECTS": {
            "summary": "FAJ向けに確認済みの課題と方針を整理した提案です。",
            "headline": ["FAJで確認された課題を整理し", "検証可能な判断と次の行動へつなげます"],
            "why_now": ["FAJで確認された業務上の課題", "確認済みデータの分散", "次の判断に必要な論点"],
            "policy": ["確認済みの範囲から整理", "実測可能な項目を定義", "段階的に検証"],
            "investment": ["確認済みの対象範囲", "必要作業と移行条件", "運用確認と継続改善"],
            "effects": ["確認済み課題の可視化", "判断材料の整理", "運用品質の確認", "次段階の条件整理"],
            "decisions": ["対象範囲の確認", "次段階の実施条件の確認"],
            "insight": "確認済みの事実を基に次の判断へ進みます。",
        },
    }[role]
    return SimpleNamespace(
        semantic_candidates={"candidates": [_candidate(role, group, value) for group, value in groups.items()]}
    )


def _proposal_summary_context() -> SimpleNamespace:
    groups = {
        "summary": "FAJ案件の背景と確認済み課題から、次の判断材料を整理します。",
        "current_state": [
            "FAJで確認された案件情報の分散",
            "確認済みの課題が複数部門にまたがる",
            "既存運用の確認事項が残っている",
            "判断に必要な情報を整理中",
        ],
        "key_measures": [
            "確認済み案件情報を整理",
            "対象範囲を確認",
            "確認済みの運用課題を可視化",
            "関係者の確認事項を整理",
            "提出前の確認手順を定義",
            "確認後に運用へ反映",
        ],
        "expected_effects": [
            "判断材料の整理", "確認後に確定", "確認済み範囲の可視化",
            "次の判断条件", "運用確認の継続", "確認後に確定",
        ],
        "decision_items": [
            "対象範囲を確認", "現状課題を確認", "必要な確認事項を整理",
            "関係者を確認", "証拠の出所を確認", "未確認項目を整理",
            "次段階の条件を確認", "判断時期を確認", "確認後に確定",
        ],
        "insight": [
            "確認済み情報のみ表示します。",
            "未確認項目は確認後に確定します。",
            "FAJ案件の確認範囲から判断します。",
            "追加情報は確認後に反映します。",
        ],
    }
    return SimpleNamespace(
        semantic_candidates={"candidates": [_candidate("PROPOSAL_SUMMARY", group, value) for group, value in groups.items()]}
    )


def _package_text(pptx_bytes: bytes) -> str:
    with ZipFile(BytesIO(pptx_bytes)) as package:
        return "\n".join(
            package.read(name).decode("utf-8", errors="ignore")
            for name in package.namelist()
            if name.endswith(".xml")
        )


def test_faj_summary_native_binding_has_no_canonical_sample_strings() -> None:
    for role, title, slide_id in (
        ("EXECUTIVE_SUMMARY", "経営判断の要点", "S02"),
        ("DECISION_AND_EXPECTED_EFFECTS", "本提案の結論と期待効果", "S03"),
    ):
        result = render_native_role_dry_run(
            role,
            data=SimpleNamespace(),
            context=_summary_context(role),
            slide=_slide(title),
            surface="summary",
            slide_id=slide_id,
        )
        assert result.success is True
        assert result.validation["injected"]["sample_content_leak_free"] is True
        text = _package_text(result.pptx_bytes or b"")
        assert all(sample not in text for sample in SAMPLE_STRINGS[role])
        assert "FAJ" in text


def test_generated_or_unknown_summary_content_falls_back_before_native_clone() -> None:
    generated = _summary_context("EXECUTIVE_SUMMARY")
    generated.semantic_candidates["candidates"][0] = _candidate(
        "EXECUTIVE_SUMMARY",
        "summary_context",
        "AIが生成した仮説",
        authority="AI_PROPOSED",
        review_state="UNCONFIRMED",
    )
    result = render_native_role_dry_run(
        "EXECUTIVE_SUMMARY",
        data=SimpleNamespace(),
        context=generated,
        slide=_slide("経営判断の要点"),
        surface="summary",
        slide_id="S02",
    )
    assert result.success is False
    assert result.failure_reason == FailureReason.EVIDENCE_REQUIRED.value

    unknown = render_native_role_dry_run(
        "DECISION_AND_EXPECTED_EFFECTS",
        data=SimpleNamespace(),
        context=SimpleNamespace(),
        slide=_slide("本提案の結論と期待効果"),
        surface="summary",
        slide_id="S03",
    )
    assert unknown.success is False
    assert unknown.failure_reason == FailureReason.EVIDENCE_REQUIRED.value


def test_missing_semantic_group_cannot_render_native_summary() -> None:
    context = _summary_context("EXECUTIVE_SUMMARY")
    context.semantic_candidates["candidates"] = [
        item for item in context.semantic_candidates["candidates"]
        if not str(item["semantic_type"]).endswith("decision_items")
    ]
    result = render_native_role_dry_run(
        "EXECUTIVE_SUMMARY",
        data=SimpleNamespace(),
        context=context,
        slide=_slide("経営判断の要点"),
        surface="summary",
        slide_id="S02",
    )
    assert result.success is False
    assert result.failure_reason == FailureReason.EVIDENCE_REQUIRED.value


def test_implementation_configuration_fallback_is_semantic_safe() -> None:
    prs = Presentation()
    trace = dispatch_approved_native_slide(
        prs,
        _slide("導入構成"),
        SimpleNamespace(),
        SimpleNamespace(),
        8,
        role="IMPLEMENTATION_CONFIGURATION",
        surface="summary",
        visual_fallback=True,
    )
    assert trace["NATIVE_RENDERED"] is False
    assert trace["FALLBACK_USED"] is True
    assert trace["FAILURE_REASON"] == FailureReason.EVIDENCE_REQUIRED.value

    text = "\n".join(
        shape.text for shape in prs.slides[0].shapes if getattr(shape, "has_text_frame", False)
    )
    forbidden = (
        "確認中を構築",
        "業務の標準化",
        "標準化",
        "データの一元管理",
        "部門間の連携",
        "コストを最適化",
        "持続的な成長",
        "ERP",
    )
    assert all(fragment not in text for fragment in forbidden)
    assert "確認済み情報のみ表示し、未確認の項目は確認後に確定します。" in text


def test_proposal_summary_verified_binding_has_no_generic_samples() -> None:
    result = render_native_role_dry_run(
        "PROPOSAL_SUMMARY",
        data=SimpleNamespace(),
        context=_proposal_summary_context(),
        slide=_slide("提案サマリー"),
        surface="summary",
        slide_id="S02",
    )
    assert result.success is True
    assert result.validation["injected"]["sample_content_leak_free"] is True
    text = _package_text(result.pptx_bytes or b"")
    assert all(sample not in text for sample in SAMPLE_STRINGS["PROPOSAL_SUMMARY"])
    assert "ProposalPilot" not in text
    assert "AI営業秘書" not in text
    assert all(value not in text for value in ("約70%削減", "約1.5倍", "約20%向上"))


def test_proposal_summary_missing_evidence_uses_safe_master_fallback() -> None:
    prs = Presentation()
    trace = dispatch_approved_native_slide(
        prs,
        _slide("提案サマリー"),
        SimpleNamespace(),
        SimpleNamespace(),
        3,
        role="PROPOSAL_SUMMARY",
        surface="summary",
        visual_fallback=True,
    )
    assert trace["NATIVE_RENDERED"] is False
    assert trace["FALLBACK_USED"] is True
    assert trace["FALLBACK_RENDERED"] is True
    assert trace["FAILURE_REASON"] == FailureReason.EVIDENCE_REQUIRED.value
    text = "\n".join(shape.text for shape in prs.slides[0].shapes if getattr(shape, "has_text_frame", False))
    assert "ProposalPilot" not in text
    assert "AI営業秘書" not in text
    assert "確認済み情報を整理し、未確認項目は確認後に確定します。" in text
    assert all(value not in text for value in ("提案書自動生成AI", "ナレッジ活用", "レビュー運用", "約70%削減", "約1.5倍", "約20%向上"))


def test_proposal_summary_insight_bar_is_compacted_and_page_number_is_four() -> None:
    prs = Presentation()
    trace = dispatch_approved_native_slide(
        prs,
        _slide("提案サマリー"),
        SimpleNamespace(),
        _proposal_summary_context(),
        3,
        role="PROPOSAL_SUMMARY",
        surface="summary",
        visual_fallback=True,
    )
    assert trace["NATIVE_RENDERED"] is True
    slide = prs.slides[0]
    by_name = {shape.name: shape for shape in slide.shapes}
    assert by_name["trace:S02:content.auto"].text == "04"
    assert by_name["trace:S02:content.auto.135"].text == "04"
    assert by_name["trace:S02:content.auto.130"].text
    assert by_name["trace:S02:content.auto.131"].text
    assert by_name["trace:S02:content.auto.132"].text == ""
    assert by_name["trace:S02:content.auto.133"].text == ""
    assert by_name["trace:S02:content.auto.130"].top + by_name["trace:S02:content.auto.130"].height <= by_name["trace:S02:content.auto.131"].top
