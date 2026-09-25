from dataclasses import replace
from io import BytesIO
from types import SimpleNamespace

from pptx import Presentation

from app.config import settings
from app.models import PowerPointData, PowerPointSlide, PptxDownloadRequest
from app.services.pptx_parts import slides
from app.services.pptx_parts.native_trace_registry import (
    APPROVED_NATIVE_REGISTRY,
    NEW_NATIVE_TRACE_ROLES,
    UNAPPROVED_NATIVE_ROLES,
    get_native_role_spec,
    native_role_gate,
    resolve_approved_native_role,
)
from app.services.pptx_parts.native_trace_renderers import dispatch_approved_native_slide, render_approved_native_slide
from app.services.pptx_service import build_pptx_context, build_pptx_result, build_summary_slides


def _slide(title: str, layout: str = "") -> SimpleNamespace:
    return SimpleNamespace(title=title, layout=layout, bullets=[], slide_no=1)


def _verified_proposal_summary_context() -> SimpleNamespace:
    groups = {
        "summary": ("proposal_summary.summary", ["FAJで確認済みの提案論点"]),
        "current_state": (
            "proposal_summary.current_state",
            ["確認済み現状1", "確認済み現状2", "確認済み現状3", "確認済み現状4"],
        ),
        "key_measures": (
            "proposal_summary.key_measure",
            ["確認済み施策1", "確認済み施策2", "確認済み施策3", "確認済み施策4", "確認済み施策5", "確認済み施策6"],
        ),
        "expected_effects": (
            "proposal_summary.expected_effect",
            ["確認済み効果1", "確認済み効果2", "確認済み効果3", "確認済み効果4", "確認済み効果5", "確認済み効果6"],
        ),
        "decision_items": (
            "proposal_summary.decision",
            [
                "確認済み判断事項1",
                "確認済み判断事項2",
                "確認済み判断事項3",
                "確認済み判断事項4",
                "確認済み判断事項5",
                "確認済み判断事項6",
                "確認済み判断事項7",
                "確認済み判断事項8",
                "確認済み判断事項9",
            ],
        ),
        "insight": ("proposal_summary.insight", ["確認済み示唆1", "確認済み示唆2", "確認済み示唆3", "確認済み示唆4"]),
    }
    return SimpleNamespace(
        semantic_candidates=[
            {
                "semantic_type": semantic_type,
                "value": values,
                "source_type": "customer_input",
                "source_field": f"faj.proposal_summary.{group}",
                "source_reference": f"faj://verified/proposal_summary/{group}",
                "authority": "USER_EXPLICIT",
                "review_state": "CONFIRMED",
                "admissible_as_evidence": True,
                "inferred": False,
            }
            for group, (semantic_type, values) in groups.items()
        ]
    )


def test_registry_contains_only_human_approved_native_roles() -> None:
    assert APPROVED_NATIVE_REGISTRY
    assert {"CASE_STUDY", "ROI_OR_EFFECT", "MARKET_ANALYSIS", "TARGET_ANALYSIS"}.isdisjoint(
        APPROVED_NATIVE_REGISTRY
    )
    assert UNAPPROVED_NATIVE_ROLES == {
        "CASE_STUDY",
        "ROI_OR_EFFECT",
        "MARKET_ANALYSIS",
        "TARGET_ANALYSIS",
    }
    assert all(spec.human_approved for spec in APPROVED_NATIVE_REGISTRY.values())
    assert len(APPROVED_NATIVE_REGISTRY) == 18
    assert sum(spec.classification == "A" for spec in APPROVED_NATIVE_REGISTRY.values()) == 6
    assert sum(spec.classification == "B" for spec in APPROVED_NATIVE_REGISTRY.values()) == 8
    assert sum(spec.classification == "C" for spec in APPROVED_NATIVE_REGISTRY.values()) == 4
    assert {role for role, spec in APPROVED_NATIVE_REGISTRY.items() if spec.classification == "A"} == set(
        NEW_NATIVE_TRACE_ROLES
    )
    assert sum(spec.alias_of is not None for spec in APPROVED_NATIVE_REGISTRY.values()) == 4
    assert all("phase3f10e_sanitized_native_assets" in spec.source_artifact or "phase3f08_native_trace_gap" in spec.source_artifact for spec in APPROVED_NATIVE_REGISTRY.values())


def test_role_gate_requires_all_runtime_checks() -> None:
    eligible, checks = native_role_gate("KPI", template_available=True)
    assert eligible is True
    assert all(checks.values())

    eligible, checks = native_role_gate("KPI", template_available=True, text_fit_valid=False)
    assert eligible is False
    assert checks["TEXT_FIT_VALID"] is False

    eligible, checks = native_role_gate("CASE_STUDY", template_available=True)
    assert eligible is False
    assert checks["HUMAN_APPROVED"] is False


def test_unapproved_roles_resolve_to_existing_renderer_only() -> None:
    assert resolve_approved_native_role(_slide("市場分析と競合環境"), 2) is None
    assert resolve_approved_native_role(_slide("ターゲット分析"), 3) is None
    assert resolve_approved_native_role(_slide("ROIと費用対効果"), 4) is None
    assert resolve_approved_native_role(_slide("導入事例"), 5) is None
    assert get_native_role_spec("CASE_STUDY") is None


def test_approved_trace_templates_respect_adapter_and_render_generic_role() -> None:
    for role in ("ESTIMATE", "KPI", "COMPETITIVE_COMPARISON", "WIN_PROBABILITY", "SCHEDULE"):
        prs = Presentation()
        prs.slide_width = 12192000
        prs.slide_height = 6858000
        rendered = render_approved_native_slide(
            prs,
            _slide(role),
            SimpleNamespace(),
            SimpleNamespace(win_probability=None),
            0,
            role=role,
        )
        assert rendered is False
        assert len(prs.slides) == 0

    prs = Presentation()
    prs.slide_width = 12192000
    prs.slide_height = 6858000
    rendered = render_approved_native_slide(
        prs,
        _slide("提案サマリー"),
        SimpleNamespace(),
        _verified_proposal_summary_context(),
        1,
        role="PROPOSAL_SUMMARY",
        surface="summary",
    )
    assert rendered is True
    assert len(prs.slides) == 1
    assert len(prs.slides[0].shapes) > 0
    assert not any(shape.shape_type == 13 for shape in prs.slides[0].shapes)


def test_flag_off_keeps_existing_renderer_path(monkeypatch) -> None:
    called = {"native": 0, "legacy": 0}

    def native(*args, **kwargs):
        called["native"] += 1
        return True

    def legacy(*args, **kwargs):
        called["legacy"] += 1
        args[0].slides.add_slide(args[0].slide_layouts[6])

    monkeypatch.setattr(slides, "settings", replace(settings, pptx_approved_native_renderer_enabled=False))
    monkeypatch.setattr(slides, "dispatch_approved_native_slide", native)
    monkeypatch.setattr(slides, "render_v5_masterpiece_slide", legacy)
    prs = Presentation()
    slides.add_designed_slide(prs, _slide("KPI設計と効果測定"), SimpleNamespace(), 0, SimpleNamespace())
    assert called == {"native": 0, "legacy": 1}


def test_native_failure_isolated_to_existing_renderer(monkeypatch) -> None:
    called = {"legacy": 0}

    def failing_native(*args, **kwargs):
        args[0].slides.add_slide(args[0].slide_layouts[6])
        raise RuntimeError("synthetic native failure")

    def legacy(*args, **kwargs):
        called["legacy"] += 1
        args[0].slides.add_slide(args[0].slide_layouts[6])

    monkeypatch.setattr(slides, "settings", replace(settings, pptx_approved_native_renderer_enabled=True))
    monkeypatch.setattr(slides, "dispatch_approved_native_slide", failing_native)
    monkeypatch.setattr(slides, "render_v5_masterpiece_slide", legacy)
    prs = Presentation()
    slides.add_designed_slide(prs, _slide("KPI設計と効果測定"), SimpleNamespace(), 0, SimpleNamespace())
    assert called["legacy"] == 1
    assert len(prs.slides) == 1


def test_summary_role_smoke_remains_eleven_and_detail_unapproved_falls_back() -> None:
    summary_titles = [
        "表紙",
        "提案サマリー",
        "現状理解",
        "主要課題",
        "提案コンセプト",
        "Web戦略",
        "システム構成",
        "KPI設計と効果測定",
        "導入スケジュールと推進体制",
        "概算見積と判断条件",
        "次のアクション",
    ]
    assert len(summary_titles) == 11
    assert resolve_approved_native_role(_slide("ROIと費用対効果"), 10) is None
    assert resolve_approved_native_role(_slide("受注確度と次の判断"), 10) == "WIN_PROBABILITY"
    assert resolve_approved_native_role(_slide("費用概算"), 9) == "ESTIMATE"


def test_detail_related_titles_do_not_collapse_into_problem_or_kpi_native_roles() -> None:
    assert resolve_approved_native_role(_slide("主要課題"), 3) == "PROBLEM_ANALYSIS"
    assert resolve_approved_native_role(_slide("課題から導入判断までの流れ"), 4) is None
    assert resolve_approved_native_role(_slide("KPI設計と効果測定"), 14) == "KPI"
    assert resolve_approved_native_role(_slide("KPIは現状値から測定します"), 5) is None
    assert resolve_approved_native_role(_slide("次回は範囲・KPI・体制を合意します"), 8) is None


def _summary_payload() -> PptxDownloadRequest:
    return PptxDownloadRequest(
        powerpoint_generation_data=PowerPointData(
            deck_title="Summary contract fixture",
            client_name="検証社",
            slides=[
                PowerPointSlide(
                    slide_no=1,
                    layout="title",
                    title="表紙",
                    bullets=["TEST FIXTURE"],
                    speaker_notes="",
                    visual_suggestion="",
                )
            ],
        ),
        project_brief="TEST FIXTURE / NOT REAL CUSTOMER DATA / Webサイト改善案件",
        client_company_info="検証社",
        summary=True,
    )


def test_summary_role_order_contract() -> None:
    payload = _summary_payload()
    context = build_pptx_context(payload)
    summary = build_summary_slides(payload.powerpoint_generation_data.slides, context)
    assert len(summary) == 11
    assert [slide.title for slide in summary] == [
        "",
        "提案サマリー",
        "現状理解",
        "主要課題",
        "提案コンセプト",
        "Web戦略",
        "サイトマップ",
        "KPI設計",
        "スケジュール",
        "費用概算",
        "今後の進め方",
    ]


def test_flag_on_summary_generates_exactly_eleven_slides(monkeypatch) -> None:
    monkeypatch.setattr(slides, "settings", replace(settings, pptx_approved_native_renderer_enabled=True))
    result = build_pptx_result(_summary_payload(), summary_mode=True)
    deck = Presentation(BytesIO(result.pptx_bytes))
    assert len(deck.slides) == 11


def test_flag_off_summary_keeps_existing_renderer_and_eleven_slide_cap(monkeypatch) -> None:
    monkeypatch.setattr(slides, "settings", replace(settings, pptx_approved_native_renderer_enabled=False))
    result = build_pptx_result(_summary_payload(), summary_mode=True)
    deck = Presentation(BytesIO(result.pptx_bytes))
    assert len(deck.slides) == 11
