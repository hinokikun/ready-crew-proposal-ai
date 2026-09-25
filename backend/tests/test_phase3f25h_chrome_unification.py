from __future__ import annotations

from pptx import Presentation
from pptx.util import Inches

from app.services.pptx_parts.native_trace_renderers import _unify_proposal_master_chrome


def _named_textbox(slide, name: str, text: str, left: float, top: float, width: float, height: float):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    shape._element.nvSpPr.cNvPr.set("name", name)
    shape.text = text
    return shape


def test_proposal_master_chrome_unifies_page_footer_and_preserves_body() -> None:
    prs = Presentation()
    reference = prs.slides.add_slide(prs.slide_layouts[6])
    reference_date = _named_textbox(reference, "trace:REF:footer.date", "2026.08.26", 11.666, 0.204, 1.296, 0.213)
    reference_date.text_frame.paragraphs[0].runs[0].font.name = "Yu Gothic UI"
    reference_date.text_frame.paragraphs[0].runs[0].font.size = Inches(0.12)
    _named_textbox(reference, "trace:REF:footer.tagline.jp", "旧JP", 2.0, 7.0, 2.0, 0.2)
    _named_textbox(reference, "trace:REF:footer.tagline.en", "旧EN", 2.0, 7.2, 2.0, 0.2)
    _named_textbox(reference, "trace:REF:footer.brand.2", "旧Copyright", 9.0, 7.0, 2.0, 0.2)

    target = prs.slides.add_slide(prs.slide_layouts[6])
    _named_textbox(target, "trace:T:content.auto", "03", 0.1, 0.1, 0.3, 0.2)
    _named_textbox(target, "trace:T:content.auto.3", "カテゴリ", 0.8, 0.1, 1.0, 0.2)
    target_date = _named_textbox(target, "trace:T:footer.date", "2025.06.22", 11.0, 0.1, 1.0, 0.2)
    target_date.text_frame.paragraphs[0].runs[0].font.name = "Arial"
    _named_textbox(target, "trace:T:content.auto.4", "", 2.0, 0.3, 5.0, 0.02)
    _named_textbox(target, "trace:T:footer.rule", "", 0.3, 6.9, 12.0, 0.01)
    _named_textbox(target, "trace:T:footer.page", "03", 0.1, 7.1, 0.3, 0.2)
    _named_textbox(target, "trace:T:footer.brand", "旧ブランド", 1.0, 7.1, 1.0, 0.2)
    body = _named_textbox(target, "trace:T:body", "案件本文は変更しない", 2.0, 1.5, 4.0, 0.4)

    _unify_proposal_master_chrome(prs, target, 2)

    assert target.shapes[0].text == "02"
    assert body.text == "案件本文は変更しない"
    assert any(shape.name.endswith(":footer.tagline.jp") and "人とテクノロジー" in shape.text for shape in target.shapes)
    assert any(shape.name.endswith(":footer.tagline.en") and "Think Together" in shape.text for shape in target.shapes)
    assert any(shape.name.endswith(":footer.brand.2") and "提案クエスト" in shape.text for shape in target.shapes)
    footer_page = next(shape for shape in target.shapes if shape.name == "trace:T:footer.page")
    assert footer_page.text == "02"
    assert round(footer_page.left / Inches(1), 3) == 0.37
    assert target_date.text == "2026.08.26"
    assert round(target_date.left / Inches(1), 3) == 11.666
    assert target_date.text_frame.paragraphs[0].runs[0].font.name == "Yu Gothic UI"
