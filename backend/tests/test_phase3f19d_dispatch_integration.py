from __future__ import annotations

from dataclasses import replace
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

from pptx import Presentation

from app.config import settings
from app.models import PowerPointData, PowerPointSlide, PptxDownloadRequest
from app.services.pptx_parts import native_trace_renderers, slides
from app.services.pptx_parts.native_trace_content_adapter import FailureReason, render_native_role_dry_run
from app.services.pptx_parts.native_trace_registry import load_runtime_native_registry, resolve_approved_native_role
from app.services.pptx_parts.native_trace_validation import sha256_file
from app.services.pptx_service import build_pptx_result


ROOT = Path(__file__).resolve().parents[2]


def _slide(title: str, layout: str = "") -> SimpleNamespace:
    return SimpleNamespace(title=title, layout=layout, bullets=["案件固有の確認事項"], slide_no=1)


def _payload(summary: bool = True) -> PptxDownloadRequest:
    return PptxDownloadRequest(
        powerpoint_generation_data=PowerPointData(
            deck_title="Phase3F19D fixture",
            client_name="検証社",
            slides=[PowerPointSlide(slide_no=1, layout="title", title="表紙", bullets=["TEST FIXTURE"], speaker_notes="", visual_suggestion="")],
        ),
        project_brief="TEST FIXTURE / NOT REAL CUSTOMER DATA / Webサイト改善案件",
        client_company_info="検証社",
        summary=summary,
    )


def test_flag_off_never_requests_native_dispatch(monkeypatch) -> None:
    calls = {"native": 0, "legacy": 0}

    def native(*args, **kwargs):
        calls["native"] += 1
        return {"NATIVE_RENDERED": True}

    def legacy(*args, **kwargs):
        calls["legacy"] += 1
        args[0].slides.add_slide(args[0].slide_layouts[6])

    monkeypatch.setattr(slides, "settings", replace(settings, pptx_approved_native_renderer_enabled=False))
    monkeypatch.setattr(slides, "dispatch_approved_native_slide", native)
    monkeypatch.setattr(slides, "render_v5_masterpiece_slide", legacy)
    prs = Presentation()
    slides.add_designed_slide(prs, _slide("提案サマリー"), SimpleNamespace(), 1, SimpleNamespace(), surface="summary")
    assert calls == {"native": 0, "legacy": 1}


def test_flag_on_eligible_role_uses_native_renderer() -> None:
    prs = Presentation()
    prs.slide_width = 12192000
    prs.slide_height = 6858000
    trace = native_trace_renderers.dispatch_approved_native_slide(
        prs,
        _slide("提案サマリー"),
        SimpleNamespace(),
        SimpleNamespace(),
        1,
        role="PROPOSAL_SUMMARY",
        surface="summary",
    )
    assert trace["NATIVE_RENDERED"] is True
    assert trace["FALLBACK_USED"] is False
    assert trace["SAMPLE_LEAK_STATUS"] == "PASS"
    assert len(prs.slides) == 1


def test_human_unapproved_title_falls_back_before_native_request() -> None:
    assert resolve_approved_native_role(_slide("市場分析と競合環境"), 2) is None


def test_evidence_sensitive_role_is_blocked_with_current_data() -> None:
    prs = Presentation()
    trace = native_trace_renderers.dispatch_approved_native_slide(
        prs,
        _slide("KPI設計と効果測定"),
        SimpleNamespace(),
        SimpleNamespace(),
        7,
        role="KPI",
        surface="summary",
    )
    assert trace["NATIVE_RENDERED"] is False
    assert trace["FALLBACK_USED"] is True
    assert trace["FAILURE_REASON"] == FailureReason.EVIDENCE_REQUIRED.value


def test_sample_leak_regression_blocks_kpi_before_package_creation() -> None:
    result = render_native_role_dry_run("KPI", data=SimpleNamespace(), context=SimpleNamespace(), slide=_slide("KPI設計と効果測定"), surface="summary")
    assert result.success is False
    assert result.failure_reason == FailureReason.EVIDENCE_REQUIRED.value
    assert result.package_path is None


def test_text_fit_failure_is_reported_before_native_render() -> None:
    long_title = "長いタイトル" * 20
    result = render_native_role_dry_run("PROPOSAL_SUMMARY", data=SimpleNamespace(), context=SimpleNamespace(), slide=_slide(long_title), surface="summary")
    assert result.success is False
    assert result.failure_reason == FailureReason.TEXT_OVERFLOW_RISK.value


def test_missing_template_falls_back(monkeypatch) -> None:
    original = native_trace_renderers.get_runtime_native_role_spec
    monkeypatch.setattr(
        native_trace_renderers,
        "get_runtime_native_role_spec",
        lambda *args, **kwargs: {"role": "PROPOSAL_SUMMARY", "slide_id": "S02", "runtime_asset": "missing/native.pptx", "human_approved": True},
    )
    prs = Presentation()
    trace = native_trace_renderers.dispatch_approved_native_slide(prs, _slide("提案サマリー"), SimpleNamespace(), SimpleNamespace(), 1, role="PROPOSAL_SUMMARY", surface="summary")
    assert trace["NATIVE_RENDERED"] is False
    assert trace["FALLBACK_USED"] is True
    assert trace["FAILURE_REASON"] == "TEMPLATE_AVAILABLE"
    monkeypatch.setattr(native_trace_renderers, "get_runtime_native_role_spec", original)


def test_renderer_exception_isolated(monkeypatch) -> None:
    def explode(*args, **kwargs):
        raise RuntimeError("synthetic native error")

    monkeypatch.setattr(native_trace_renderers, "render_native_role_dry_run", explode)
    trace = native_trace_renderers.dispatch_approved_native_slide(Presentation(), _slide("提案サマリー"), SimpleNamespace(), SimpleNamespace(), 1, role="PROPOSAL_SUMMARY", surface="summary")
    assert trace["NATIVE_RENDERED"] is False
    assert trace["FALLBACK_USED"] is True
    assert trace["FAILURE_REASON"] == FailureReason.TEMPLATE_CLONE_FAILED.value


def test_evidence_safe_visual_fallback_continues_deck_generation(monkeypatch) -> None:
    calls = {"legacy": 0}

    def legacy(*args, **kwargs):
        calls["legacy"] += 1
        args[0].slides.add_slide(args[0].slide_layouts[6])

    monkeypatch.setattr(slides, "settings", replace(settings, pptx_approved_native_renderer_enabled=True))
    monkeypatch.setattr(slides, "render_v5_masterpiece_slide", legacy)
    prs = Presentation()
    slides.add_designed_slide(prs, _slide("KPI設計と効果測定"), SimpleNamespace(), 7, SimpleNamespace(), surface="summary")
    assert calls["legacy"] == 0
    assert len(prs.slides) == 1


def test_summary_and_detail_paths_keep_existing_slide_counts(monkeypatch) -> None:
    monkeypatch.setattr(slides, "settings", replace(settings, pptx_approved_native_renderer_enabled=True))
    summary = Presentation(BytesIO(build_pptx_result(_payload(True), summary_mode=True).pptx_bytes))
    detail = Presentation(BytesIO(build_pptx_result(_payload(False), summary_mode=False).pptx_bytes))
    assert len(summary.slides) == 11
    assert len(detail.slides) > 0


def test_runtime_source_template_checksum_is_unchanged() -> None:
    registry = load_runtime_native_registry()
    role = next(item for item in registry["roles"] if item["role"] == "PROPOSAL_SUMMARY" and item["surface"] == "summary")
    source = ROOT / role["runtime_asset"]
    before = sha256_file(source)
    result = render_native_role_dry_run("PROPOSAL_SUMMARY", data=SimpleNamespace(), context=SimpleNamespace(), slide=_slide("提案サマリー"), surface="summary")
    assert result.success is True
    assert sha256_file(source) == before
