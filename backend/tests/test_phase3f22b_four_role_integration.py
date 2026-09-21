"""Focused Phase3F-22B tests for the four Human-approved native roles."""

from __future__ import annotations

from types import SimpleNamespace
from dataclasses import replace
import unittest
from unittest.mock import patch

from pptx import Presentation

from app.config import settings
from app.services.pptx_parts import slides
from app.services.pptx_parts.native_trace_content_adapter import (
    FailureReason,
    adapt_role,
    render_native_role_dry_run,
)
from app.services.pptx_parts.native_trace_registry import (
    canonical_runtime_role,
    get_runtime_native_role_spec,
    runtime_native_role_gate,
)
from app.services.pptx_parts.native_trace_renderers import dispatch_approved_native_slide


def _candidate(semantic_type: str) -> dict[str, object]:
    return {
        "id": f"phase3f22b:{semantic_type}",
        "semantic_type": semantic_type,
        "source_type": "user_input",
        "authority": "USER_PROVIDED",
        "review_state": "CONFIRMED",
        "admissible_as_evidence": True,
        "source_reference": f"synthetic://phase3f22b/{semantic_type}",
        "source_field": f"fixture.{semantic_type}",
    }


def _slide(title: str) -> SimpleNamespace:
    return SimpleNamespace(title=title, bullets=[])


def _competition_context() -> SimpleNamespace:
    return SimpleNamespace(
        competitor_a_name="比較対象A",
        competitor_b_name="比較対象B",
        competitor_rows=[
            {"criterion": "導入スピード", "own_value": "検証済み3か月", "competitor_a_value": "検証済み4〜5か月", "competitor_b_value": "検証済み3〜4か月", "evaluation": "条件確認済み"},
            {"criterion": "提案の具体性", "own_value": "検証済みKPI・運用設計", "competitor_a_value": "検証済み概略中心", "competitor_b_value": "検証済み一部詳細", "evaluation": "条件確認済み"},
            {"criterion": "AI活用実績", "own_value": "検証済み実績あり", "competitor_a_value": "検証済み実績限定", "competitor_b_value": "検証済みPoC中心", "evaluation": "条件確認済み"},
            {"criterion": "運用伴走", "own_value": "検証済み定例支援", "competitor_a_value": "検証済み限定対応", "competitor_b_value": "検証済み別契約", "evaluation": "条件確認済み"},
            {"criterion": "コスト透明性", "own_value": "検証済み前提明示", "competitor_a_value": "検証済み粒度粗い", "competitor_b_value": "検証済み追加条件", "evaluation": "条件確認済み"},
        ],
        differentiation_bullets=["検証済み支援範囲", "検証済み運用設計", "検証済み費用条件", "検証済み改善支援"],
        caution_items=["検証済みの比較条件を確認", "未確認事項は次回確認"],
        next_confirmation_items=["決裁条件を確認", "契約条件を確認", "実施時期を確認"],
        recommendation="検証済み比較根拠に基づき、条件を確認して判断します。",
    )


def _win_context() -> SimpleNamespace:
    return SimpleNamespace(
        win_probability=SimpleNamespace(
            probability=68,
            period="2026年下期",
            confidence="中",
            confidence_summary="検証済みの中程度の確度です。",
            reason="検証済みの課題適合性と決裁者接点を確認しています。",
            positive_factors=["検証済み課題", "検証済み決裁者接点", "検証済み予算感", "検証済み導入時期"],
            risk_factors=["予算承認", "競合比較", "体制確認", "開始時期"],
            next_action_cards=[
                {"title": "決裁条件の確認", "body": "検証済みの決裁条件を確認します。"},
                {"title": "予算の合意", "body": "検証済みの予算条件を合意します。"},
                {"title": "導入時期の確定", "body": "検証済みの導入時期を確定します。"},
            ],
            decision_direction="次回確認で導入条件を確定します。",
            evidence_rows=[
                {"item": "課題の明確性", "content": "検証済み課題を確認", "weight": "高", "evaluation": "強い"},
                {"item": "決裁者接点", "content": "検証済み接点を確認", "weight": "高", "evaluation": "強い"},
                {"item": "予算感", "content": "検証済み予算レンジ", "weight": "中", "evaluation": "普通"},
                {"item": "導入時期", "content": "検証済み時期候補", "weight": "中", "evaluation": "普通"},
                {"item": "競合優位性", "content": "検証済み比較根拠", "weight": "中", "evaluation": "確認中"},
            ],
        ),
    )


def _schedule_context() -> SimpleNamespace:
    durations = ["2週間", "3週間", "4週間", "2週間", "継続"]
    phases = [
        {"duration": duration, "owner": f"検証担当{index}", "milestone": f"確認マイルストーン{index}"}
        for index, duration in enumerate(durations, start=1)
    ]
    return SimpleNamespace(
        verified_schedule_phases=phases,
        verified_schedule_start_date="2026.09.01",
    )


def _roadmap_context() -> SimpleNamespace:
    phases = [
        {
            "number": index,
            "name": name,
            "duration": f"{index + 2}週間",
            "description": f"第{index}段階の実施内容を確認します。",
            "actions": [f"確認項目{index}-{action}" for action in range(1, 5)],
        }
        for index, name in enumerate(("準備", "構築", "試行", "展開", "定着"), start=1)
    ]
    return SimpleNamespace(
        verified_roadmap={
            "phases": phases,
            "milestones": [
                {"title": f"確認{index}", "timing": timing}
                for index, timing in enumerate(("Week 2", "Week 7", "Week 9", "Week 14", "継続確認"), start=1)
            ],
            "success_factors": [
                {"title": f"成功要因{index}", "body": "検証済みの推進条件です。"}
                for index in range(1, 5)
            ],
            "objective": "検証済みの計画に基づき、継続的な価値を創出します。",
        },
    )


def _data(*semantic_types: str) -> SimpleNamespace:
    return SimpleNamespace(semantic_candidates=[_candidate(value) for value in semantic_types])


class FourRoleIntegrationTests(unittest.TestCase):
    def test_aliases_resolve_to_one_runtime_identity(self) -> None:
        self.assertEqual(canonical_runtime_role("COMPETITIVE_COMPARISON"), "COMPETITION")
        self.assertEqual(canonical_runtime_role("SCHEDULE_GOVERNANCE"), "SCHEDULE")
        self.assertEqual(canonical_runtime_role("IMPLEMENTATION_SCHEDULE"), "ROADMAP")
        self.assertEqual(canonical_runtime_role("WIN_PROBABILITY"), "WIN_PROBABILITY")

    def test_four_runtime_registry_entries_are_human_approved(self) -> None:
        for role, surface in (
            ("COMPETITION", "conditional"),
            ("WIN_PROBABILITY", "conditional"),
            ("SCHEDULE", "summary"),
            ("ROADMAP", "conditional"),
        ):
            spec = get_runtime_native_role_spec(role, surface=surface)
            self.assertIsNotNone(spec)
            self.assertIs(spec["human_approved"], True)
            self.assertEqual(spec["registry_status"], "FOUNDATION_ONLY_NOT_ACTIVATED")
            self.assertTrue(spec["source_frozen"])

    def test_current_data_fails_closed_for_all_four_roles(self) -> None:
        cases = (
            ("COMPETITIVE_COMPARISON", "conditional", _slide("競合比較と差別化ポイント")),
            ("WIN_PROBABILITY", "conditional", _slide("受注確度と次の判断")),
            ("SCHEDULE_GOVERNANCE", "summary", _slide("導入スケジュールと推進体制")),
            ("IMPLEMENTATION_SCHEDULE", "conditional", _slide("導入ステップとスケジュール")),
        )
        for role, surface, slide in cases:
            payload = adapt_role(role, surface=surface, slide=slide)
            self.assertFalse(payload.renderable, role)
            self.assertEqual(payload.failure_reason, FailureReason.EVIDENCE_REQUIRED.value, role)

    def test_verified_fixtures_native_render_all_four(self) -> None:
        cases = (
            ("COMPETITIVE_COMPARISON", "conditional", _data("competition evidence"), _competition_context(), _slide("競合比較と差別化ポイント")),
            ("WIN_PROBABILITY", "conditional", _data("win probability"), _win_context(), _slide("受注確度と次の判断")),
            ("SCHEDULE_GOVERNANCE", "summary", _data("schedule evidence"), _schedule_context(), _slide("導入スケジュールと推進体制")),
            ("IMPLEMENTATION_SCHEDULE", "conditional", _data("roadmap implementation"), _roadmap_context(), _slide("導入ステップとスケジュール")),
        )
        for role, surface, data, context, slide in cases:
            result = render_native_role_dry_run(role, data=data, context=context, slide=slide, surface=surface)
            self.assertTrue(result.success, f"{role}: {result.reason}; {result.failure_reason}")
            self.assertEqual(result.validation["injected"]["prohibited_sample_strings_found"], [], role)
            if role == "COMPETITIVE_COMPARISON":
                self.assertEqual(len(result.payload.table_cell_bindings["trace:COMPETITION:title.primary.4"]), 22)
                self.assertIn("trace:COMPETITION:content.auto.36", result.injected_slots)
            if role == "WIN_PROBABILITY":
                self.assertEqual(len(result.payload.table_cell_bindings["trace:WIN_PROBABILITY:content.auto.30"]), 20)
                self.assertIn("trace:WIN_PROBABILITY:content.auto.84", result.injected_slots)

    def test_incomplete_verified_competition_fixture_falls_back_before_native(self) -> None:
        context = _competition_context()
        context.differentiation_bullets = ["検証済み支援範囲"]
        payload = adapt_role(
            "COMPETITIVE_COMPARISON",
            surface="conditional",
            data=_data("competition evidence"),
            context=context,
            slide=_slide("競合比較と差別化ポイント"),
        )
        self.assertFalse(payload.renderable)
        self.assertEqual(payload.failure_reason, FailureReason.UNBOUND_REQUIRED_CONTENT.value)

    def test_incomplete_verified_win_fixture_falls_back_before_native(self) -> None:
        context = _win_context()
        context.win_probability.next_action_cards = context.win_probability.next_action_cards[:2]
        payload = adapt_role(
            "WIN_PROBABILITY",
            surface="conditional",
            data=_data("win probability"),
            context=context,
            slide=_slide("受注確度と次の判断"),
        )
        self.assertFalse(payload.renderable)
        self.assertEqual(payload.failure_reason, FailureReason.UNBOUND_REQUIRED_CONTENT.value)

    def test_actual_dispatcher_renders_verified_fixture_and_keeps_editable_slide(self) -> None:
        prs = Presentation()
        cases = (
            ("COMPETITIVE_COMPARISON", "conditional", _data("competition evidence"), _competition_context(), _slide("競合比較と差別化ポイント")),
            ("WIN_PROBABILITY", "conditional", _data("win probability"), _win_context(), _slide("受注確度と次の判断")),
            ("SCHEDULE_GOVERNANCE", "summary", _data("schedule evidence"), _schedule_context(), _slide("導入スケジュールと推進体制")),
            ("IMPLEMENTATION_SCHEDULE", "conditional", _data("roadmap implementation"), _roadmap_context(), _slide("導入ステップとスケジュール")),
        )
        for index, (role, surface, data, context, slide) in enumerate(cases, start=1):
            trace = dispatch_approved_native_slide(prs, slide, data, context, index, role=role, surface=surface)
            self.assertTrue(trace["NATIVE_RENDERED"], trace)
            self.assertFalse(trace["FALLBACK_USED"], trace)
            self.assertEqual(trace["SAMPLE_LEAK_STATUS"], "PASS", trace)
        self.assertEqual(len(prs.slides), 4)

    def test_runtime_gate_fails_closed_when_template_is_missing(self) -> None:
        eligible, checks = runtime_native_role_gate("IMPLEMENTATION_SCHEDULE", surface="conditional", template_available=False)
        self.assertFalse(eligible)
        self.assertFalse(checks["TEMPLATE_AVAILABLE"])

    def test_generated_only_fixture_is_rejected_for_all_four_roles(self) -> None:
        generated = SimpleNamespace(
            semantic_candidates=[
                {
                    **_candidate("four-role generated evidence"),
                    "source_type": "ai_generated",
                    "authority": "AI_PROPOSED",
                    "review_state": "UNCONFIRMED",
                    "admissible_as_evidence": False,
                    "inferred": True,
                }
            ]
        )
        contexts = (
            ("COMPETITIVE_COMPARISON", "conditional", _competition_context(), _slide("競合比較と差別化ポイント")),
            ("WIN_PROBABILITY", "conditional", _win_context(), _slide("受注確度と次の判断")),
            ("SCHEDULE_GOVERNANCE", "summary", _schedule_context(), _slide("導入スケジュールと推進体制")),
            ("IMPLEMENTATION_SCHEDULE", "conditional", _roadmap_context(), _slide("導入ステップとスケジュール")),
        )
        for role, surface, context, slide in contexts:
            payload = adapt_role(role, surface=surface, data=generated, context=context, slide=slide)
            self.assertFalse(payload.renderable, role)
            self.assertEqual(payload.failure_reason, FailureReason.EVIDENCE_REQUIRED.value, role)

    def test_text_fit_blocks_long_title_before_native_render(self) -> None:
        cases = (
            ("COMPETITIVE_COMPARISON", "conditional", _data("competition evidence"), _competition_context()),
            ("WIN_PROBABILITY", "conditional", _data("win probability"), _win_context()),
            ("SCHEDULE_GOVERNANCE", "summary", _data("schedule evidence"), _schedule_context()),
            ("IMPLEMENTATION_SCHEDULE", "conditional", _data("roadmap implementation"), _roadmap_context()),
        )
        for role, surface, data, context in cases:
            result = render_native_role_dry_run(
                role,
                data=data,
                context=context,
                slide=_slide("長いタイトル" * 20),
                surface=surface,
            )
            self.assertFalse(result.success, role)
            self.assertEqual(result.failure_reason, FailureReason.TEXT_OVERFLOW_RISK.value, role)

    def test_native_exception_isolated_to_role(self) -> None:
        from app.services.pptx_parts import native_trace_renderers

        with patch.object(native_trace_renderers, "render_native_role_dry_run", side_effect=RuntimeError("synthetic")):
            trace = dispatch_approved_native_slide(
                Presentation(),
                _slide("導入ステップとスケジュール"),
                _data("roadmap implementation"),
                _roadmap_context(),
                1,
                role="IMPLEMENTATION_SCHEDULE",
                surface="conditional",
            )
        self.assertFalse(trace["NATIVE_RENDERED"])
        self.assertTrue(trace["FALLBACK_USED"])
        self.assertEqual(trace["FAILURE_REASON"], FailureReason.TEMPLATE_CLONE_FAILED.value)

    def test_flag_off_preserves_existing_renderer_path(self) -> None:
        prs = Presentation()
        legacy_calls = []

        def legacy(prs_arg, *args, **kwargs):
            legacy_calls.append(True)
            prs_arg.slides.add_slide(prs_arg.slide_layouts[6])

        with patch.object(slides, "settings", replace(settings, pptx_approved_native_renderer_enabled=False)):
            with patch.object(slides, "dispatch_approved_native_slide", side_effect=AssertionError("native must not be requested")):
                with patch.object(slides, "render_v5_masterpiece_slide", side_effect=legacy):
                    slides.add_designed_slide(
                        prs,
                        _slide("競合比較と差別化ポイント"),
                        _data("competition evidence"),
                        1,
                        _competition_context(),
                        surface="conditional",
                    )
        self.assertEqual(legacy_calls, [True])
        self.assertEqual(len(prs.slides), 1)

    def test_existing_add_designed_slide_insertion_point_uses_native_when_flag_on(self) -> None:
        prs = Presentation()
        with patch.object(slides, "settings", replace(settings, pptx_approved_native_renderer_enabled=True)):
            slides.add_designed_slide(
                prs,
                _slide("競合比較と差別化ポイント"),
                _data("competition evidence"),
                1,
                _competition_context(),
                surface="conditional",
            )
        self.assertEqual(len(prs.slides), 1)


if __name__ == "__main__":
    unittest.main()
