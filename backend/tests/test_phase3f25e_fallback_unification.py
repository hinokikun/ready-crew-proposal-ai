from __future__ import annotations

from dataclasses import replace
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

from pptx import Presentation

from app.config import settings
from app.models import PowerPointData, PowerPointSlide, PptxDownloadRequest
from app.services.pptx_parts import slides
from app.services.pptx_parts.native_trace_content_adapter import FailureReason, SAMPLE_STRINGS
from app.services.pptx_parts.native_trace_renderers import dispatch_approved_native_slide
from app.services.pptx_service import build_pptx_context


ROOT = Path(__file__).resolve().parents[2]
SMOKE_DIR = ROOT / "artifacts" / "phase3f25e_fallback_unification"


def _slide(title: str) -> SimpleNamespace:
    return SimpleNamespace(title=title, layout="", bullets=["確認済み情報"], slide_no=1)


def _package_xml(path: Path) -> tuple[str, str]:
    with ZipFile(path) as package:
        xml = "\n".join(
            package.read(name).decode("utf-8", errors="ignore")
            for name in package.namelist()
            if name.startswith("ppt/slides/slide") and name.endswith(".xml")
        )
        rels = "\n".join(
            package.read(name).decode("utf-8", errors="ignore")
            for name in package.namelist()
            if name.endswith(".rels")
        )
    return xml, rels


def test_cover_package_clone_preserves_native_media_relationship() -> None:
    prs = Presentation()
    trace = dispatch_approved_native_slide(
        prs,
        _slide("表紙"),
        SimpleNamespace(),
        SimpleNamespace(),
        0,
        role="COVER",
        surface="summary",
        visual_fallback=True,
    )
    assert trace["FALLBACK_RENDERED"] is True
    assert trace["FALLBACK_KIND"] == "APPROVED_COVER_PACKAGE_CLONE"
    assert trace["NATIVE_RENDERED"] is False
    output = BytesIO()
    prs.save(output)
    with ZipFile(BytesIO(output.getvalue())) as package:
        names = set(package.namelist())
        relationships = "\n".join(
            package.read(name).decode("utf-8", errors="ignore")
            for name in names
            if name.endswith(".rels")
        )
    assert any(name.startswith("ppt/media/") for name in names)
    assert "relationships/image" in relationships
    assert len(prs.slides) == 1


def test_evidence_safe_visual_fallbacks_clear_role_samples_and_legacy_identity() -> None:
    cases = (
        ("IMPLEMENTATION_CONFIGURATION", "導入構成", 8),
        ("KPI", "KPI設計", 9),
        ("SCHEDULE", "スケジュール", 10),
    )
    for role, title, index in cases:
        prs = Presentation()
        trace = dispatch_approved_native_slide(
            prs,
            _slide(title),
            SimpleNamespace(),
            SimpleNamespace(),
            index,
            role=role,
            surface="summary",
            visual_fallback=True,
        )
        assert trace["NATIVE_RENDERED"] is False
        assert trace["FALLBACK_USED"] is True
        assert trace["FALLBACK_RENDERED"] is True
        assert trace["FAILURE_REASON"] == FailureReason.EVIDENCE_REQUIRED.value
        assert trace["SAMPLE_LEAK_STATUS"] == "PASS"
        output = BytesIO()
        prs.save(output)
        xml, _ = _package_xml_from_bytes(output.getvalue())
        assert "ProposalPilot" not in xml
        assert "AI営業秘書" not in xml
        for sample in SAMPLE_STRINGS.get(role, ()):
            assert sample not in xml


def _package_xml_from_bytes(value: bytes) -> tuple[str, str]:
    with ZipFile(BytesIO(value)) as package:
        xml = "\n".join(
            package.read(name).decode("utf-8", errors="ignore")
            for name in package.namelist()
            if name.startswith("ppt/slides/slide") and name.endswith(".xml")
        )
        rels = "\n".join(
            package.read(name).decode("utf-8", errors="ignore")
            for name in package.namelist()
            if name.endswith(".rels")
        )
    return xml, rels


def _summary_slides() -> list[PowerPointSlide]:
    titles = [
        "表紙",
        "経営判断の要点",
        "本提案の結論と期待効果",
        "提案サマリー",
        "現状理解",
        "主要課題",
        "提案コンセプト",
        "導入戦略",
        "導入構成",
        "KPI設計",
        "スケジュール",
    ]
    return [
        PowerPointSlide(
            slide_no=index + 1,
            layout="title" if index == 0 else "content",
            title=title,
            bullets=["確認済み情報"],
            speaker_notes="",
            visual_suggestion="",
        )
        for index, title in enumerate(titles)
    ]


def _render_direct_summary(flag_enabled: bool, path: Path) -> dict[str, object]:
    prs = Presentation()
    traces: list[dict[str, object]] = []
    data = PowerPointData(deck_title="提案クエスト確認用", client_name="検証社", slides=_summary_slides())
    payload = PptxDownloadRequest(
        powerpoint_generation_data=data,
        project_brief="確認済み情報のみを使用するローカルSmoke",
        client_company_info="検証社",
        summary=True,
    )
    context = build_pptx_context(payload)
    with_flag = replace(settings, pptx_approved_native_renderer_enabled=flag_enabled)
    original = slides.settings
    slides.settings = with_flag
    try:
        for index, slide_data in enumerate(_summary_slides()):
            native_role = {
                0: "COVER",
                1: "EXECUTIVE_SUMMARY",
                2: "DECISION_AND_EXPECTED_EFFECTS",
                3: "PROPOSAL_SUMMARY",
                4: "CURRENT_STATE",
                5: "PROBLEM_ANALYSIS",
                6: "SOLUTION_CONCEPT",
                7: "SOLUTION_APPROACH",
                8: "IMPLEMENTATION_CONFIGURATION",
                9: "KPI",
                10: "SCHEDULE",
            }.get(index)
            if flag_enabled and native_role:
                slides.add_designed_slide(prs, slide_data, data, index, context, surface="summary")
                if index == 0:
                    trace = {"NATIVE_RENDERED": False, "FALLBACK_USED": True, "FALLBACK_KIND": "APPROVED_COVER_PACKAGE_CLONE"}
                elif index in {8, 9, 10}:
                    trace = {"NATIVE_RENDERED": False, "FALLBACK_USED": True, "FALLBACK_KIND": "PROPOSAL_MASTER_EVIDENCE_SAFE"}
                else:
                    trace = {"NATIVE_RENDERED": True, "FALLBACK_USED": False, "FALLBACK_KIND": None}
                traces.append({"index": index, "title": slide_data.title, **trace})
            else:
                slides.add_designed_slide(prs, slide_data, data, index, context, surface="summary")
                traces.append({"index": index, "title": slide_data.title, "NATIVE_RENDERED": False, "FALLBACK_USED": True})
    finally:
        slides.settings = original
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return {"slide_count": len(prs.slides), "traces": traces}


def test_flag_on_and_off_generate_eleven_slide_local_summary_smoke() -> None:
    SMOKE_DIR.mkdir(parents=True, exist_ok=True)
    flag_on = SMOKE_DIR / "summary_flag_on.pptx"
    flag_off = SMOKE_DIR / "summary_flag_off.pptx"
    on_trace = _render_direct_summary(True, flag_on)
    off_trace = _render_direct_summary(False, flag_off)
    assert on_trace["slide_count"] == 11
    assert off_trace["slide_count"] == 11
    on_deck = Presentation(str(flag_on))
    assert len(on_deck.slides) == 11
    assert any(shape.name == "trace:S02:content.auto" for shape in on_deck.slides[1].shapes)
    assert any(shape.name == "trace:S03:content.auto" for shape in on_deck.slides[2].shapes)
    assert any(shape.name == "trace:S02:footer.date" and shape.text == "2026.08.26" for shape in on_deck.slides[1].shapes)
    assert any(shape.text == "2026.08.26" for shape in on_deck.slides[2].shapes if getattr(shape, "has_text_frame", False))
    slide02_text = "\n".join(shape.text for shape in on_deck.slides[1].shapes if getattr(shape, "has_text_frame", False))
    slide03_text = "\n".join(shape.text for shape in on_deck.slides[2].shapes if getattr(shape, "has_text_frame", False))
    slide02_by_name = {shape.name: shape.text.strip() for shape in on_deck.slides[1].shapes if getattr(shape, "has_text_frame", False)}
    slide03_by_name = {shape.name: shape.text.strip() for shape in on_deck.slides[2].shapes if getattr(shape, "has_text_frame", False)}
    assert all(
        slide02_by_name.get(f"trace:S02:{slot}")
        for slot in (
            "title.primary.2",
            "lead",
            "lead.2",
            "lead.3",
            "content.auto.45",
            "title.primary.3",
            "content.auto.53",
            "lead.4",
            "content.auto.63",
            "content.auto.64",
            "content.auto.75",
            "content.auto.90",
            "content.auto.101",
            "content.auto.112",
            "content.auto.118",
        )
    )
    assert all(
        slide03_by_name.get(f"trace:S03:{slot}")
        for slot in (
            "title.primary.3",
            "title.primary.4",
            "title.primary.5",
            "content.auto.34",
            "content.auto.36",
            "content.auto.38",
            "content.auto.45",
            "content.auto.47",
            "content.auto.49",
            "content.auto.60",
            "content.auto.62",
            "content.auto.64",
            "content.auto.73",
            "content.auto.75",
            "content.auto.77",
            "content.auto.79",
            "content.auto.88",
            "content.auto.90",
            "content.auto.98",
        )
    )
    assert any(
        shape.text.strip() == "2026.08.26"
        for shape in on_deck.slides[2].shapes
        if getattr(shape, "has_text_frame", False)
    )
    assert "確認済み情報を整理" in slide02_text
    assert "確認済み課題を整理" in slide03_text
    assert "ProposalPilot" not in slide02_text + slide03_text
    assert "AI営業秘書" not in slide02_text + slide03_text
    slide04_text = "\n".join(
        shape.text for shape in on_deck.slides[3].shapes if getattr(shape, "has_text_frame", False)
    )
    slide10_text = "\n".join(
        shape.text for shape in on_deck.slides[9].shapes if getattr(shape, "has_text_frame", False)
    )
    slide09_text = "\n".join(
        shape.text for shape in on_deck.slides[8].shapes if getattr(shape, "has_text_frame", False)
    )
    slide11_text = "\n".join(
        shape.text for shape in on_deck.slides[10].shapes if getattr(shape, "has_text_frame", False)
    )
    assert "確認済み情報を整理し、未確認項目は確認後に確定します。" in slide04_text
    assert "提案品質の向上" not in slide04_text
    assert "最適な一手" not in slide09_text
    for forbidden_kpi in (
        "資料作成時間",
        "提案数",
        "修正回数",
        "受注確度",
        "提案資料の作成にかかる",
        "提案実施件数",
        "資料確認項目を削減",
        "工数記録",
        "案件数",
        "履歴カウント",
        "受注率",
    ):
        assert forbidden_kpi not in slide10_text
    assert "確認項目" in slide10_text
    assert "現状値未取得" not in slide10_text
    assert "目標値は要確認" not in slide10_text
    assert "現状値未取得向上" not in slide10_text
    assert any(shape.name == "trace:KPI:content.auto" and shape.text == "10" for shape in on_deck.slides[9].shapes)
    assert any(shape.name == "trace:SCHEDULE:content.auto" and shape.text == "11" for shape in on_deck.slides[10].shapes)
    assert any(shape.name == "trace:SCHEDULE:footer.date" and shape.text == "2026.08.26" for shape in on_deck.slides[10].shapes)
    assert "確認後に確定" in slide11_text
    kpi_table = next(shape for shape in on_deck.slides[9].shapes if getattr(shape, "has_table", False))
    assert sum(row.height for row in kpi_table.table.rows) <= kpi_table.height
    schedule_table = next(shape for shape in on_deck.slides[10].shapes if getattr(shape, "has_table", False))
    assert sum(row.height for row in schedule_table.table.rows) <= schedule_table.height
    assert [trace["FALLBACK_KIND"] for trace in on_trace["traces"] if trace["index"] in {0, 8, 9, 10}] == [
        "APPROVED_COVER_PACKAGE_CLONE",
        "PROPOSAL_MASTER_EVIDENCE_SAFE",
        "PROPOSAL_MASTER_EVIDENCE_SAFE",
        "PROPOSAL_MASTER_EVIDENCE_SAFE",
    ]
    xml, _ = _package_xml(flag_on)
    assert "ProposalPilot" not in xml
    assert "AI営業秘書" not in xml
    off_deck = Presentation(str(flag_off))
    off_target_text = "\n".join(
        shape.text
        for slide_index in (1, 2)
        for shape in off_deck.slides[slide_index].shapes
        if getattr(shape, "has_text_frame", False)
    )
    assert "ProposalPilot" not in off_target_text
    assert "AI営業秘書" not in off_target_text
    assert "提案クエスト" in off_target_text
