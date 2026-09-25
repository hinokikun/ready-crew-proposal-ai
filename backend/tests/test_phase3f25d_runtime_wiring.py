from __future__ import annotations

from dataclasses import replace
from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from pptx import Presentation

from app.config import settings
from app.services.pptx_parts import native_trace_renderers, slides
from app.services.pptx_parts.native_trace_content_adapter import FailureReason, render_native_role_dry_run
from app.services.pptx_parts.native_trace_registry import resolve_approved_native_role


def _slide(title: str) -> SimpleNamespace:
    return SimpleNamespace(title=title, layout="", bullets=[], slide_no=1)


def _verified_configuration_context(*, omit: str | None = None) -> SimpleNamespace:
    values = {
        "implementation_summary": ["FAJの受注プロセスを案件単位で可視化し、確認済みの運用範囲から段階導入します。", "FAJの確認済み運用情報を基に、案件改善を継続します。"],
        "implementation_scope": ["営業管理", "受注承認", "在庫可視化", "入出庫管理", "購買承認", "支払管理", "製造計画", "工程進捗", "財務報告", "予算管理"],
        "integration_targets": ["顧客マスタ", "商品API", "請求サービス", "決済ゲートウェイ", "人事基盤", "勤怠連携", "BI基盤", "経営ダッシュボード"],
        "implementation_architecture": ["案件データセット", "確認済みオペレーション標準化", "FAJ案件基盤", "案件データベース", "確認済み業務フロー", "確認済みマスタ", "FAJ権限管理", "FAJ案件基盤", "確認済みデータ連係", "確認済み外部連係"],
        "estimate": ["980 万円〜", "年間 240 万円〜", "280 万円〜", "1,520 万円〜"],
        "roi": ["約 3,200 時間", "約 540 万円", "約 1.8 年", "約 1,600 万円"],
        "schedule": ["2026年10月1日", "要件確認", "5週間", "構築", "10週間", "検証", "本番開始", "2週間"],
    }
    candidates = [
        {
            "semantic_type": key,
            "value": value,
            "review_state": "CONFIRMED",
            "authority": "EXTERNAL_VERIFIED",
            "admissible_as_evidence": True,
            "source_reference": f"fixture:{key}",
        }
        for key, value in values.items()
        if key != omit
    ]
    return SimpleNamespace(semantic_candidates=candidates)


def _verified_semantic_context(role: str) -> SimpleNamespace:
    values = {
        "EXECUTIVE_SUMMARY": {
            "summary_context": ["FAJ案件の確認済み論点を整理し、判断材料を揃えます。"],
            "background": ["FAJで確認済みの業務課題を優先して整理します。"],
            "current_state": ["FAJの現状データで確認された分散課題を整理します。"],
            "conclusion": ["確認済みのFAJ課題に対する実行可能な方針を選定します。"],
            "expected_effects": ["確認済みの工数課題を可視化", "判断材料を整理", "運用確認を継続"],
            "decision_items": ["対象範囲を確認", "測定方法を確認", "費用条件を確認", "推進体制を確認"],
            "insight": ["確認済みの範囲から段階的に判断します。"],
        },
        "DECISION_AND_EXPECTED_EFFECTS": {
            "summary": ["FAJ向けに確認済みの課題と方針を整理した提案です。"],
            "headline": ["FAJで確認された課題を整理し", "検証可能な判断と次の行動へつなげます"],
            "why_now": ["FAJで確認された業務上の課題", "確認済みデータの分散", "次の判断に必要な論点"],
            "policy": ["確認済みの範囲から整理", "実測可能な項目を定義", "段階的に検証"],
            "investment": ["確認済みの対象範囲", "必要作業と移行条件", "運用確認と継続改善"],
            "effects": ["確認済み課題の可視化", "判断材料の整理", "運用品質の確認", "次段階の条件整理"],
            "decisions": ["対象範囲の確認", "次段階の実施条件の確認"],
            "insight": ["確認済みの事実を基に次の判断へ進みます。"],
        },
    }[role]
    candidates = [
        {
            "semantic_type": f"{role.lower()}.{key}",
            "value": value,
            "review_state": "CONFIRMED",
            "authority": "USER_EXPLICIT",
            "admissible_as_evidence": True,
            "source_reference": f"faj://verified/{key}",
            "source_field": f"faj.{key}",
        }
        for key, value in values.items()
    ]
    return SimpleNamespace(semantic_candidates=candidates)


def test_v11_roles_resolve_without_replacing_existing_roles() -> None:
    assert resolve_approved_native_role(_slide("経営判断の要点"), 1) == "EXECUTIVE_SUMMARY"
    assert resolve_approved_native_role(_slide("本提案の結論と期待効果"), 2) == "DECISION_AND_EXPECTED_EFFECTS"
    assert resolve_approved_native_role(_slide("導入構成"), 8) == "IMPLEMENTATION_CONFIGURATION"
    assert resolve_approved_native_role(_slide("提案サマリー"), 3) == "PROPOSAL_SUMMARY"
    assert resolve_approved_native_role(_slide("現状理解"), 4) == "CURRENT_STATE"
    assert resolve_approved_native_role(_slide("システム構成"), 6) == "SYSTEM_ARCHITECTURE"


@pytest.mark.parametrize(
    ("role", "title", "slide_id"),
    [
        ("EXECUTIVE_SUMMARY", "経営判断の要点", "S02"),
        ("DECISION_AND_EXPECTED_EFFECTS", "本提案の結論と期待効果", "S03"),
    ],
)
def test_v11_role_requires_explicit_semantic_binding(role: str, title: str, slide_id: str) -> None:
    result = render_native_role_dry_run(role, data=SimpleNamespace(), context=_verified_semantic_context(role), slide=_slide(title), surface="summary", slide_id=slide_id)
    assert result.success is True
    assert result.validation["injected"]["sample_content_leak_free"] is True
    assert result.payload and result.payload.evidence_status
    assert result.pptx_bytes is not None
    with ZipFile(BytesIO(result.pptx_bytes)) as package:
        package_text = "\n".join(
            package.read(name).decode("utf-8", errors="ignore")
            for name in package.namelist()
            if name.endswith(".xml")
        )
    for sample in result.payload.prohibited_sample_strings:
        assert sample not in package_text

    missing = render_native_role_dry_run(role, data=SimpleNamespace(), context=SimpleNamespace(), slide=_slide(title), surface="summary", slide_id=slide_id)
    assert missing.success is False
    assert missing.failure_reason == FailureReason.EVIDENCE_REQUIRED.value
    unsupported = render_native_role_dry_run(role, data=SimpleNamespace(), context=SimpleNamespace(), slide=_slide("unsupported runtime title"), surface="summary", slide_id=slide_id)
    assert unsupported.success is False
    assert unsupported.failure_reason == FailureReason.UNBOUND_REQUIRED_CONTENT.value


def test_implementation_configuration_verified_evidence_uses_native_path_without_trace_samples() -> None:
    result = render_native_role_dry_run(
        "IMPLEMENTATION_CONFIGURATION",
        data=SimpleNamespace(),
        context=_verified_configuration_context(),
        slide=_slide("導入構成"),
        surface="summary",
        slide_id="S09",
    )
    assert result.success is True
    assert result.validation["injected"]["sample_content_leak_free"] is True
    assert result.pptx_bytes is not None
    with ZipFile(BytesIO(result.pptx_bytes)) as package:
        package_text = "\n".join(
            package.read(name).decode("utf-8", errors="ignore")
            for name in package.namelist()
            if name.endswith(".xml")
        )
    for sample in result.payload.prohibited_sample_strings:
        assert sample not in package_text
    trace = native_trace_renderers.dispatch_approved_native_slide(
        Presentation(),
        _slide("導入構成"),
        SimpleNamespace(),
        _verified_configuration_context(),
        8,
        role="IMPLEMENTATION_CONFIGURATION",
        surface="summary",
    )
    assert trace["NATIVE_RENDERED"] is True


@pytest.mark.parametrize("missing", ["implementation_scope", "integration_targets", "estimate", "roi", "schedule"])
def test_implementation_configuration_missing_verified_evidence_falls_back(missing: str) -> None:
    result = render_native_role_dry_run(
        "IMPLEMENTATION_CONFIGURATION",
        data=SimpleNamespace(),
        context=_verified_configuration_context(omit=missing),
        slide=_slide("導入構成"),
        surface="summary",
        slide_id="S09",
    )
    assert result.success is False
    assert result.failure_reason == FailureReason.EVIDENCE_REQUIRED.value


def test_implementation_configuration_generated_only_evidence_is_rejected() -> None:
    context = _verified_configuration_context()
    context.semantic_candidates[0] = {
        **context.semantic_candidates[0],
        "authority": "AI_PROPOSED",
        "admissible_as_evidence": False,
        "review_state": "CONFIRMED",
    }
    result = render_native_role_dry_run(
        "IMPLEMENTATION_CONFIGURATION",
        data=SimpleNamespace(),
        context=context,
        slide=_slide("導入構成"),
        surface="summary",
        slide_id="S09",
    )
    assert result.success is False
    assert result.failure_reason == FailureReason.EVIDENCE_REQUIRED.value


def test_flag_on_dispatches_new_static_role_and_flag_off_keeps_legacy(monkeypatch) -> None:
    prs = Presentation()
    monkeypatch.setattr(slides, "settings", replace(settings, pptx_approved_native_renderer_enabled=True))
    native_trace = native_trace_renderers.dispatch_approved_native_slide(
        prs,
        _slide("経営判断の要点"),
        SimpleNamespace(),
        _verified_semantic_context("EXECUTIVE_SUMMARY"),
        1,
        role="EXECUTIVE_SUMMARY",
        surface="summary",
    )
    assert native_trace["NATIVE_RENDERED"] is True

    calls = {"native": 0, "legacy": 0}

    def native(*args, **kwargs):
        calls["native"] += 1
        return {"NATIVE_RENDERED": True}

    def legacy(*args, **kwargs):
        calls["legacy"] += 1
        args[0].slides.add_slide(args[0].slide_layouts[6])

    monkeypatch.setattr(slides, "settings", replace(settings, pptx_approved_native_renderer_enabled=False))
    monkeypatch.setattr(slides, "dispatch_approved_native_slide", native)
    monkeypatch.setattr(slides, "render_v5_masterpiece_slide", legacy)
    slides.add_designed_slide(prs, _slide("経営判断の要点"), SimpleNamespace(), 1, SimpleNamespace(), surface="summary")
    assert calls == {"native": 0, "legacy": 1}
